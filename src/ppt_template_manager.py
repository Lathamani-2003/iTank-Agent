from __future__ import annotations

import hashlib
from copy import deepcopy
import json
import math
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

EMU_PER_INCH = 914400.0

# Increment this whenever the PPT export implementation changes.  The value is
# included in the session signature so a running Streamlit session cannot keep
# serving PPT bytes built by an older exporter after the code is updated.
PPT_EXPORT_IMPLEMENTATION_VERSION = "2026-09-14-worksheet-price-details-pairs-v2"

DEFAULT_PLACEMENT_CONFIG: dict[str, Any] = {
    "slide_number": 7,
    "x": 0.55,
    "y": 1.20,
    "width": 12.23,
    "height": 5.75,
    "maintain_aspect_ratio": True,
    "horizontal_align": "center",
    "vertical_align": "middle",
}

# Small expansion used only while detecting the old diagram artwork in the
# fixed template. It allows connectors sitting a few pixels outside the saved
# diagram box to be removed with the rest of the old diagram.
_REPLACEMENT_TOLERANCE_IN = 0.18

# Header/footer/title elements are template artwork and must never be deleted.
# These margins are intentionally conservative; the generated diagram itself
# is still placed using the configured diagram box.
_TOP_TEMPLATE_SAFE_IN = 1.00
_BOTTOM_TEMPLATE_SAFE_IN = 0.85



def _coerce_float(value: Any, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default)



def _normalize_config(config: dict[str, Any] | None) -> dict[str, Any]:
    source = dict(DEFAULT_PLACEMENT_CONFIG)
    if isinstance(config, dict):
        source.update(config)

    try:
        slide_number = int(source.get("slide_number", DEFAULT_PLACEMENT_CONFIG["slide_number"]))
    except (TypeError, ValueError):
        slide_number = int(DEFAULT_PLACEMENT_CONFIG["slide_number"])

    horizontal_align = str(source.get("horizontal_align", "center") or "center").lower()
    if horizontal_align not in {"left", "center", "right"}:
        horizontal_align = "center"

    vertical_align = str(source.get("vertical_align", "middle") or "middle").lower()
    if vertical_align not in {"top", "middle", "bottom"}:
        vertical_align = "middle"

    return {
        "slide_number": max(1, slide_number),
        "x": max(0.0, _coerce_float(source.get("x"), DEFAULT_PLACEMENT_CONFIG["x"])),
        "y": max(0.0, _coerce_float(source.get("y"), DEFAULT_PLACEMENT_CONFIG["y"])),
        "width": max(0.10, _coerce_float(source.get("width"), DEFAULT_PLACEMENT_CONFIG["width"])),
        "height": max(0.10, _coerce_float(source.get("height"), DEFAULT_PLACEMENT_CONFIG["height"])),
        "maintain_aspect_ratio": bool(source.get("maintain_aspect_ratio", True)),
        "horizontal_align": horizontal_align,
        "vertical_align": vertical_align,
    }



def load_placement_config(config_path: Path) -> dict[str, Any]:
    """Load the persistent fixed-template placement settings."""
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, TypeError):
        data = dict(DEFAULT_PLACEMENT_CONFIG)
    return _normalize_config(data)



def save_placement_config(config_path: Path, config: dict[str, Any]) -> dict[str, Any]:
    """Persist normalized placement settings without touching the PPT template."""
    normalized = _normalize_config(config)
    config_path.parent.mkdir(parents=True, exist_ok=True)
    config_path.write_text(json.dumps(normalized, indent=2), encoding="utf-8")
    return normalized



def get_template_metadata(template_path: Path) -> dict[str, Any]:
    """Return basic fixed-template metadata required by the placement editor."""
    if not template_path.exists():
        raise FileNotFoundError(f"Fixed PowerPoint template not found: {template_path}")
    presentation = Presentation(str(template_path))
    return {
        "slide_count": len(presentation.slides),
        "slide_width": float(presentation.slide_width) / EMU_PER_INCH,
        "slide_height": float(presentation.slide_height) / EMU_PER_INCH,
    }



def fixed_template_signature(template_path: Path, config_path: Path) -> str:
    """Hash only PPT-output inputs; diagram generation does not depend on this."""
    digest = hashlib.sha256()
    if template_path.exists():
        digest.update(template_path.read_bytes())
    else:
        digest.update(b"missing-fixed-template")

    config = load_placement_config(config_path)
    digest.update(json.dumps(config, sort_keys=True, separators=(",", ":")).encode("utf-8"))
    # Also invalidate the in-memory/session PPT whenever the exporter code changes.
    # Without this, an unchanged worksheet PNG can cause app.py to reuse the old
    # ppt_session_pptx_bytes forever, making a correct exporter update look like
    # it had "no change".
    digest.update(PPT_EXPORT_IMPLEMENTATION_VERSION.encode("utf-8"))
    return digest.hexdigest()



def _fit_inside_box(
    source_width: int,
    source_height: int,
    box_x: float,
    box_y: float,
    box_width: float,
    box_height: float,
    horizontal_align: str,
    vertical_align: str,
) -> tuple[float, float, float, float]:
    if source_width <= 0 or source_height <= 0:
        return box_x, box_y, box_width, box_height

    source_ratio = float(source_width) / float(source_height)
    box_ratio = float(box_width) / float(box_height)

    if source_ratio >= box_ratio:
        picture_width = box_width
        picture_height = box_width / source_ratio
    else:
        picture_height = box_height
        picture_width = box_height * source_ratio

    if horizontal_align == "left":
        picture_x = box_x
    elif horizontal_align == "right":
        picture_x = box_x + box_width - picture_width
    else:
        picture_x = box_x + (box_width - picture_width) / 2.0

    if vertical_align == "top":
        picture_y = box_y
    elif vertical_align == "bottom":
        picture_y = box_y + box_height - picture_height
    else:
        picture_y = box_y + (box_height - picture_height) / 2.0

    return picture_x, picture_y, picture_width, picture_height



def _shape_bounds_in(shape: Any) -> tuple[float, float, float, float]:
    left = float(shape.left) / EMU_PER_INCH
    top = float(shape.top) / EMU_PER_INCH
    width = float(shape.width) / EMU_PER_INCH
    height = float(shape.height) / EMU_PER_INCH
    return left, top, left + width, top + height



def _boxes_overlap(
    a: tuple[float, float, float, float],
    b: tuple[float, float, float, float],
) -> bool:
    return a[0] < b[2] and a[2] > b[0] and a[1] < b[3] and a[3] > b[1]



def _is_template_protected_shape(shape: Any, slide_height_in: float) -> bool:
    """Protect template-only artwork such as title/header/footer/logo/slide number."""
    name = str(getattr(shape, "name", "") or "").strip().lower()
    if any(token in name for token in ("title", "footer", "logo", "slide number", "date placeholder")):
        return True

    left, top, right, bottom = _shape_bounds_in(shape)
    width = max(0.0, right - left)
    height = max(0.0, bottom - top)
    center_y = (top + bottom) / 2.0

    # Header/title chrome is always protected.
    if center_y <= _TOP_TEMPLATE_SAFE_IN:
        return True

    # Preserve the bottom-right iTank logo/footer branding.
    if center_y >= max(0.0, slide_height_in - _BOTTOM_TEMPLATE_SAFE_IN):
        center_x = (left + right) / 2.0
        if center_x >= 11.0:
            return True

    # Protect large background/border artwork.  The example diagram content is
    # made from smaller shapes/images/connectors; the slide background/border,
    # if represented as a shape, typically covers most of the slide.
    if width >= 11.0 and height >= 5.5:
        return True

    # Ordinary text boxes are preserved. Diagram labels in the supplied fixed
    # template are AutoShapes, so removing them remains possible while keeping
    # unrelated template text completely untouched.
    if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return True

    # Placeholders belong to the template layout and are never removed.
    if bool(getattr(shape, "is_placeholder", False)):
        return True

    return False


def _remove_all_existing_diagram_artwork(slide: Any, slide_width_in: float, slide_height_in: float) -> None:
    """Remove the old/example diagram from a worksheet slide.

    This is intentionally stronger than the normal bounding-box replacement.
    It is used for the fixed template's configured diagram slide because the
    existing sample diagram can be a collection of pictures, auto-shapes and
    connectors that do not always sit inside the saved placement box.  Protected
    title/header/footer/logo/background elements are kept.
    """
    removable_types = {
        MSO_SHAPE_TYPE.PICTURE,
        MSO_SHAPE_TYPE.LINE,
        MSO_SHAPE_TYPE.AUTO_SHAPE,
        MSO_SHAPE_TYPE.GROUP,
        MSO_SHAPE_TYPE.FREEFORM,
    }

    to_remove: list[Any] = []
    for shape in list(slide.shapes):
        if shape.shape_type not in removable_types:
            continue
        if _is_template_protected_shape(shape, slide_height_in):
            continue

        left, top, right, bottom = _shape_bounds_in(shape)
        center_y = (top + bottom) / 2.0

        # Remove body-area artwork. This catches the old tank/sump/connector
        # diagram even when its bounding box differs from the configured image
        # placement box. Header/title/footer/logo areas stay protected above.
        if _TOP_TEMPLATE_SAFE_IN < center_y < max(_TOP_TEMPLATE_SAFE_IN, slide_height_in - _BOTTOM_TEMPLATE_SAFE_IN):
            to_remove.append(shape)

    for shape in to_remove:
        element = shape._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)



def _remove_old_diagram_content(
    slide: Any,
    box_x: float,
    box_y: float,
    box_width: float,
    box_height: float,
    slide_width_in: float,
    slide_height_in: float,
) -> None:
    """Remove only the old diagram artwork occupying the configured diagram area.

    The fixed template can contain an old diagram as a mixture of pictures,
    connectors and diagram-label AutoShapes instead of one single picture.
    All such diagram artwork is removed as a set, while template text, header,
    footer, logos, backgrounds, tables and other unrelated elements remain.
    """
    tol = _REPLACEMENT_TOLERANCE_IN
    replacement_box = (
        max(0.0, box_x - tol),
        max(0.0, box_y - tol),
        min(slide_width_in, box_x + box_width + tol),
        min(slide_height_in, box_y + box_height + tol),
    )

    removable_types = {
        MSO_SHAPE_TYPE.PICTURE,
        MSO_SHAPE_TYPE.LINE,
        MSO_SHAPE_TYPE.AUTO_SHAPE,
        MSO_SHAPE_TYPE.GROUP,
        MSO_SHAPE_TYPE.FREEFORM,
    }

    to_remove: list[Any] = []
    for shape in list(slide.shapes):
        if shape.shape_type not in removable_types:
            continue
        if _is_template_protected_shape(shape, slide_height_in):
            continue
        if not _boxes_overlap(_shape_bounds_in(shape), replacement_box):
            continue
        to_remove.append(shape)

    for shape in to_remove:
        element = shape._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)



def _move_picture_behind_template_shapes(slide: Any, picture: Any) -> None:
    """Keep preserved template overlays (logo/footer/etc.) above the new image."""
    sp_tree = slide.shapes._spTree
    picture_element = picture._element
    # The first two children are the group non-visual/group properties. Insert
    # immediately after them so all preserved template artwork stays on top.
    sp_tree.remove(picture_element)
    sp_tree.insert(2, picture_element)



def generate_ppt_from_fixed_template(
    template_path: Path,
    diagram_png: bytes,
    config_path: Path,
) -> bytes:
    """Replace the fixed template's old diagram with the existing generated PNG.

    The original template file is never modified. The configured target slide is
    loaded from the actual fixed .pptx, old diagram artwork occupying the target
    diagram area is removed, and the exact generated PNG is fitted into the same
    area. All unrelated template content and styling are preserved.
    """
    if not template_path.exists():
        raise FileNotFoundError(f"Fixed PowerPoint template not found: {template_path}")
    if not diagram_png:
        raise ValueError("Generated diagram image is empty.")

    presentation = Presentation(str(template_path))
    config = load_placement_config(config_path)

    slide_count = len(presentation.slides)
    if slide_count < 1:
        raise ValueError("Fixed PowerPoint template contains no slides.")

    slide_number = min(max(1, int(config["slide_number"])), slide_count)
    slide = presentation.slides[slide_number - 1]

    slide_width = float(presentation.slide_width) / EMU_PER_INCH
    slide_height = float(presentation.slide_height) / EMU_PER_INCH

    x = min(max(0.0, float(config["x"])), max(0.0, slide_width - 0.10))
    y = min(max(0.0, float(config["y"])), max(0.0, slide_height - 0.10))
    width = min(max(0.10, float(config["width"])), max(0.10, slide_width - x))
    height = min(max(0.10, float(config["height"])), max(0.10, slide_height - y))

    # Critical replacement step: remove the old template diagram before adding
    # the new one. This prevents old/new diagram overlap in the downloaded PPT.
    _remove_old_diagram_content(
        slide,
        x,
        y,
        width,
        height,
        slide_width,
        slide_height,
    )

    if bool(config.get("maintain_aspect_ratio", True)):
        with Image.open(BytesIO(diagram_png)) as image:
            source_width, source_height = image.size
        picture_x, picture_y, picture_width, picture_height = _fit_inside_box(
            source_width,
            source_height,
            x,
            y,
            width,
            height,
            str(config.get("horizontal_align", "center")),
            str(config.get("vertical_align", "middle")),
        )
    else:
        picture_x, picture_y, picture_width, picture_height = x, y, width, height

    picture = slide.shapes.add_picture(
        BytesIO(diagram_png),
        Inches(picture_x),
        Inches(picture_y),
        width=Inches(picture_width),
        height=Inches(picture_height),
    )

    # Preserve header/footer/logo/text overlays even if the configured image
    # region reaches their area.
    _move_picture_behind_template_shapes(slide, picture)

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


# =============================================================================
# SESSION-LEVEL PPT SUPPORT: FIXED SLIDES 1-6 + DYNAMIC WORKSHEET PAIRS + FIXED 29+
# =============================================================================

_GENERATED_DIAGRAM_SHAPE_PREFIX = "RTS_GENERATED_DIAGRAM::"
_FIXED_PREFIX_SLIDE_COUNT = 6
_FIXED_TAIL_START_SLIDE_NUMBER = 29


def _presentation_bytes(presentation: Presentation) -> bytes:
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def initialize_combined_ppt_session(
    template_path: Path,
    config_path: Path,
) -> bytes:
    """Load the complete fixed template once for the current PPT session.

    The source template file is never modified.  Slides 1-6 and slide 29 onward
    are treated as fixed protected template pages.  Worksheet diagram/estimate
    pages are rebuilt dynamically only inside the editable middle region between
    slide 6 and slide 29.
    """
    if not template_path.exists():
        raise FileNotFoundError(f"Fixed PowerPoint template not found: {template_path}")

    presentation = Presentation(str(template_path))
    if len(presentation.slides) < 1:
        raise ValueError("Fixed PowerPoint template contains no slides.")

    config = load_placement_config(config_path)
    slide_number = int(config["slide_number"])
    if slide_number < 1 or slide_number > len(presentation.slides):
        raise ValueError(
            f"Configured diagram slide {slide_number} is outside the fixed "
            f"template slide range 1-{len(presentation.slides)}."
        )

    return _presentation_bytes(presentation)


def _clear_slide_relationships_except_layout(slide: Any) -> None:
    """Clear destination-slide relationships that are not its slide layout."""
    for rel_id, rel in list(slide.part.rels.items()):
        if "slideLayout" in str(rel.reltype):
            continue
        try:
            slide.part.drop_rel(rel_id)
        except Exception:
            pass


def _clone_template_slide(presentation: Presentation, source_slide: Any) -> Any:
    """Append an exact visual clone of one fixed-template slide."""
    destination = presentation.slides.add_slide(source_slide.slide_layout)

    # Remove placeholders/shapes automatically created by the destination layout.
    for shape in list(destination.shapes):
        element = shape._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    _clear_slide_relationships_except_layout(destination)

    relationship_map: dict[str, str] = {}
    for old_rel_id, rel in source_slide.part.rels.items():
        reltype = str(rel.reltype)
        if "slideLayout" in reltype or "notesSlide" in reltype:
            continue
        try:
            if bool(rel.is_external):
                new_rel_id = destination.part.rels._add_relationship(
                    reltype,
                    rel.target_ref,
                    is_external=True,
                )
            else:
                new_rel_id = destination.part.rels._add_relationship(
                    reltype,
                    rel.target_part,
                    is_external=False,
                )
            relationship_map[str(old_rel_id)] = str(new_rel_id)
        except Exception:
            continue

    destination_tree = destination.shapes._spTree
    for shape in source_slide.shapes:
        copied = deepcopy(shape._element)
        for element in copied.iter():
            for attr_name, attr_value in list(element.attrib.items()):
                if attr_value in relationship_map:
                    element.set(attr_name, relationship_map[attr_value])
        destination_tree.insert_element_before(copied, "p:extLst")

    # Preserve a source-slide-specific background. Layout/master backgrounds
    # remain inherited automatically from the same fixed-template layout.
    src_cSld = source_slide._element.cSld
    dst_cSld = destination._element.cSld
    src_bg = getattr(src_cSld, "bg", None)
    dst_bg = getattr(dst_cSld, "bg", None)
    if dst_bg is not None:
        dst_cSld.remove(dst_bg)
    if src_bg is not None:
        dst_cSld.insert(0, deepcopy(src_bg))

    return destination


def _slide_index(presentation: Presentation, slide: Any) -> int:
    """Return the current zero-based index of a slide object."""
    for index, candidate in enumerate(presentation.slides):
        if candidate == slide:
            return index
    raise ValueError("Slide is not part of this presentation.")


def _move_slide_to_index(presentation: Presentation, slide: Any, new_index: int) -> None:
    """Move an existing slide to a specific zero-based position."""
    slide_count = len(presentation.slides)
    if slide_count <= 1:
        return

    current_index = _slide_index(presentation, slide)
    new_index = max(0, min(int(new_index), slide_count - 1))
    if current_index == new_index:
        return

    sld_id_list = presentation.slides._sldIdLst
    sld_id = sld_id_list[current_index]
    sld_id_list.remove(sld_id)

    if current_index < new_index:
        new_index -= 1
    new_index = max(0, min(new_index, len(sld_id_list)))
    sld_id_list.insert(new_index, sld_id)


def _delete_slide(presentation: Presentation, slide: Any) -> None:
    """Remove one slide from the presentation.

    This is used only for slides 7-28, the editable worksheet-generation region.
    Slides 1-6 and slide 29 onward are never passed to this function.
    """
    slide_id_list = presentation.slides._sldIdLst
    slide_index = _slide_index(presentation, slide)
    slide_id = slide_id_list[slide_index]
    rel_id = slide_id.rId
    slide_id_list.remove(slide_id)
    try:
        presentation.part.drop_rel(rel_id)
    except Exception:
        pass


def _set_generated_picture_name(picture: Any, diagram_id: str) -> None:
    name = f"{_GENERATED_DIAGRAM_SHAPE_PREFIX}{diagram_id}"
    try:
        picture._element.nvPicPr.cNvPr.set("name", name)
    except Exception:
        pass


def _insert_diagram_on_template_slide(
    presentation: Presentation,
    slide: Any,
    diagram_id: str,
    diagram_png: bytes,
    config: dict[str, Any],
    *,
    remove_existing_artwork: bool = True,
) -> None:
    """Insert one worksheet image into one generated image slide.

    The old/template sample diagram on the cloned image slide is removed first,
    then the newly generated worksheet PNG is embedded using add_picture().  This
    physically stores the worksheet image inside the final PPTX package, so it is
    still visible after closing and reopening PowerPoint.
    """
    if not diagram_png:
        raise ValueError(f"Worksheet diagram '{diagram_id}' has no PNG bytes.")

    slide_width = float(presentation.slide_width) / EMU_PER_INCH
    slide_height = float(presentation.slide_height) / EMU_PER_INCH

    x = min(max(0.0, float(config["x"])), max(0.0, slide_width - 0.10))
    y = min(max(0.0, float(config["y"])), max(0.0, slide_height - 0.10))
    width = min(max(0.10, float(config["width"])), max(0.10, slide_width - x))
    height = min(max(0.10, float(config["height"])), max(0.10, slide_height - y))

    if remove_existing_artwork:
        _remove_all_existing_diagram_artwork(slide, slide_width, slide_height)
        _remove_old_diagram_content(
            slide,
            x,
            y,
            width,
            height,
            slide_width,
            slide_height,
        )

    with Image.open(BytesIO(diagram_png)) as image:
        source_width, source_height = image.size

    if bool(config.get("maintain_aspect_ratio", True)):
        pic_x, pic_y, pic_w, pic_h = _fit_inside_box(
            source_width,
            source_height,
            x,
            y,
            width,
            height,
            str(config.get("horizontal_align", "center")),
            str(config.get("vertical_align", "middle")),
        )
    else:
        pic_x, pic_y, pic_w, pic_h = x, y, width, height

    picture = slide.shapes.add_picture(
        BytesIO(diagram_png),
        Inches(pic_x),
        Inches(pic_y),
        width=Inches(pic_w),
        height=Inches(pic_h),
    )
    _set_generated_picture_name(picture, diagram_id)

    # Keep the generated worksheet image visible above old diagram-area artwork.
    sp_tree = slide.shapes._spTree
    picture_element = picture._element
    sp_tree.remove(picture_element)
    sp_tree.insert_element_before(picture_element, "p:extLst")


def _replace_slide_with_full_page_image(
    presentation: Presentation,
    slide: Any,
    page_png: bytes,
    page_id: str,
) -> None:
    """Replace one generated pair's second slide with the supplied page image.

    This is applied only to cloned dynamic pages. Fixed template slides remain
    completely untouched.
    """
    if not page_png:
        return

    for shape in list(slide.shapes):
        element = shape._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    slide_width = float(presentation.slide_width) / EMU_PER_INCH
    slide_height = float(presentation.slide_height) / EMU_PER_INCH

    picture = slide.shapes.add_picture(
        BytesIO(page_png),
        Inches(0),
        Inches(0),
        width=Inches(slide_width),
        height=Inches(slide_height),
    )
    try:
        picture._element.nvPicPr.cNvPr.set(
            "name",
            f"RTS_PRICE_DETAILS::{page_id}",
        )
    except Exception:
        pass


def _slide_all_text(slide: Any) -> str:
    """Return normalized text from every text-bearing shape on one slide."""
    chunks: list[str] = []
    for shape in slide.shapes:
        if not bool(getattr(shape, "has_text_frame", False)):
            continue
        try:
            text = str(shape.text or "").strip()
        except Exception:
            text = ""
        if text:
            chunks.append(text)
    return "\n".join(chunks)


def _resolve_estimate_source_slide(
    presentation: Presentation,
    diagram_slide_number: int,
) -> Any:
    """Find the SYSTEM REQUIREMENTS AND ESTIMATE source slide in the editable region."""
    slide_count = len(presentation.slides)
    middle_start = _FIXED_PREFIX_SLIDE_COUNT
    middle_end = min(_FIXED_TAIL_START_SLIDE_NUMBER - 1, slide_count)

    adjacent_index = diagram_slide_number  # 1-based diagram -> next 0-based index
    if middle_start <= adjacent_index < middle_end:
        adjacent = presentation.slides[adjacent_index]
        if "SYSTEM REQUIREMENTS AND ESTIMATE" in _slide_all_text(adjacent).upper():
            return adjacent

    for index in range(middle_start, middle_end):
        slide = presentation.slides[index]
        if "SYSTEM REQUIREMENTS AND ESTIMATE" in _slide_all_text(slide).upper():
            return slide

    raise ValueError(
        "The editable PPT template region between slide 6 and slide 29 does "
        "not contain a SYSTEM REQUIREMENTS AND ESTIMATE slide."
    )


def _editable_middle_slides(presentation: Presentation) -> list[Any]:
    """Return original slides 7-28 only. These are the only slides we modify."""
    slide_count = len(presentation.slides)
    start = _FIXED_PREFIX_SLIDE_COUNT
    end = min(_FIXED_TAIL_START_SLIDE_NUMBER - 1, slide_count)
    if end <= start:
        return []
    return [presentation.slides[i] for i in range(start, end)]


def _open_presentation_from_template_bytes(session_base_ppt: bytes) -> Presentation:
    if not session_base_ppt:
        raise ValueError("PPT session base is empty.")
    presentation = Presentation(BytesIO(session_base_ppt))
    if len(presentation.slides) < _FIXED_PREFIX_SLIDE_COUNT:
        raise ValueError("Fixed PPT template must contain at least 6 slides.")
    return presentation


def generate_combined_ppt_from_session(
    session_base_ppt: bytes,
    diagrams: list[tuple[str, bytes]],
    config_path: Path,
    price_pages: dict[str, bytes] | None = None,
) -> bytes:
    """Build the final PPT with fixed slides preserved and dynamic worksheet pairs.

    Required output structure:

    - Slides 1-6: kept exactly from the fixed template.
    - Editable region after slide 6: rebuilt dynamically as
      Worksheet image slide -> worksheet-specific Price Details slide pairs.
    - Original slide 29 onward: kept exactly from the fixed template.

    The original fixed template file is never overwritten.  The old/sample image
    on each generated image page is removed, and each worksheet PNG is embedded
    physically in its corresponding image slide.  The loop processes every
    worksheet in order until the caller resets the PPT session via Start New PPT.
    """
    presentation = _open_presentation_from_template_bytes(session_base_ppt)
    config = load_placement_config(config_path)
    slide_count = len(presentation.slides)

    diagram_slide_number = int(config["slide_number"])
    if diagram_slide_number <= _FIXED_PREFIX_SLIDE_COUNT or diagram_slide_number >= _FIXED_TAIL_START_SLIDE_NUMBER:
        # The user's requirement allows modification only between slide 6 and 29.
        # If an old config points outside that region, use slide 7 as the image
        # source while preserving the user's placement box values.
        diagram_slide_number = _FIXED_PREFIX_SLIDE_COUNT + 1

    if diagram_slide_number > slide_count:
        raise ValueError(
            f"Configured diagram slide {diagram_slide_number} is outside the fixed "
            f"template slide range 1-{slide_count}."
        )

    source_diagram_slide = presentation.slides[diagram_slide_number - 1]
    source_estimate_slide = _resolve_estimate_source_slide(presentation, diagram_slide_number)

    valid_diagrams: list[tuple[str, bytes]] = []
    for diagram_id, diagram_png in diagrams:
        if not diagram_png:
            continue
        valid_diagrams.append((str(diagram_id), bytes(diagram_png)))

    if not valid_diagrams:
        return _presentation_bytes(presentation)

    # Clone every required pair first while the original editable region is still
    # available as a clean source.  This supports 1, 2, 5, 10 or any number of
    # worksheets without relying on a fixed number of template pages.
    price_pages = dict(price_pages or {})

    generated_pairs: list[tuple[Any, Any, str, bytes, bytes | None]] = []
    for diagram_id, diagram_png in valid_diagrams:
        image_slide = _clone_template_slide(
            presentation,
            source_diagram_slide,
        )
        price_slide = _clone_template_slide(
            presentation,
            source_estimate_slide,
        )
        generated_pairs.append(
            (
                image_slide,
                price_slide,
                diagram_id,
                diagram_png,
                price_pages.get(str(diagram_id)),
            )
        )

    # Remove only original slides 7-28.  Slides 1-6 and original slide 29 onward
    # are never removed, reordered, resized, or modified.
    for slide in _editable_middle_slides(presentation):
        _delete_slide(presentation, slide)

    # Move generated pairs into the editable middle region directly after slide 6.
    insert_index = _FIXED_PREFIX_SLIDE_COUNT
    for (
        image_slide,
        price_slide,
        diagram_id,
        diagram_png,
        price_page_png,
    ) in generated_pairs:
        _move_slide_to_index(presentation, image_slide, insert_index)
        insert_index += 1
        _move_slide_to_index(presentation, price_slide, insert_index)
        insert_index += 1

        _insert_diagram_on_template_slide(
            presentation,
            image_slide,
            diagram_id,
            diagram_png,
            config,
            remove_existing_artwork=True,
        )

        if price_page_png:
            _replace_slide_with_full_page_image(
                presentation,
                price_slide,
                price_page_png,
                diagram_id,
            )

    return _presentation_bytes(presentation)
