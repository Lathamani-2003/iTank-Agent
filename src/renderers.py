from __future__ import annotations

import heapq
import math
import textwrap
import threading
from contextlib import contextmanager
from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
from itertools import count
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .models import DiagramNode, DiagramPoint, DiagramSpec
from .routing_engine import analyze_routing_space, plan_connection_routes
from .component_image_paths import (
    COMPONENT_IMAGE_ALIASES,
    COMPONENT_IMAGE_PATHS,
    canonical_component_name,
    resolve_component_image_path,
)


RENDERER_VERSION = "2026-09-02-fixed-upright-arrow-v71.0"


@lru_cache(maxsize=64)
def load_font(size: int, bold: bool = False):
    """Load and cache a usable font for diagram rendering.

    Prefers common Windows fonts and falls back safely to Pillow's
    bundled/default fonts.
    """
    size = max(8, int(size))

    font_candidates = (
        [
            r"C:\\Windows\\Fonts\\arialbd.ttf",
            r"C:\\Windows\\Fonts\\calibrib.ttf",
            r"C:\\Windows\\Fonts\\segoeuib.ttf",
            "arialbd.ttf",
            "DejaVuSans-Bold.ttf",
        ]
        if bold
        else [
            r"C:\\Windows\\Fonts\\arial.ttf",
            r"C:\\Windows\\Fonts\\calibri.ttf",
            r"C:\\Windows\\Fonts\\segoeui.ttf",
            "arial.ttf",
            "DejaVuSans.ttf",
        ]
    )

    for font_path in font_candidates:
        try:
            return ImageFont.truetype(font_path, size=size)
        except (OSError, IOError):
            continue

    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()

def ensure_component_assets(assets_dir: Path | str, diagram=None) -> None:
    """Compatibility hook for the existing renderer.

    Component images are resolved lazily from ``src.component_image_paths``.
    This helper keeps the legacy rendering calls valid without generating or
    replacing any artwork.
    """
    try:
        Path(assets_dir).mkdir(parents=True, exist_ok=True)
    except (TypeError, OSError, ValueError):
        pass



# =============================================================================
# PAGE / STYLE
# =============================================================================

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5

CONTENT_LEFT = 0.64
CONTENT_RIGHT = 12.70
CONTENT_TOP = 1.18
CONTENT_BOTTOM = 6.82


BASE_SLIDE_WIDTH_IN = 13.333
BASE_SLIDE_HEIGHT_IN = 7.5
BASE_CONTENT_LEFT = 0.64
BASE_CONTENT_RIGHT = 12.70
BASE_CONTENT_TOP = 1.18
BASE_CONTENT_BOTTOM = 6.82

_CANVAS_LOCK = threading.RLock()


@dataclass(frozen=True)
class _CanvasSpec:
    width: float
    height: float
    content_left: float
    content_right: float
    content_top: float
    content_bottom: float


def _diagram_component_count(diagram: DiagramSpec) -> int:
    node_ids = [node.id for node in diagram.nodes if node.node_type != "junction"]
    if not node_ids:
        return 0
    adjacency = {node_id: set() for node_id in node_ids}
    for edge in diagram.edges:
        if edge.source in adjacency and edge.target in adjacency:
            adjacency[edge.source].add(edge.target)
            adjacency[edge.target].add(edge.source)
    remaining = set(node_ids)
    count_components = 0
    while remaining:
        count_components += 1
        seed = remaining.pop()
        stack = [seed]
        while stack:
            current = stack.pop()
            for neighbour in adjacency[current]:
                if neighbour in remaining:
                    remaining.remove(neighbour)
                    stack.append(neighbour)
    return count_components


def _alignment_load(diagram: DiagramSpec) -> tuple[int, int]:
    """Return max physical nodes sharing a horizontal / vertical visual band.

    Large fixed-reference fan-outs intentionally remain in one row/column instead
    of wrapping into a new structure. The canvas therefore grows along that axis
    so physical card and routing spacing stay readable as branch count increases.
    """
    nodes = [node for node in diagram.nodes if node.node_type != "junction"]
    if not nodes:
        return (0, 0)

    def max_bucket(values: list[float], tolerance: float = 0.035) -> int:
        ordered = sorted(values)
        best = 1
        start = 0
        for end in range(len(ordered)):
            while ordered[end] - ordered[start] > tolerance:
                start += 1
            best = max(best, end - start + 1)
        return best

    return (
        max_bucket([float(node.y) for node in nodes]),
        max_bucket([float(node.x) for node in nodes]),
    )


def _dynamic_canvas_spec(diagram: DiagramSpec) -> _CanvasSpec:
    """Return a scalable logical canvas for dense selected-component diagrams.

    Sketch/AI mode keeps the original fixed page. Component Selection mode grows
    the logical page from graph size and routing pressure, so 50-100+ nodes are
    given real geometric room instead of merely being drawn smaller.
    """
    selection_mode = "Component Selection mode" in (getattr(diagram, "style_notes", []) or [])
    if not selection_mode:
        return _CanvasSpec(
            BASE_SLIDE_WIDTH_IN, BASE_SLIDE_HEIGHT_IN,
            BASE_CONTENT_LEFT, BASE_CONTENT_RIGHT,
            BASE_CONTENT_TOP, BASE_CONTENT_BOTTOM,
        )

    node_count = max(1, sum(1 for node in diagram.nodes if node.node_type != "junction"))
    junction_count = sum(1 for node in diagram.nodes if node.node_type == "junction")
    edge_count = len(diagram.edges)
    component_count = max(1, _diagram_component_count(diagram))
    routing_pressure = edge_count / node_count

    # Physical area grows with real components, routing junctions and edges.
    # Distribution spines intentionally add junctions to simplify topology; the
    # canvas therefore reserves real geometric space for those trunks rather than
    # squeezing them into the old fixed page.
    routing_space = analyze_routing_space(diagram)
    required_area = (
        96.0
        + node_count * 3.00
        + junction_count * 0.56
        + edge_count * 0.82
        + component_count * 0.62
        + routing_space.area_bonus
    )
    aspect = 1.58
    if component_count >= 10:
        aspect = 1.48
    if routing_pressure >= 2.0:
        aspect = 1.42

    width = max(BASE_SLIDE_WIDTH_IN, math.sqrt(required_area * aspect))
    height = max(BASE_SLIDE_HEIGHT_IN, required_area / width)

    # Fixed-reference fan-outs do not wrap when branch count grows. Reserve real
    # physical width/height for long aligned rows/columns so the exact same
    # manifold-and-branch structure remains readable for 2, 6, 15 or more loads.
    max_row_load, max_col_load = _alignment_load(diagram)
    if max_row_load >= 6:
        width = max(width, 2.60 + max_row_load * 1.52)
    if max_col_load >= 6:
        height = max(height, 2.20 + max_col_load * 1.20)

    # Keep within the practical PowerPoint page-size ceiling while still allowing
    # very large PNG/PDF canvases.  The selected graph layout tiles internally, so
    # this cap does not force all nodes into one row.
    width = min(52.0, width)
    height = min(52.0, height)

    left = 0.72
    right = width - 0.72
    top = 1.22
    bottom = height - 0.72
    return _CanvasSpec(width, height, left, right, top, bottom)


@contextmanager
def _canvas_scope(diagram: DiagramSpec):
    """Temporarily apply diagram-specific logical page bounds.

    Rendering is serialized through a re-entrant lock so concurrent Streamlit
    sessions cannot observe another session's canvas constants.
    """
    global SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN
    global CONTENT_LEFT, CONTENT_RIGHT, CONTENT_TOP, CONTENT_BOTTOM

    spec = _dynamic_canvas_spec(diagram)
    with _CANVAS_LOCK:
        old = (
            SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN,
            CONTENT_LEFT, CONTENT_RIGHT, CONTENT_TOP, CONTENT_BOTTOM,
        )
        SLIDE_WIDTH_IN = spec.width
        SLIDE_HEIGHT_IN = spec.height
        CONTENT_LEFT = spec.content_left
        CONTENT_RIGHT = spec.content_right
        CONTENT_TOP = spec.content_top
        CONTENT_BOTTOM = spec.content_bottom
        try:
            yield spec
        finally:
            (
                SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN,
                CONTENT_LEFT, CONTENT_RIGHT, CONTENT_TOP, CONTENT_BOTTOM,
            ) = old

WHITE = (255, 255, 255)
BLACK = (36, 36, 36)

TITLE_BLUE = (8, 54, 112)
SET_BLUE = (10, 102, 227)

LABEL_BLUE = (7, 93, 215)
CARD_BG = (255, 255, 255)
CARD_BORDER = (205, 222, 244)

PIPE_COLOR = (0, 122, 255)
PIPE_DARK = (0, 82, 204)
PIPE_HALO = (215, 237, 255)

# Connection visual styles.  The routing geometry is unchanged; only the final
# stroke/arrow appearance is selected from edge topology metadata.
WATER_FLOW_COLOR = PIPE_COLOR
WATER_FLOW_DARK = PIPE_DARK
NON_WATER_COLOR = (26, 26, 26)
NON_WATER_HALO = (255, 255, 255)
ARROW_COLOR = (245, 96, 24)
ARROW_OUTLINE = (255, 255, 255)

# Wireless destination badge. This affects only the visual destination marker;
# routing geometry, component layout, line styles and arrow styles remain unchanged.
WIFI_SYMBOL_COLOR = (7, 93, 215)
WIFI_SYMBOL_BG = (255, 255, 255)
WIFI_SYMBOL_BORDER = (178, 207, 240)

RED = (6, 74, 170)
TAG_BG = (242, 248, 255)
TAG_BORDER = (178, 207, 240)
BORDER_GRAY = (210, 225, 242)


# =============================================================================
# RENDERING COMPATIBILITY HELPERS
# =============================================================================

def _normalize_rect(box):
    """Normalize Pillow rectangle coordinates and guarantee non-zero size."""
    x0, y0, x1, y1 = [int(round(value)) for value in box]

    if x1 < x0:
        x0, x1 = x1, x0
    if y1 < y0:
        y0, y1 = y1, y0

    if x1 == x0:
        x1 += 1
    if y1 == y0:
        y1 += 1

    return x0, y0, x1, y1


def _safe_rectangle(draw, box, **kwargs):
    """Draw a Pillow rectangle using validated coordinates."""
    draw.rectangle(_normalize_rect(box), **kwargs)


def _safe_rounded_rectangle(draw, box, radius=0, **kwargs):
    """Draw a Pillow rounded rectangle while keeping radius/bounds valid."""
    x0, y0, x1, y1 = _normalize_rect(box)
    max_radius = max(0, min((x1 - x0) // 2, (y1 - y0) // 2))

    draw.rounded_rectangle(
        (x0, y0, x1, y1),
        radius=min(max(0, int(radius)), max_radius),
        **kwargs,
    )


def _edge_is_water_flow(edge) -> bool:
    """Return True when edge metadata represents a physical water-flow line."""
    channel = str(getattr(edge, "topology_channel", "") or "").strip().lower()
    if channel:
        return channel in {
            "process",
            "water",
            "water_flow",
            "water-flow",
            "hydraulic",
            "fluid",
            "liquid",
            "pipeline",
        }

    label = str(getattr(edge, "label", "") or "").strip().lower()
    if not label:
        return False

    water_terms = (
        "water flow",
        "process flow",
        "water supply",
        "water line",
        "inlet",
        "outlet",
        "pipeline",
        "pipe flow",
        "pump flow",
    )
    return any(term in label for term in water_terms)


def _edge_line_style(edge) -> tuple[tuple[int, int, int], bool]:
    """Return line color/dotted style without changing route geometry."""
    if _edge_is_water_flow(edge):
        return WATER_FLOW_COLOR, False
    return NON_WATER_COLOR, True


def _wireless_destination_node_ids(diagram: DiagramSpec) -> set[str]:
    """Return visible component ids that terminate at least one wireless edge.

    The resolved ``connection_medium`` is already produced by the existing
    Automatic/Wired/Wireless feature. This function only converts that metadata
    into a renderer-side destination badge decision. Multiple wireless edges that
    terminate at the same component intentionally collapse to one id, preventing
    duplicate Wi-Fi symbols.
    """
    node_lookup = {node.id: node for node in diagram.nodes}
    destinations: set[str] = set()

    for edge in diagram.edges:
        medium = str(getattr(edge, "connection_medium", "") or "").strip().lower()
        if medium != "wireless":
            continue

        # Follow the actual flow direction. Unknown/bidirectional links retain the
        # stored target endpoint as the visual destination, which is deterministic
        # and does not modify any existing routing or topology behavior.
        destination_id = edge.source if edge.direction == "target_to_source" else edge.target
        node = node_lookup.get(destination_id)
        if node is None or node.node_type == "junction":
            continue
        destinations.add(destination_id)

    return destinations


@lru_cache(maxsize=4)
def _wifi_badge_png_bytes(size: int = 160) -> bytes:
    """Create a small transparent Wi-Fi badge for PowerPoint placement."""
    size = max(64, int(size))
    image = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)

    border_width = max(2, size // 40)
    draw.ellipse(
        (border_width, border_width, size - border_width, size - border_width),
        fill=(*WIFI_SYMBOL_BG, 248),
        outline=(*WIFI_SYMBOL_BORDER, 255),
        width=border_width,
    )

    stroke = max(4, size // 16)
    cx = size // 2
    base_y = int(size * 0.76)
    for radius in (int(size * 0.36), int(size * 0.26), int(size * 0.17)):
        bbox = (cx - radius, base_y - radius, cx + radius, base_y + radius)
        draw.arc(
            bbox,
            start=205,
            end=335,
            fill=(*WIFI_SYMBOL_COLOR, 255),
            width=stroke,
        )

    dot_radius = max(4, size // 20)
    draw.ellipse(
        (
            cx - dot_radius,
            base_y - dot_radius,
            cx + dot_radius,
            base_y + dot_radius,
        ),
        fill=(*WIFI_SYMBOL_COLOR, 255),
    )

    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


# Professional card sizes. Dense diagrams scale these down automatically.
NODE_SIZES = {
    "bore": (1.55, 1.24),
    "sump": (1.78, 1.36),
    "oht": (1.55, 1.24),
    "motor": (0.92, 0.78),
    "valve": (0.74, 0.62),
    "sensor": (0.74, 0.62),
    "controller": (1.00, 0.82),
    "junction": (0.10, 0.10),
    "other": (1.34, 1.08),
}

MIN_CARD_GAP_X = 0.82
MIN_CARD_GAP_Y = 0.78

PORT_MARGIN = 0.10
BASE_STUB = 0.34
STUB_STEP = 0.12

# A finer grid provides more independent routing lanes in dense manifolds.
GRID_STEP = 0.055
ROUTE_CLEARANCE = 0.26

# =============================================================================
# COMPONENT IMAGE RESOLUTION
# =============================================================================
# Manually copied component images live here:
#     <project>/assets/component_images/
#
# This mapping is intentionally kept in renderers.py as the final authoritative
# lookup used during diagram rendering.  The filenames below match the user's
# manually copied files, including spaces and capitalization.
_MANUAL_COMPONENT_IMAGE_DIR = Path(__file__).resolve().parent.parent / "assets" / "component_images"

_MANUAL_COMPONENT_IMAGE_FILES: dict[str, str] = {
    "Sump": "Sump.png",
    "Motor (Pump)": "motor.png",
    "Bore Well": "Bore well.png",
    "Well": "Well.png",
    "OHT Tank": "OHT Tank.png",
    "OHT Tank with Valve": "OHT Tank with Valve.png",
    "OHT Tank without Valve": "OHT Tank without Valve.png",
    "Non-Return Valve (NRV)": "NRV.png",
    "Pressure Relief Valve (PRV)": "PRV.png",
    "Motorized Valve (MV)": "Motorized Valve.png",
    "Display with GSM (DWG)": "DWG.png",
    "Auto Change Over Unit": "Auto Change over.png",
    "Ultrasonic Flow Meter": "Flow meter.png",
    "Flush Flow Meter": "Flow meter.png",
    "Electromagnetic Flow Meter": "Flow meter.png",
    "Transmitter": "Transmitter.png",
    "Master": "Master.png",
    "Smart Motor Controller (SMC)": "SMC.png",
    "Valve Controller (VCT)": "VCT.png",
    "Repeater": "Repeater.png",
    "Linear Level Sensor (LLS)": "LLS.png",
    "Valve Control Unit (VCU)": "VCU.png",
    "Display (D)": "Display.png",
    "Data Logger": "Data Logger.png",
}

# Generic assets are retained ONLY for unknown/custom AI nodes.  Recognized
# selectable components never fall back to these when a manual mapping exists.
ASSET_CANDIDATES = {
    "bore": [
        "bore.png", "bore_real.jpeg", "bore_real.jpg", "bore_real.png",
        "bore_photo.jpeg", "bore_photo.jpg", "bore_photo.png",
        "bore.jpeg", "bore.jpg",
    ],
    "sump": [
        "sump.png", "sump_real.jpeg", "sump_real.jpg", "sump_real.png",
        "sump_photo.jpeg", "sump_photo.jpg", "sump_photo.png",
        "sump.jpeg", "sump.jpg",
    ],
    "oht": [
        "oht.jpeg", "oht.png", "oht.jpg", "tank_real.jpeg", "tank_real.jpg",
        "tank_real.png", "oht_real.jpeg", "oht_real.jpg", "oht_real.png",
        "tank.jpeg", "tank.jpg", "tank.png",
    ],
    "motor": [
        "motor.png", "motor_real.jpeg", "motor_real.jpg", "motor_real.png",
        "motor_photo.jpeg", "motor_photo.jpg", "motor_photo.png",
        "motor.jpeg", "motor.jpg",
    ],
    "valve": [
        "mv.png", "prv.png", "nrv.png", "valve_real.jpeg", "valve_real.jpg",
        "valve_real.png", "valve.jpeg", "valve.jpg", "valve.png",
    ],
    "sensor": [
        "lls.jpeg", "flowmeter.png", "sensor_real.jpeg", "sensor_real.jpg",
        "sensor_real.png", "sensor.jpeg", "sensor.jpg", "sensor.png",
    ],
    "controller": [
        "master.png", "transmitter.png", "repeater.png", "smc.png", "vct.png",
        "auto_change_over.png", "dwg.png", "vcu.png", "display.png",
        "controller_real.jpeg", "controller_real.jpg", "controller_real.png",
        "controller.jpeg", "controller.jpg", "controller.png",
    ],
}

# Compatibility aliases used elsewhere in this renderer.
COMPONENT_IMAGE_FILES = COMPONENT_IMAGE_PATHS
COMPONENT_LABEL_ALIASES = COMPONENT_IMAGE_ALIASES


def _normalize_manual_filename(value: str) -> str:
    """Normalize a filename/component label for tolerant matching."""
    text = Path(str(value or "").strip()).stem
    text = text.replace("_", " ").replace("-", " ")
    text = text.replace("(", " ").replace(")", " ")
    text = " ".join(text.lower().split())
    return text


def _find_manual_file(filename: str) -> Path | None:
    """Find a manually copied image, tolerating case/spacing differences."""
    folder = _MANUAL_COMPONENT_IMAGE_DIR
    if not folder.is_dir():
        return None

    requested = Path(str(filename or "").strip())
    if not requested.name:
        return None

    # Exact filename first.
    exact = folder / requested.name
    if exact.is_file():
        return exact

    requested_norm = _normalize_manual_filename(requested.name)
    requested_stem_norm = _normalize_manual_filename(requested.stem)
    supported = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}

    try:
        files = list(folder.iterdir())
    except OSError:
        return None

    for path in files:
        if not path.is_file() or path.suffix.lower() not in supported:
            continue
        name_norm = _normalize_manual_filename(path.name)
        stem_norm = _normalize_manual_filename(path.stem)
        if (
            name_norm == requested_norm
            or stem_norm == requested_norm
            or name_norm == requested_stem_norm
            or stem_norm == requested_stem_norm
        ):
            return path
    return None


def _candidate_paths_for_stems(assets_dir: Path, stems: list[str]) -> list[Path]:
    candidates: list[Path] = []
    exts = ("png", "jpg", "jpeg", "webp", "bmp")
    for stem in stems:
        if not stem:
            continue
        for ext in exts:
            candidates.append(assets_dir / f"{stem}.{ext}")
            candidates.append(assets_dir / f"{stem}_latest.{ext}")
    return candidates


def _select_best_asset(candidates: list[Path]) -> Path | None:
    """Return the first existing image path from candidates."""
    for candidate in candidates:
        try:
            path = Path(candidate)
            if path.exists() and path.is_file():
                return path
        except (TypeError, OSError, ValueError):
            continue
    return None


def _select_first_stem_asset(assets_dir: Path, stems: list[str]) -> Path | None:
    exts = ("png", "jpg", "jpeg", "webp", "bmp")
    for stem in stems:
        if not stem:
            continue
        group: list[Path] = []
        for ext in exts:
            group.append(assets_dir / f"{stem}.{ext}")
            group.append(assets_dir / f"{stem}_latest.{ext}")
        chosen = _select_best_asset(group)
        if chosen is not None:
            return chosen
    return None


def _component_code(node: DiagramNode) -> str:
    """Return a compact code for vector fallback artwork when no image exists."""
    for detail in list(getattr(node, "details", []) or []):
        text = str(detail).strip()
        if text.lower().startswith("code:"):
            code = text.split(":", 1)[1].strip().upper()
            if code:
                return code[:7]

    label = str(getattr(node, "label", "") or "").strip()
    if "(" in label and ")" in label:
        inside = label.rsplit("(", 1)[-1].split(")", 1)[0].strip().upper()
        if inside and len(inside) <= 7:
            return inside

    words = [word for word in label.replace("-", " ").split() if word]
    if not words:
        return str(getattr(node, "node_type", "COMP")).upper()[:5]
    if len(words) == 1:
        return words[0][:5].upper()
    return "".join(word[0] for word in words[:5]).upper()


def _normalize_component_name(value: str) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _manual_component_name(node: DiagramNode) -> str | None:
    """Resolve a node/instance label to the authoritative manual image mapping."""
    label = str(getattr(node, "label", "") or "").strip()

    # Existing canonical resolver already understands numbered instances such as
    # "Sump 1" and "Bore Well 2".
    canonical = canonical_component_name(label)
    if canonical is not None:
        # Normalize equivalent flow-meter names to our manual file mapping.
        if canonical in _MANUAL_COMPONENT_IMAGE_FILES:
            return canonical

    code = _component_code(node).strip().upper()
    code_map = {
        "MASTER": "Master",
        "TX": "Transmitter",
        "RPT": "Repeater",
        "DWG": "Display with GSM (DWG)",
        "LLS": "Linear Level Sensor (LLS)",
        "MV": "Motorized Valve (MV)",
        "VCU": "Valve Control Unit (VCU)",
        "SMC": "Smart Motor Controller (SMC)",
        "VCT": "Valve Controller (VCT)",
        "D": "Display (D)",
        "UFM": "Ultrasonic Flow Meter",
        "FFM": "Flush Flow Meter",
        "PFM": "Flush Flow Meter",
        "EMFM": "Electromagnetic Flow Meter",
        "DL": "Data Logger",
        "PRV": "Pressure Relief Valve (PRV)",
        "NRV": "Non-Return Valve (NRV)",
        "OHT": "OHT Tank",
        "OHTV": "OHT Tank with Valve",
        "OHTNV": "OHT Tank without Valve",
        "SUMP": "Sump",
        "BORE": "Bore Well",
        "WELL": "Well",
        "MOTOR": "Motor (Pump)",
        "ACOU": "Auto Change Over Unit",
    }
    return code_map.get(code)


def _manual_asset_for_component(component_name: str) -> Path | None:
    """Resolve the exact file from assets/component_images/."""
    configured = _MANUAL_COMPONENT_IMAGE_FILES.get(component_name)
    if not configured:
        return None
    return _find_manual_file(configured)


def component_asset_path(assets_dir: Path, node_type: str) -> Path | None:
    """Generic fallback for unknown/custom AI nodes only."""
    node_type = (node_type or "other").lower().strip()
    candidates = [Path(assets_dir) / filename for filename in ASSET_CANDIDATES.get(node_type, [])]
    return _select_best_asset(candidates)


def node_component_asset_path(assets_dir: Path, node: DiagramNode) -> Path | None:
    """Return the exact manually copied image for a known component.

    Known selectable components are STRICT: when they map to a manual image,
    the renderer does not use the old/default asset set.  If the expected manual
    file is missing, ``None`` is returned and the existing vector fallback card
    is drawn instead of silently showing stale artwork.
    """
    component_name = _manual_component_name(node)
    if component_name is not None:
        return _manual_asset_for_component(component_name)

    # Preserve AssetFile compatibility only for genuinely unknown/custom nodes.
    for detail in list(getattr(node, "details", []) or []):
        text = str(detail).strip()
        if text.lower().startswith("assetfile:"):
            raw_value = text.split(":", 1)[1].strip()
            if raw_value:
                requested = Path(raw_value).expanduser()
                if not requested.is_absolute():
                    requested = Path(assets_dir) / requested.name
                if requested.exists() and requested.is_file():
                    return requested
                return None

    return component_asset_path(Path(assets_dir), getattr(node, "node_type", "other"))


def _contain_size(
    source_w: int,
    source_h: int,
    target_w: int,
    target_h: int,
) -> tuple[int, int]:
    if (
        source_w <= 0
        or source_h <= 0
        or target_w <= 0
        or target_h <= 0
    ):
        return (
            max(1, target_w),
            max(1, target_h),
        )

    scale = min(
        target_w / source_w,
        target_h / source_h,
    )

    return (
        max(1, int(source_w * scale)),
        max(1, int(source_h * scale)),
    )


def contain_image(
    source: Image.Image,
    target_width: int,
    target_height: int,
) -> Image.Image:
    # Preserve the user's transparent PNG assets correctly in the white diagram
    # cards. A direct RGBA -> RGB conversion composites transparent pixels against
    # black, which changes the provided artwork. Composite on white instead.
    if source.mode in {"RGBA", "LA"} or "transparency" in source.info:
        rgba = source.convert("RGBA")
        white_bg = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
        white_bg.alpha_composite(rgba)
        source = white_bg.convert("RGB")
    else:
        source = source.convert("RGB")

    target_width = max(
        1,
        int(target_width),
    )

    target_height = max(
        1,
        int(target_height),
    )

    output = Image.new(
        "RGB",
        (target_width, target_height),
        WHITE,
    )

    new_w, new_h = _contain_size(
        source.width,
        source.height,
        target_width,
        target_height,
    )

    resized = source.resize(
        (new_w, new_h),
        Image.Resampling.LANCZOS,
    )

    output.paste(
        resized,
        (
            (target_width - new_w) // 2,
            (target_height - new_h) // 2,
        ),
    )

    return output


@lru_cache(maxsize=256)
def _cached_contained_asset(
    path_text: str,
    mtime_ns: int,
    file_size: int,
    target_width: int,
    target_height: int,
) -> Image.Image:
    """Return the exact same contained component bitmap, cached by file version/size.

    The cache key includes mtime and file size, so replacing an asset invalidates the
    cached raster automatically.  The returned pixels are identical to contain_image().
    """
    del mtime_ns, file_size  # values intentionally participate in the cache key
    with Image.open(path_text) as source:
        return contain_image(source, target_width, target_height)


def _contained_asset_for_path(path: Path, target_width: int, target_height: int) -> Image.Image:
    try:
        stat = path.stat()
        return _cached_contained_asset(
            str(path.resolve()),
            int(stat.st_mtime_ns),
            int(stat.st_size),
            int(target_width),
            int(target_height),
        )
    except OSError:
        with Image.open(path) as source:
            return contain_image(source, target_width, target_height)


def add_picture_contained_ppt(
    slide,
    image_path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
) -> None:
    with Image.open(image_path) as image:
        source_w, source_h = image.size

    target_w = max(
        1,
        int(w * 1000),
    )

    target_h = max(
        1,
        int(h * 1000),
    )

    fit_w_px, fit_h_px = _contain_size(
        source_w,
        source_h,
        target_w,
        target_h,
    )

    fit_w = w * fit_w_px / target_w
    fit_h = h * fit_h_px / target_h

    picture_x = x + (w - fit_w) / 2
    picture_y = y + (h - fit_h) / 2

    slide.shapes.add_picture(
        str(image_path),
        Inches(picture_x),
        Inches(picture_y),
        width=Inches(fit_w),
        height=Inches(fit_h),
    )


# =============================================================================
# TEXT
# =============================================================================

def get_display_title(
    diagram: DiagramSpec,
) -> str:
    title = (
        (diagram.title or "")
        .strip()
    )

    if not title:
        title = (
            "SITE WATER "
            "AUTOMATION DIAGRAM"
        )

    if not title.upper().startswith(
        "FULLY AUTOMATION"
    ):
        title = (
            "FULLY AUTOMATION - "
            + title
        )

    return title.upper()


def split_title_lines(
    text: str,
    width: int = 62,
) -> list[str]:
    lines = textwrap.wrap(
        text,
        width=width,
        break_long_words=False,
        break_on_hyphens=False,
    )

    if len(lines) <= 2:
        return lines or [text]

    return [
        lines[0],
        " ".join(lines[1:]),
    ]


def _numeric_suffix(*values: str) -> str:
    for value in values:
        digits = "".join(ch for ch in str(value or "") if ch.isdigit())
        if digits:
            return digits
    return ""


def _expand_component_label(
    text: str,
    node_type: str = "other",
    node_id: str = "",
) -> str:
    clean = " ".join(str(text or "").replace("_", " ").split())
    if not clean:
        clean = str(node_id or "").replace("_", " ").strip()

    upper = clean.upper().replace(".", "")
    digits = _numeric_suffix(clean, node_id)

    if node_type == "motor":
        if upper in {"M", "MOTOR", "PUMP"}:
            return f"Motor {digits}".strip()
        if upper.startswith("M") and upper[1:].isdigit():
            return f"Motor {upper[1:]}"
        if upper.startswith("MOTOR"):
            suffix = _numeric_suffix(upper[5:], digits)
            return f"Motor {suffix}".strip() if suffix else "Motor"

    if "OHT" in upper:
        words: list[str] = []
        for word in clean.replace(".", "").split():
            if word.upper() == "OHT":
                # Avoid results such as "Tank Overhead Tank" or duplicated
                # "Overhead Tank Overhead Tank".
                if (
                    len(words) >= 2
                    and words[-2].lower() == "overhead"
                    and words[-1].lower() == "tank"
                ):
                    continue
                if words and words[-1].lower() == "tank":
                    words.pop()
                words.extend(["Overhead", "Tank"])
            else:
                words.append(word.title() if word.isupper() else word)
        return " ".join(words).strip()

    if upper == "SUMP":
        return "Sump"
    if upper == "BORE":
        return "Bore"

    return clean.title() if clean.isupper() else clean


def component_display_label(node: DiagramNode) -> str:
    return _expand_component_label(
        getattr(node, "label", ""),
        getattr(node, "node_type", "other"),
        getattr(node, "id", ""),
    )


def wrap_component_label(
    text: str,
    width: int = 19,
) -> list[str]:
    clean = " ".join(
        (text or "").split()
    ).upper()

    if not clean:
        clean = "UNNAMED"

    lines = textwrap.wrap(
        clean,
        width=width,
        break_long_words=False,
        break_on_hyphens=False,
    )

    if len(lines) <= 2:
        return lines or [clean]

    return [
        lines[0],
        " ".join(lines[1:]),
    ]


def _semantic_size_text(
    text: str,
) -> str:
    return (
        (text or "")
        .lower()
        .replace("inches", "")
        .replace("inch", "")
        .replace("pipe", "")
        .replace('"', "")
        .replace("'", "")
        .replace(" ", "")
    )


def format_edge_label(edge) -> str:
    pipe_size = (
        edge.pipe_size or ""
    ).strip()

    label = (
        edge.label or ""
    ).strip()

    if pipe_size and label:
        size_a = _semantic_size_text(
            pipe_size
        )

        size_b = _semantic_size_text(
            label
        )

        if (
            size_a
            and size_b
            and (
                size_a in size_b
                or size_b in size_a
            )
        ):
            return label

        return f"{pipe_size} | {label}"

    return pipe_size or label


# =============================================================================
# LAYOUT
# =============================================================================

@dataclass
class _LayoutNode:
    node: DiagramNode
    x: float
    y: float
    w: float
    h: float
    desired_x: float
    desired_y: float


class _LayoutBoxes(dict):
    """
    Dict-compatible layout result plus a smooth sketch-to-layout geometry warp.

    The previous renderer moved component cards to remove overlaps but continued
    mapping pipe waypoints with the original unmodified coordinate transform. That
    is the main reason a correctly extracted OpenCV route could visually connect to
    the wrong place after layout. This class keeps the public dict behavior while
    mapping route waypoints through the same displacement field as the components.
    """

    def __init__(self, values, warp_points):
        super().__init__(values)
        self._warp_points = list(warp_points)

    def map_normalized_point(self, point: DiagramPoint) -> tuple[float, float]:
        """Map a hand-sketch waypoint through the same displacement as nearby cards.

        This is the critical geometry link between the extracted route and the final
        layout. Cards may move slightly to remove collisions, so their nearby pipe
        bends must move with them. The topology and waypoint order are unchanged.
        """
        content_w = CONTENT_RIGHT - CONTENT_LEFT
        content_h = CONTENT_BOTTOM - CONTENT_TOP
        raw_x = CONTENT_LEFT + min(max(float(point.x), 0.0), 1.0) * content_w
        raw_y = CONTENT_TOP + min(max(float(point.y), 0.0), 1.0) * content_h

        if not self._warp_points:
            return raw_x, raw_y

        weighted_dx = 0.0
        weighted_dy = 0.0
        total_weight = 0.0

        for anchor_x, anchor_y, delta_x, delta_y in self._warp_points:
            distance_sq = (raw_x - anchor_x) ** 2 + (raw_y - anchor_y) ** 2
            weight = 1.0 / (distance_sq + 0.28 ** 2)
            weighted_dx += delta_x * weight
            weighted_dy += delta_y * weight
            total_weight += weight

        # Fixed page-corner anchors prevent a long outer pipeline from being
        # dragged too far by a dense cluster of moved cards.
        for anchor_x, anchor_y in (
            (CONTENT_LEFT, CONTENT_TOP),
            (CONTENT_RIGHT, CONTENT_TOP),
            (CONTENT_LEFT, CONTENT_BOTTOM),
            (CONTENT_RIGHT, CONTENT_BOTTOM),
        ):
            distance_sq = (raw_x - anchor_x) ** 2 + (raw_y - anchor_y) ** 2
            total_weight += 0.70 / (distance_sq + 0.52 ** 2)

        if total_weight > 0.0:
            raw_x += weighted_dx / total_weight
            raw_y += weighted_dy / total_weight

        return (
            min(max(raw_x, CONTENT_LEFT + 0.04), CONTENT_RIGHT - 0.04),
            min(max(raw_y, CONTENT_TOP + 0.04), CONTENT_BOTTOM - 0.04),
        )


def _node_size(
    node: DiagramNode,
    scale: float,
) -> tuple[float, float]:
    base_w, base_h = NODE_SIZES.get(
        node.node_type,
        NODE_SIZES["other"],
    )

    if node.node_type == "junction":
        return base_w, base_h

    return (
        base_w * scale,
        base_h * scale,
    )


def _layout_scale(
    diagram: DiagramSpec,
) -> float:
    count_visible = sum(
        1
        for node in diagram.nodes
        if node.node_type != "junction"
    )

    if count_visible <= 8:
        return 1.22

    if count_visible <= 12:
        return 1.10

    if count_visible <= 16:
        return 1.00

    if count_visible <= 20:
        return 0.90

    return 0.82

def _clamp_center(
    item: _LayoutNode,
) -> None:
    item.x = min(
        max(
            item.x,
            CONTENT_LEFT + item.w / 2,
        ),
        CONTENT_RIGHT - item.w / 2,
    )

    item.y = min(
        max(
            item.y,
            CONTENT_TOP + item.h / 2,
        ),
        CONTENT_BOTTOM - item.h / 2,
    )


def _overlap(
    a: _LayoutNode,
    b: _LayoutNode,
) -> tuple[float, float]:
    overlap_x = (
        (a.w + b.w) / 2
        + MIN_CARD_GAP_X
        - abs(a.x - b.x)
    )

    overlap_y = (
        (a.h + b.h) / 2
        + MIN_CARD_GAP_Y
        - abs(a.y - b.y)
    )

    return overlap_x, overlap_y



def _rects_overlap_with_gap(a, b, gap_x=0.10, gap_y=0.10):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return not (
        ax + aw + gap_x <= bx
        or bx + bw + gap_x <= ax
        or ay + ah + gap_y <= by
        or by + bh + gap_y <= ay
    )


def _resolve_layout_overlaps(
    items: list[_LayoutNode],
    gap_x: float = 0.12,
    gap_y: float = 0.10,
    iterations: int = 180,
) -> None:
    """Spread cards apart while staying close to their sketch positions.

    The old renderer kept every card at its extracted coordinate even when the
    extracted boxes overlapped. That makes it impossible for any pipe router to
    produce a readable engineering diagram. This solver moves ONLY the cards;
    edges are routed afterwards from the final card positions.
    """
    if len(items) < 2:
        return

    for _ in range(iterations):
        moved = False

        # Weak pull toward the original sketch position preserves the visual
        # left/right/top/bottom ordering without allowing overlaps.
        for item in items:
            item.x += (item.desired_x - item.x) * 0.018
            item.y += (item.desired_y - item.y) * 0.018
            _clamp_center(item)

        for i in range(len(items)):
            a = items[i]
            for j in range(i + 1, len(items)):
                b = items[j]

                ox = (
                    (a.w + b.w) / 2.0
                    + gap_x
                    - abs(a.x - b.x)
                )
                oy = (
                    (a.h + b.h) / 2.0
                    + gap_y
                    - abs(a.y - b.y)
                )

                if ox <= 0 or oy <= 0:
                    continue

                moved = True

                # Push along the axis requiring the least displacement. This
                # preserves the original 2-D arrangement better than a one-axis
                # row/column packing algorithm.
                if ox <= oy:
                    direction = 1.0 if a.x <= b.x else -1.0
                    shift = ox * 0.52
                    a.x -= direction * shift
                    b.x += direction * shift
                else:
                    direction = 1.0 if a.y <= b.y else -1.0
                    shift = oy * 0.52
                    a.y -= direction * shift
                    b.y += direction * shift

                _clamp_center(a)
                _clamp_center(b)

        if not moved:
            break


def _layout_has_overlap(items: list[_LayoutNode], gap_x=0.055, gap_y=0.055) -> bool:
    for i, a in enumerate(items):
        for b in items[i + 1:]:
            if _rects_overlap_with_gap(
                (a.x - a.w / 2, a.y - a.h / 2, a.w, a.h),
                (b.x - b.w / 2, b.y - b.h / 2, b.w, b.h),
                gap_x,
                gap_y,
            ):
                return True
    return False


def compute_layout_boxes(diagram: DiagramSpec):
    """Create a dynamically spread, collision-free layout.

    The extracted x/y ordering is preserved, but a sketch that occupies only a
    small central area is expanded to use the available drawing canvas.  This is
    important for routing: a collision solver alone can remove card overlap while
    still leaving too little whitespace for independent pipelines.
    """
    visible_nodes = [node for node in diagram.nodes if node.node_type != "junction"]
    visible_count = len(visible_nodes)
    selection_mode = "Component Selection mode" in (getattr(diagram, "style_notes", []) or [])

    if selection_mode:
        # The logical canvas itself grows for dense diagrams, so cards no longer
        # need to be aggressively shrunk just because the node count is large.
        if visible_count <= 60:
            initial_scale = 1.0
        elif visible_count <= 100:
            initial_scale = 0.92
        elif visible_count <= 160:
            initial_scale = 0.84
        else:
            initial_scale = 0.76
    elif visible_count <= 8:
        initial_scale = 0.98
    elif visible_count <= 12:
        initial_scale = 0.88
    elif visible_count <= 16:
        initial_scale = 0.78
    elif visible_count <= 20:
        initial_scale = 0.70
    else:
        initial_scale = 0.62

    content_w = CONTENT_RIGHT - CONTENT_LEFT
    content_h = CONTENT_BOTTOM - CONTENT_TOP

    degree = {node.id: 0 for node in diagram.nodes}
    for edge in diagram.edges:
        degree[edge.source] = degree.get(edge.source, 0) + 1
        degree[edge.target] = degree.get(edge.target, 0) + 1
    routing_space = analyze_routing_space(diagram)

    if visible_nodes:
        raw_xs = [CONTENT_LEFT + min(max(float(node.x), 0.01), 0.99) * content_w for node in visible_nodes]
        raw_ys = [CONTENT_TOP + min(max(float(node.y), 0.01), 0.99) * content_h for node in visible_nodes]
        min_x, max_x = min(raw_xs), max(raw_xs)
        min_y, max_y = min(raw_ys), max(raw_ys)
        center_x = (min_x + max_x) / 2.0
        center_y = (min_y + max_y) / 2.0
        span_x = max(0.01, max_x - min_x)
        span_y = max(0.01, max_y - min_y)

        # Expand only when the extracted drawing is crowded into a small region.
        # Axis order and relative spacing remain unchanged.
        # Reserve real whitespace for route channels before cards are finalized.
        # More connections increase the target span instead of squeezing pipes into
        # the same narrow corridors after layout.
        route_x_reserve = min(content_w * 0.10, routing_space.component_gap_x * 0.85)
        route_y_reserve = min(content_h * 0.12, routing_space.component_gap_y * 0.85)
        target_span_x = min(content_w * 0.90, max(span_x + route_x_reserve, content_w * 0.70))
        target_span_y = min(content_h * 0.86, max(span_y + route_y_reserve, content_h * 0.66))
        expand_x = min(2.25, max(1.0, target_span_x / span_x))
        expand_y = min(2.10, max(1.0, target_span_y / span_y))
    else:
        center_x = (CONTENT_LEFT + CONTENT_RIGHT) / 2.0
        center_y = (CONTENT_TOP + CONTENT_BOTTOM) / 2.0
        expand_x = expand_y = 1.0

    items: list[_LayoutNode] = []

    for attempt in range(7):
        scale = max(0.52, initial_scale - attempt * 0.045)
        items = []

        for node in visible_nodes:
            if selection_mode:
                # Keep the same professional card proportions on expanded canvases.
                bulk_card_factor = 1.00 if visible_count <= 100 else 0.92
                base_w, base_h = (1.42 * bulk_card_factor, 1.08 * bulk_card_factor)
                width = base_w * scale
                height = base_h * scale
            else:
                base_w, base_h = NODE_SIZES.get(node.node_type, NODE_SIZES["other"])
                width = min(base_w * scale, 1.66 if node.node_type in {"sump", "bore", "oht"} else 1.38)
                height = min(base_h * scale, 1.34 if node.node_type in {"sump", "bore", "oht"} else 1.14)

            raw_x = CONTENT_LEFT + min(max(float(node.x), 0.015), 0.985) * content_w
            raw_y = CONTENT_TOP + min(max(float(node.y), 0.015), 0.985) * content_h
            desired_x = center_x + (raw_x - center_x) * expand_x
            desired_y = center_y + (raw_y - center_y) * expand_y

            item = _LayoutNode(
                node=node,
                x=desired_x,
                y=desired_y,
                w=width,
                h=height,
                desired_x=desired_x,
                desired_y=desired_y,
            )
            _clamp_center(item)
            item.desired_x, item.desired_y = item.x, item.y
            items.append(item)

        # High-degree components need more corridor space because many independent
        # pipes must leave the card without immediately sharing a lane.
        for _ in range(360):
            moved = False
            for item in items:
                item.x += (item.desired_x - item.x) * 0.010
                item.y += (item.desired_y - item.y) * 0.010
                _clamp_center(item)

            for i, a in enumerate(items):
                for b in items[i + 1:]:
                    degree_a = degree.get(a.node.id, 0)
                    degree_b = degree.get(b.node.id, 0)
                    # High-degree cards need extra breathing room for the dynamic
                    # two-side fan-out rule.  This only affects diagrams where a
                    # selected component actually has several required links.
                    high_degree = max(degree_a, degree_b)
                    fanout_bonus = min(0.14, max(0, high_degree - 2) * 0.05)
                    connection_pressure = min(
                        0.34,
                        0.022 * (degree_a + degree_b) + fanout_bonus,
                    )
                    gap_x = max(
                        routing_space.component_gap_x,
                        max(0.16, 0.30 - visible_count * 0.006) + connection_pressure,
                    )
                    gap_y = max(
                        routing_space.component_gap_y,
                        max(0.14, 0.26 - visible_count * 0.005) + connection_pressure * 0.82,
                    )
                    ox = (a.w + b.w) / 2.0 + gap_x - abs(a.x - b.x)
                    oy = (a.h + b.h) / 2.0 + gap_y - abs(a.y - b.y)
                    if ox <= 0 or oy <= 0:
                        continue
                    moved = True

                    original_dx = b.desired_x - a.desired_x
                    original_dy = b.desired_y - a.desired_y
                    if abs(original_dx) >= abs(original_dy):
                        sign = 1.0 if original_dx >= 0 else -1.0
                        if abs(original_dx) < 1e-9:
                            sign = 1.0 if b.x >= a.x else -1.0
                        shift = ox * 0.54 + 0.02
                        a.x -= sign * shift
                        b.x += sign * shift
                    else:
                        sign = 1.0 if original_dy >= 0 else -1.0
                        if abs(original_dy) < 1e-9:
                            sign = 1.0 if b.y >= a.y else -1.0
                        shift = oy * 0.54 + 0.02
                        a.y -= sign * shift
                        b.y += sign * shift
                    _clamp_center(a)
                    _clamp_center(b)

            if not moved:
                break

        if not _layout_has_overlap(items, gap_x=0.10, gap_y=0.09):
            break

    values = {
        item.node.id: (item.x - item.w / 2, item.y - item.h / 2, item.w, item.h)
        for item in items
    }
    warp_points = [
        (
            CONTENT_LEFT + min(max(float(item.node.x), 0.015), 0.985) * content_w,
            CONTENT_TOP + min(max(float(item.node.y), 0.015), 0.985) * content_h,
            item.x - (CONTENT_LEFT + min(max(float(item.node.x), 0.015), 0.985) * content_w),
            item.y - (CONTENT_TOP + min(max(float(item.node.y), 0.015), 0.985) * content_h),
        )
        for item in items
    ]

    layout = _LayoutBoxes(values, warp_points)
    for node in diagram.nodes:
        if node.node_type != "junction":
            continue
        size = NODE_SIZES["junction"][0]
        mapped_x, mapped_y = layout.map_normalized_point(DiagramPoint(x=float(node.x), y=float(node.y)))
        values[node.id] = (mapped_x - size / 2, mapped_y - size / 2, size, size)

    return _LayoutBoxes(values, warp_points)

def _center_final_layout_horizontally(boxes, routes, labels):
    """Translate the completed diagram horizontally so it is centered.

    IMPORTANT: this runs only *after* ``build_edge_routes`` and edge-label
    placement have finished.  It therefore does not participate in routing and
    cannot change which ports, bends, lanes, source/target relationships or
    route shapes were chosen.  Every rendered geometry item receives the same
    horizontal offset.
    """
    min_x = float("inf")
    max_x = float("-inf")

    for box in boxes.values():
        x, _y, w, _h = box
        min_x = min(min_x, float(x))
        max_x = max(max_x, float(x) + float(w))

    for route in routes:
        if route is None:
            continue
        points, _direction = route
        for point in points:
            min_x = min(min_x, float(point[0]))
            max_x = max(max_x, float(point[0]))

    for placement in labels:
        if placement is None:
            continue
        _text, rect, leader_anchor = placement
        rx, _ry, rw, _rh = rect
        min_x = min(min_x, float(rx))
        max_x = max(max_x, float(rx) + float(rw))
        if leader_anchor is not None:
            min_x = min(min_x, float(leader_anchor[0]))
            max_x = max(max_x, float(leader_anchor[0]))

    if not math.isfinite(min_x) or not math.isfinite(max_x) or max_x <= min_x:
        return boxes, routes, labels

    target_center_x = (CONTENT_LEFT + CONTENT_RIGHT) / 2.0
    current_center_x = (min_x + max_x) / 2.0
    dx = target_center_x - current_center_x

    # Never move any already-routed geometry outside the existing drawing area.
    left_limit = CONTENT_LEFT + 0.02
    right_limit = CONTENT_RIGHT - 0.02
    min_allowed_dx = left_limit - min_x
    max_allowed_dx = right_limit - max_x
    dx = max(min_allowed_dx, min(dx, max_allowed_dx))

    if abs(dx) < 1e-9:
        return boxes, routes, labels

    centered_boxes = {
        node_id: (float(x) + dx, y, w, h)
        for node_id, (x, y, w, h) in boxes.items()
    }

    centered_routes = []
    for route in routes:
        if route is None:
            centered_routes.append(None)
            continue
        points, direction = route
        centered_routes.append(
            ([(float(x) + dx, y) for x, y in points], direction)
        )

    centered_labels = []
    for placement in labels:
        if placement is None:
            centered_labels.append(None)
            continue
        text_value, rect, leader_anchor = placement
        x, y, w, h = rect
        moved_anchor = None
        if leader_anchor is not None:
            moved_anchor = (float(leader_anchor[0]) + dx, leader_anchor[1])
        centered_labels.append(
            (text_value, (float(x) + dx, y, w, h), moved_anchor)
        )

    return centered_boxes, centered_routes, centered_labels


def _box_center(box):
    x, y, w, h = box

    return (
        x + w / 2,
        y + h / 2,
    )


def _choose_sides(
    source_box,
    target_box,
) -> tuple[str, str]:
    sx, sy = _box_center(source_box)
    tx, ty = _box_center(target_box)

    dx = tx - sx
    dy = ty - sy

    if abs(dx) >= abs(dy):
        if dx >= 0:
            return "right", "left"

        return "left", "right"

    if dy >= 0:
        return "bottom", "top"

    return "top", "bottom"


def _port_point(
    box,
    side: str,
    index: int,
    total: int,
    fraction_override: float | None = None,
) -> tuple[float, float]:
    x, y, w, h = box

    if fraction_override is not None:
        fraction = min(max(float(fraction_override), PORT_MARGIN), 1.0 - PORT_MARGIN)
    elif total <= 1:
        fraction = 0.5
    else:
        fraction = (
            PORT_MARGIN
            + (1.0 - 2.0 * PORT_MARGIN) * ((index + 1) / (total + 1))
        )

    if side == "left":
        return x, y + h * fraction
    if side == "right":
        return x + w, y + h * fraction
    if side == "top":
        return x + w * fraction, y
    return x + w * fraction, y + h


def _anchor_fraction(
    node: DiagramNode,
    side: str,
    anchor: DiagramPoint | None,
) -> float | None:
    """Map an image-derived endpoint anchor to a fraction along the card side."""
    if anchor is None:
        return None

    left = float(node.x) - float(node.width) / 2.0
    top = float(node.y) - float(node.height) / 2.0

    if side in {"left", "right"}:
        size = max(0.01, float(node.height))
        return min(max((float(anchor.y) - top) / size, PORT_MARGIN), 1.0 - PORT_MARGIN)

    size = max(0.01, float(node.width))
    return min(max((float(anchor.x) - left) / size, PORT_MARGIN), 1.0 - PORT_MARGIN)


def _stub_distance(
    index: int,
    total: int,
) -> float:
    if total <= 1:
        return BASE_STUB

    centered = (
        index
        - (total - 1) / 2
    )

    return (
        BASE_STUB
        + abs(centered) * STUB_STEP
    )


def _stub_point(
    point,
    side: str,
    distance: float,
) -> tuple[float, float]:
    x, y = point

    if side == "left":
        return x - distance, y

    if side == "right":
        return x + distance, y

    if side == "top":
        return x, y - distance

    return x, y + distance


def _port_sort_key(
    edge,
    role: str,
    side: str,
    boxes,
) -> float:
    anchor = edge.source_anchor if role == "source" else edge.target_anchor
    if anchor is not None:
        return float(anchor.y) if side in {"left", "right"} else float(anchor.x)

    other_id = edge.target if role == "source" else edge.source
    other_box = boxes.get(other_id)
    if other_box is None:
        return 0.0

    other_x, other_y = _box_center(other_box)
    return other_y if side in {"left", "right"} else other_x


# =============================================================================
# ROUTING GRID
# =============================================================================

def _snap(value: float) -> int:
    return int(
        round(
            value
            / GRID_STEP
        )
    )


def _unsnap(value: int) -> float:
    return value * GRID_STEP


def _point_inside_expanded_box(
    point,
    box,
    clearance: float = ROUTE_CLEARANCE,
) -> bool:
    px, py = point
    x, y, w, h = box

    return (
        x - clearance
        <= px
        <= x + w + clearance
        and
        y - clearance
        <= py
        <= y + h + clearance
    )


def _compress_polyline(
    points: list[
        tuple[float, float]
    ],
) -> list[
    tuple[float, float]
]:
    if len(points) <= 2:
        return points

    deduped = []

    for point in points:
        if (
            not deduped
            or (
                abs(
                    deduped[-1][0]
                    - point[0]
                )
                > 1e-8
                or
                abs(
                    deduped[-1][1]
                    - point[1]
                )
                > 1e-8
            )
        ):
            deduped.append(point)

    if len(deduped) <= 2:
        return deduped

    result = [
        deduped[0]
    ]

    for index in range(
        1,
        len(deduped) - 1,
    ):
        previous = result[-1]
        current = deduped[index]
        following = deduped[
            index + 1
        ]

        same_x = (
            abs(previous[0] - current[0])
            < 1e-8
            and
            abs(current[0] - following[0])
            < 1e-8
        )

        same_y = (
            abs(previous[1] - current[1])
            < 1e-8
            and
            abs(current[1] - following[1])
            < 1e-8
        )

        if same_x or same_y:
            continue

        result.append(current)

    result.append(
        deduped[-1]
    )

    return result


def _route_orientation(
    a,
    b,
) -> str:
    if abs(
        b[0] - a[0]
    ) >= abs(
        b[1] - a[1]
    ):
        return "h"

    return "v"


def _lane_conflict_nearby(cell, orientation: str, occupancy: dict, radius: int = 1) -> bool:
    """True when ANY already routed pipeline occupies this local grid area.

    Earlier versions only blocked a matching orientation, which still allowed a
    new horizontal line to cross a previously routed vertical line.  For the clean
    engineering output requested here, both overlap and crossing are treated as a
    routing conflict except at the physical edge endpoints.
    """
    gx, gy = cell
    for dx in range(-radius, radius + 1):
        for dy in range(-radius, radius + 1):
            if abs(dx) + abs(dy) > radius:
                continue
            use = occupancy.get((gx + dx, gy + dy), {"h": 0, "v": 0})
            if use.get("h", 0) > 0 or use.get("v", 0) > 0:
                return True
    return False

def _orthogonal_segment_conflict(a, b, c, d, tolerance: float = 0.015):
    """Return (conflict, point) for H/V segments that overlap or cross."""
    k1 = _segment_kind(a, b)
    k2 = _segment_kind(c, d)
    if k1 == "d" or k2 == "d":
        return False, None

    if k1 == "h" and k2 == "h":
        if abs(a[1] - c[1]) > tolerance:
            return False, None
        left = max(min(a[0], b[0]), min(c[0], d[0]))
        right = min(max(a[0], b[0]), max(c[0], d[0]))
        if right - left > tolerance:
            return True, ((left + right) / 2.0, (a[1] + c[1]) / 2.0)
        return False, None

    if k1 == "v" and k2 == "v":
        if abs(a[0] - c[0]) > tolerance:
            return False, None
        top = max(min(a[1], b[1]), min(c[1], d[1]))
        bottom = min(max(a[1], b[1]), max(c[1], d[1]))
        if bottom - top > tolerance:
            return True, ((a[0] + c[0]) / 2.0, (top + bottom) / 2.0)
        return False, None

    if k1 == "v" and k2 == "h":
        return _orthogonal_segment_conflict(c, d, a, b, tolerance)

    # k1 horizontal, k2 vertical
    x = c[0]
    y = a[1]
    if (
        min(a[0], b[0]) - tolerance <= x <= max(a[0], b[0]) + tolerance
        and min(c[1], d[1]) - tolerance <= y <= max(c[1], d[1]) + tolerance
    ):
        return True, (x, y)
    return False, None


def _mark_route_occupancy(
    points,
    occupancy: dict,
) -> None:
    # Keep exact geometry as well as grid occupancy. The exact segment list is
    # necessary because component ports/stubs are not always perfectly grid-snapped.
    segment_store = occupancy.setdefault("__segments__", [])

    for a, b in zip(points, points[1:]):
        if _segment_length(a, b) <= 1e-8:
            continue
        segment_store.append((a, b))
        orientation = _route_orientation(a, b)
        ax, ay = _snap(a[0]), _snap(a[1])
        bx, by = _snap(b[0]), _snap(b[1])

        if orientation == "h":
            for gx in range(min(ax, bx), max(ax, bx) + 1):
                entry = occupancy.setdefault((gx, ay), {"h": 0, "v": 0})
                entry["h"] += 1
        else:
            for gy in range(min(ay, by), max(ay, by) + 1):
                entry = occupancy.setdefault((ax, gy), {"h": 0, "v": 0})
                entry["v"] += 1

def _nearby_occupancy_penalty(
    cell,
    orientation: str,
    occupancy: dict,
) -> tuple[float, float]:
    """Penalty for already-used exact and neighboring lanes."""
    gx, gy = cell
    same = 0.0
    crossing = 0.0

    for dx in range(-3, 4):
        for dy in range(-3, 4):
            distance = abs(dx) + abs(dy)
            if distance > 3:
                continue
            weight = 1.35 if distance == 0 else (0.75 if distance == 1 else (0.30 if distance == 2 else 0.12))
            use = occupancy.get((gx + dx, gy + dy), {"h": 0, "v": 0})
            same += weight * use[orientation]
            crossing += weight * (use["v"] if orientation == "h" else use["h"])

    return same, crossing


def _astar_route(
    start,
    end,
    obstacles,
    occupancy,
    occupancy_radius: int = 1,
) -> list[tuple[float, float]] | None:
    """Orthogonal A* with already-routed pipelines as hard geometric obstacles."""
    min_gx = _snap(0.24)
    max_gx = _snap(SLIDE_WIDTH_IN - 0.24)
    min_gy = _snap(0.96)
    max_gy = _snap(SLIDE_HEIGHT_IN - 0.24)
    start_cell = (_snap(start[0]), _snap(start[1]))
    end_cell = (_snap(end[0]), _snap(end[1]))

    blocked = set()
    for gx in range(min_gx, max_gx + 1):
        for gy in range(min_gy, max_gy + 1):
            cell = (gx, gy)
            if cell in {start_cell, end_cell}:
                continue
            point = (_unsnap(gx), _unsnap(gy))
            if any(_point_inside_expanded_box(point, box) for box in obstacles):
                blocked.add(cell)

    existing_segments = occupancy.get("__segments__", [])

    def near_endpoint_cell(cell) -> bool:
        return (
            abs(cell[0] - start_cell[0]) + abs(cell[1] - start_cell[1]) <= 3
            or abs(cell[0] - end_cell[0]) + abs(cell[1] - end_cell[1]) <= 3
        )

    def near_endpoint_point(point) -> bool:
        return min(math.hypot(point[0] - start[0], point[1] - start[1]), math.hypot(point[0] - end[0], point[1] - end[1])) <= 0.14

    start_state = (start_cell[0], start_cell[1], -1)
    serial = count()
    queue = [(0.0, next(serial), start_state)]
    best = {start_state: 0.0}
    parent = {}
    directions = [(1, 0, "h"), (-1, 0, "h"), (0, 1, "v"), (0, -1, "v")]
    final_state = None

    while queue:
        _estimated, _serial, state = heapq.heappop(queue)
        gx, gy, previous_direction = state
        if (gx, gy) == end_cell:
            final_state = state
            break
        current_cost = best[state]
        current_point = (_unsnap(gx), _unsnap(gy))

        for direction_index, (dx, dy, orientation) in enumerate(directions):
            nx, ny = gx + dx, gy + dy
            if nx < min_gx or nx > max_gx or ny < min_gy or ny > max_gy:
                continue
            next_cell = (nx, ny)
            if next_cell in blocked:
                continue
            next_point = (_unsnap(nx), _unsnap(ny))

            if _lane_conflict_nearby(next_cell, orientation, occupancy, radius=occupancy_radius) and not near_endpoint_cell(next_cell):
                continue

            geometric_conflict = False
            for a, b in existing_segments:
                conflict, point = _orthogonal_segment_conflict(current_point, next_point, a, b)
                if conflict and point is not None and not near_endpoint_point(point):
                    geometric_conflict = True
                    break
            if geometric_conflict:
                continue

            turn_penalty = 0.0 if previous_direction in {-1, direction_index} else 1.15
            same_usage, perpendicular_usage = _nearby_occupancy_penalty(next_cell, orientation, occupancy)
            next_cost = current_cost + 1.0 + turn_penalty + 7.0 * same_usage + 9.0 * perpendicular_usage
            next_state = (nx, ny, direction_index)
            if next_cost >= best.get(next_state, float("inf")):
                continue
            best[next_state] = next_cost
            parent[next_state] = state
            heuristic = abs(end_cell[0] - nx) + abs(end_cell[1] - ny)
            heapq.heappush(queue, (next_cost + heuristic, next(serial), next_state))

    if final_state is None:
        return None

    cells = []
    state = final_state
    while True:
        cells.append((state[0], state[1]))
        if state == start_state:
            break
        state = parent[state]
    cells.reverse()
    return _compress_polyline([(_unsnap(gx), _unsnap(gy)) for gx, gy in cells])

def _fallback_route(
    start,
    end,
) -> list[
    tuple[float, float]
]:
    sx, sy = start
    ex, ey = end

    if abs(ex - sx) >= abs(ey - sy):
        mid_x = (sx + ex) / 2

        return _compress_polyline(
            [
                start,
                (mid_x, sy),
                (mid_x, ey),
                end,
            ]
        )

    mid_y = (sy + ey) / 2

    return _compress_polyline(
        [
            start,
            (sx, mid_y),
            (ex, mid_y),
            end,
        ]
    )


def _outer_guides(
    hint: str,
    start,
    end,
) -> list[
    tuple[float, float]
]:
    if hint == "top_outer":
        channel_y = CONTENT_TOP + 0.18

        return [
            (start[0], channel_y),
            (end[0], channel_y),
        ]

    if hint == "bottom_outer":
        channel_y = CONTENT_BOTTOM - 0.18

        return [
            (start[0], channel_y),
            (end[0], channel_y),
        ]

    if hint == "left_outer":
        channel_x = CONTENT_LEFT + 0.18

        return [
            (channel_x, start[1]),
            (channel_x, end[1]),
        ]

    if hint == "right_outer":
        channel_x = CONTENT_RIGHT - 0.18

        return [
            (channel_x, start[1]),
            (channel_x, end[1]),
        ]

    return []


def _map_sketch_point(
    point: DiagramPoint,
    boxes=None,
) -> tuple[float, float]:
    """Map a sketch point through the same layout displacement as the cards."""
    if boxes is not None and hasattr(boxes, "map_normalized_point"):
        return boxes.map_normalized_point(point)

    return (
        CONTENT_LEFT + float(point.x) * (CONTENT_RIGHT - CONTENT_LEFT),
        CONTENT_TOP + float(point.y) * (CONTENT_BOTTOM - CONTENT_TOP),
    )


def _point_inside_box(point, box, clearance=0.02) -> bool:
    px, py = point
    x, y, w, h = box
    return (x - clearance) <= px <= (x + w + clearance) and (y - clearance) <= py <= (y + h + clearance)


def _nudge_guide_outside(point, obstacles):
    px, py = point
    for box in obstacles:
        if _point_inside_box((px, py), box, 0.02):
            x, y, w, h = box
            # push to nearest outside side with a small margin
            candidates = [
                (x - 0.10, py),
                (x + w + 0.10, py),
                (px, y - 0.10),
                (px, y + h + 0.10),
            ]
            px, py = min(candidates, key=lambda c: abs(c[0]-point[0]) + abs(c[1]-point[1]))
    px = min(max(px, CONTENT_LEFT + 0.06), CONTENT_RIGHT - 0.06)
    py = min(max(py, CONTENT_TOP + 0.06), CONTENT_BOTTOM - 0.06)
    return (px, py)


def _guides_from_edge(edge, obstacles, boxes=None):
    guides = []
    if getattr(edge, "waypoints", None):
        for wp in edge.waypoints:
            mapped = _map_sketch_point(wp, boxes)
            guides.append(_nudge_guide_outside(mapped, obstacles))
        return _compress_polyline(guides) if len(guides) >= 2 else guides
    return []


def _segment_hits_box(a, b, box, clearance=ROUTE_CLEARANCE) -> bool:
    ax, ay = a
    bx, by = b
    x, y, w, h = box

    left = x - clearance
    right = x + w + clearance
    top = y - clearance
    bottom = y + h + clearance

    if abs(ay - by) < 1e-8:
        seg_left = min(ax, bx)
        seg_right = max(ax, bx)
        return (top <= ay <= bottom and not (seg_right < left or seg_left > right))

    if abs(ax - bx) < 1e-8:
        seg_top = min(ay, by)
        seg_bottom = max(ay, by)
        return (left <= ax <= right and not (seg_bottom < top or seg_top > bottom))

    return False


def _candidate_score(points, obstacles, occupancy) -> tuple[float, bool]:
    score = 0.0
    blocked = False

    for a, b in zip(points, points[1:]):
        length = abs(b[0] - a[0]) + abs(b[1] - a[1])
        score += length

        for box in obstacles:
            if _segment_hits_box(a, b, box):
                blocked = True
                score += 12000.0

        orientation = _route_orientation(a, b)
        ax = _snap(a[0])
        ay = _snap(a[1])
        bx = _snap(b[0])
        by = _snap(b[1])

        if orientation == "h":
            for gx in range(min(ax, bx), max(ax, bx) + 1):
                cell = (gx, ay)
                use = occupancy.get(cell, {"h": 0, "v": 0})
                if use["h"] > 0 or _lane_conflict_nearby(cell, "h", occupancy, radius=1):
                    blocked = True
                    score += 7000.0
                if use["v"] > 0:
                    score += 95.0 * use["v"]
        else:
            for gy in range(min(ay, by), max(ay, by) + 1):
                cell = (ax, gy)
                use = occupancy.get(cell, {"h": 0, "v": 0})
                if use["v"] > 0 or _lane_conflict_nearby(cell, "v", occupancy, radius=1):
                    blocked = True
                    score += 7000.0
                if use["h"] > 0:
                    score += 95.0 * use["h"]

    score += max(0, len(points) - 2) * 0.55
    return score, blocked


def _direct_candidates(start, end) -> list[list[tuple[float, float]]]:
    sx, sy = start
    ex, ey = end

    candidates = [
        _compress_polyline([start, (ex, sy), end]),
        _compress_polyline([start, (sx, ey), end]),
    ]

    mid_x = (sx + ex) / 2
    mid_y = (sy + ey) / 2

    candidates.append(
        _compress_polyline([start, (mid_x, sy), (mid_x, ey), end])
    )
    candidates.append(
        _compress_polyline([start, (sx, mid_y), (ex, mid_y), end])
    )

    # Small channel offsets provide separate lanes without producing giant loops.
    lane = 0.16
    for sign in (-1, 1):
        channel_y = mid_y + sign * lane
        candidates.append(
            _compress_polyline([start, (sx, channel_y), (ex, channel_y), end])
        )

        channel_x = mid_x + sign * lane
        candidates.append(
            _compress_polyline([start, (channel_x, sy), (channel_x, ey), end])
        )

    unique = []
    seen = set()
    for candidate in candidates:
        key = tuple((round(x, 3), round(y, 3)) for x, y in candidate)
        if key not in seen:
            seen.add(key)
            unique.append(candidate)

    return unique


def _orthogonalize_polyline(
    points: list[tuple[float, float]],
    obstacles,
    occupancy,
) -> list[tuple[float, float]]:
    """Convert tiny grid/warp diagonals into clean horizontal/vertical elbows."""
    if len(points) <= 1:
        return points

    output = [points[0]]
    for point in points[1:]:
        previous = output[-1]
        if abs(previous[0] - point[0]) < 1e-8 or abs(previous[1] - point[1]) < 1e-8:
            output.append(point)
            continue

        corners = [
            (point[0], previous[1]),
            (previous[0], point[1]),
        ]
        scored = []
        for corner in corners:
            candidate = [previous, corner, point]
            score, blocked = _candidate_score(candidate, obstacles, occupancy)
            scored.append((blocked, score, corner))
        _blocked, _score, corner = min(scored, key=lambda item: (item[0], item[1]))
        if abs(corner[0] - previous[0]) > 1e-8 or abs(corner[1] - previous[1]) > 1e-8:
            output.append(corner)
        output.append(point)

    return _compress_polyline(output)


def _best_direct_route(start, end, obstacles, occupancy):
    best_points = None
    best_score = float("inf")
    best_blocked = True

    for candidate in _direct_candidates(start, end):
        score, blocked = _candidate_score(candidate, obstacles, occupancy)

        if (blocked, score) < (best_blocked, best_score):
            best_points = candidate
            best_score = score
            best_blocked = blocked

    if best_points is not None and not best_blocked:
        return best_points

    return None

def _route_via_points(
    start,
    end,
    guides,
    obstacles,
    occupancy,
) -> list[
    tuple[float, float]
]:
    anchors = [start, *guides, end]
    complete = [anchors[0]]

    for segment_start, segment_end in zip(anchors, anchors[1:]):
        segment = _astar_route(segment_start, segment_end, obstacles, occupancy)
        if segment is None:
            segment = _best_direct_route(segment_start, segment_end, obstacles, occupancy)
        if segment is None:
            segment = _fallback_route(segment_start, segment_end)
        complete.extend(segment[1:])

    return _compress_polyline(complete)


def _final_direction(
    points,
) -> str:
    if len(points) < 2:
        return "right"

    x1, y1 = points[-2]
    x2, y2 = points[-1]

    if abs(x2 - x1) >= abs(y2 - y1):
        return (
            "right"
            if x2 >= x1
            else "left"
        )

    return (
        "down"
        if y2 >= y1
        else "up"
    )



def _route_geometry_key(edge) -> tuple:
    return (
        edge.source,
        edge.target,
        edge.source_side,
        edge.target_side,
        tuple((round(point.x, 2), round(point.y, 2)) for point in edge.waypoints),
    )


def _parallel_route_offsets(diagram: DiagramSpec) -> dict[int, float]:
    """Small lane offsets only for genuinely coincident distinct physical pipes."""
    groups: dict[tuple, list[int]] = {}

    for index, edge in enumerate(diagram.edges):
        groups.setdefault(_route_geometry_key(edge), []).append(index)

    offsets: dict[int, float] = {}

    for indices in groups.values():
        if len(indices) <= 1:
            offsets[indices[0]] = 0.0
            continue

        center = (len(indices) - 1) / 2
        for lane_index, edge_index in enumerate(indices):
            offsets[edge_index] = (lane_index - center) * 0.085

    return offsets


def _offset_guides_for_parallel_pipe(
    guides: list[tuple[float, float]],
    start,
    end,
    offset: float,
) -> list[tuple[float, float]]:
    if abs(offset) < 1e-9 or not guides:
        return guides

    dx = end[0] - start[0]
    dy = end[1] - start[1]
    if abs(dx) >= abs(dy):
        return [(x, y + offset) for x, y in guides]
    return [(x + offset, y) for x, y in guides]


def _offset_internal_route(points, offset: float) -> list[tuple[float, float]]:
    if abs(offset) < 1e-9 or len(points) <= 4:
        return points

    first = points[1]
    last = points[-2]
    dx = last[0] - first[0]
    dy = last[1] - first[1]

    # Offset perpendicular to the dominant route direction. Endpoints stay fixed.
    if abs(dx) >= abs(dy):
        shifted = [points[0], points[1]]
        shifted.extend((x, y + offset) for x, y in points[2:-2])
        shifted.extend([points[-2], points[-1]])
    else:
        shifted = [points[0], points[1]]
        shifted.extend((x + offset, y) for x, y in points[2:-2])
        shifted.extend([points[-2], points[-1]])

    return shifted


def _orthogonal_locked_path(points, obstacles):
    """Turn a locked sketch polyline into clean H/V engineering-pipe geometry.

    The ordered sketch points are still honored.  We only insert a single elbow
    between two diagonal points; we never delete/reorder a real bend.
    """
    if len(points) <= 1:
        return points

    result = [points[0]]
    for target in points[1:]:
        current = result[-1]
        if abs(current[0] - target[0]) < 1e-8 or abs(current[1] - target[1]) < 1e-8:
            result.append(target)
            continue

        c1 = (target[0], current[1])       # horizontal then vertical
        c2 = (current[0], target[1])       # vertical then horizontal

        def score(corner):
            candidate = [current, corner, target]
            blocked = sum(
                1 for a, b in zip(candidate, candidate[1:])
                for box in obstacles
                if _segment_hits_box(a, b, box)
            )
            length = _segment_length(current, corner) + _segment_length(corner, target)
            return blocked, length

        corner = min((c1, c2), key=score)
        if corner != current and corner != target:
            result.append(corner)
        result.append(target)

    return _compress_polyline(result)


def _strict_edge_polyline(
    edge,
    start,
    start_stub,
    end_stub,
    end,
    boxes=None,
    lane_offset: float = 0.0,
) -> list[tuple[float, float]]:
    """Render image-derived geometry without obstacle rerouting or global warping."""
    mapped = [_map_sketch_point(wp, None) for wp in edge.waypoints] if edge.waypoints else []

    if mapped:
        anchors = [start, start_stub, *mapped, end_stub, end]
    else:
        sx, sy = start_stub
        ex, ey = end_stub
        if abs(ex - sx) >= abs(ey - sy):
            anchors = [start, start_stub, (ex, sy), end_stub, end]
        else:
            anchors = [start, start_stub, (sx, ey), end_stub, end]

    # Only apply a lane offset to the internal route. Ports remain fixed on the
    # component boundary so multiple connections enter at different anchors.
    if abs(lane_offset) > 1e-9 and len(anchors) > 4:
        shifted = [anchors[0], anchors[1]]
        for x, y in anchors[2:-2]:
            shifted.append((x, y + lane_offset))
        shifted.extend(anchors[-2:])
        anchors = shifted

    # Convert each diagonal pair to a deterministic 90-degree elbow. We DO NOT
    # inspect component obstacles here; the waypoints came from the actual image.
    result = [anchors[0]]
    for target in anchors[1:]:
        current = result[-1]
        if abs(current[0] - target[0]) < 1e-8 or abs(current[1] - target[1]) < 1e-8:
            result.append(target)
            continue
        # Prefer the elbow that follows the dominant direction of the segment.
        if abs(target[0] - current[0]) >= abs(target[1] - current[1]):
            corner = (target[0], current[1])
        else:
            corner = (current[0], target[1])
        result.extend([corner, target])
    return _compress_polyline(result)


# =============================================================================
# BUILD EDGE ROUTES
# =============================================================================

def _route_candidates_for_edge(start, start_stub, end_stub, end, source_side, target_side):
    """Generate Manhattan candidates with multiple independent routing channels."""
    sx, sy = start_stub
    ex, ey = end_stub
    candidates = [
        [start, start_stub, (ex, sy), end_stub, end],
        [start, start_stub, (sx, ey), end_stub, end],
    ]

    mid_x = (sx + ex) / 2.0
    mid_y = (sy + ey) / 2.0
    lane_offsets = (-0.72, -0.54, -0.36, -0.18, 0.18, 0.36, 0.54, 0.72)

    # Horizontal corridor alternatives.
    for delta in lane_offsets:
        channel_y = mid_y + delta
        if CONTENT_TOP + 0.08 <= channel_y <= CONTENT_BOTTOM - 0.08:
            candidates.append(
                [start, start_stub, (sx, channel_y), (ex, channel_y), end_stub, end]
            )

    # Vertical corridor alternatives.
    for delta in lane_offsets:
        channel_x = mid_x + delta
        if CONTENT_LEFT + 0.08 <= channel_x <= CONTENT_RIGHT - 0.08:
            candidates.append(
                [start, start_stub, (channel_x, sy), (channel_x, ey), end_stub, end]
            )

    # Mixed-side Z routes at the exact midpoint remain useful for short edges.
    candidates.extend(
        [
            [start, start_stub, (mid_x, sy), (mid_x, ey), end_stub, end],
            [start, start_stub, (sx, mid_y), (ex, mid_y), end_stub, end],
        ]
    )

    unique = []
    seen = set()
    for candidate in candidates:
        candidate = _compress_polyline(candidate)
        key = tuple((round(x, 3), round(y, 3)) for x, y in candidate)
        if key not in seen:
            seen.add(key)
            unique.append(candidate)
    return unique

def _route_score(points, obstacles, occupancy, own_endpoints):
    """Score a candidate; exact crossings and shared segments are hard conflicts."""
    score = 0.0
    blocked = 0
    existing_segments = occupancy.get("__segments__", [])

    def near_endpoint_point(point):
        return min(
            math.hypot(point[0] - own_endpoints[0][0], point[1] - own_endpoints[0][1]),
            math.hypot(point[0] - own_endpoints[1][0], point[1] - own_endpoints[1][1]),
        ) <= 0.14

    start_cell = (_snap(own_endpoints[0][0]), _snap(own_endpoints[0][1]))
    end_cell = (_snap(own_endpoints[1][0]), _snap(own_endpoints[1][1]))
    def near_endpoint_cell(cell):
        return (
            abs(cell[0] - start_cell[0]) + abs(cell[1] - start_cell[1]) <= 3
            or abs(cell[0] - end_cell[0]) + abs(cell[1] - end_cell[1]) <= 3
        )

    for a, b in zip(points, points[1:]):
        length = _segment_length(a, b)
        if length < 0.02:
            continue
        score += length

        for box in obstacles:
            if _segment_hits_box(a, b, box, clearance=ROUTE_CLEARANCE):
                blocked += 1
                score += 50000.0

        for c, d in existing_segments:
            conflict, point = _orthogonal_segment_conflict(a, b, c, d)
            if conflict and point is not None and not near_endpoint_point(point):
                blocked += 1
                score += 60000.0

        orientation = _route_orientation(a, b)
        ax, ay = _snap(a[0]), _snap(a[1])
        bx, by = _snap(b[0]), _snap(b[1])
        cells = (
            ((gx, ay) for gx in range(min(ax, bx), max(ax, bx) + 1))
            if orientation == "h"
            else ((ax, gy) for gy in range(min(ay, by), max(ay, by) + 1))
        )
        for cell in cells:
            if _lane_conflict_nearby(cell, orientation, occupancy, radius=1) and not near_endpoint_cell(cell):
                blocked += 1
                score += 40000.0
            same_near, crossing_near = _nearby_occupancy_penalty(cell, orientation, occupancy)
            score += 16.0 * same_near + 22.0 * crossing_near

    score += max(0, len(points) - 2) * 0.30
    return blocked, score

def _edge_route_priority_distance(edge, boxes) -> float:
    a = boxes.get(edge.source)
    b = boxes.get(edge.target)
    if a is None or b is None:
        return 999.0
    ax, ay = _box_center(a)
    bx, by = _box_center(b)
    return abs(ax - bx) + abs(ay - by)



def _route_from_final_layout(
    edge,
    start,
    start_stub,
    end_stub,
    end,
    source_side,
    target_side,
    obstacles,
    occupancy,
    boxes,
    lane_offset: float = 0.0,
):
    """Route one edge without changing source/target/topology.

    Image-derived waypoints are kept as ordered soft guides.  The actual visible
    pipe is allowed to detour around blocks and already-routed pipes so the final
    drawing remains readable instead of reproducing overlapping ink literally.
    """
    guides = _guides_from_edge(edge, obstacles, boxes)
    guides = _offset_guides_for_parallel_pipe(guides, start_stub, end_stub, lane_offset)

    # If a guide lands directly on an existing route (common in a hand sketch
    # where lines cross), move that guide to the nearest free grid point while
    # keeping its order and local position.
    clean_guides = []
    for gx, gy in guides:
        base = (_snap(gx), _snap(gy))
        if not _lane_conflict_nearby(base, "h", occupancy, radius=1):
            clean_guides.append((gx, gy))
            continue
        found = None
        for radius in range(1, 8):
            for dx in range(-radius, radius + 1):
                for dy in range(-radius, radius + 1):
                    if abs(dx) + abs(dy) != radius:
                        continue
                    cell = (base[0] + dx, base[1] + dy)
                    if _lane_conflict_nearby(cell, "h", occupancy, radius=1):
                        continue
                    point = (_unsnap(cell[0]), _unsnap(cell[1]))
                    if any(_point_inside_expanded_box(point, box) for box in obstacles):
                        continue
                    found = point
                    break
                if found:
                    break
            if found:
                break
        clean_guides.append(found or (gx, gy))
    guides = clean_guides

    if guides:
        core = _route_via_points(start_stub, end_stub, guides, obstacles, occupancy)
        if core:
            candidate = _compress_polyline([start, start_stub, *core[1:-1], end_stub, end])
            blocked, _ = _route_score(candidate, obstacles, occupancy, (start, end))
            if blocked == 0:
                return candidate

    candidates = _route_candidates_for_edge(
        start, start_stub, end_stub, end, source_side, target_side
    )
    outer = _outer_guides(edge.route_hint, start_stub, end_stub)
    if outer:
        candidates.append(_compress_polyline([start, start_stub, *outer, end_stub, end]))

    scored = sorted(
        [(_route_score(candidate, obstacles, occupancy, (start, end)), candidate) for candidate in candidates],
        key=lambda item: item[0],
    )
    for (blocked, _score), candidate in scored:
        if blocked == 0:
            return candidate

    # Hard-separated A*: first demand one empty grid-cell corridor around every
    # pipe. If that is too restrictive, retry with exact-cell separation only.
    for radius in (1, 0):
        core = _astar_route(start_stub, end_stub, obstacles, occupancy, occupancy_radius=radius)
        if core:
            candidate = _compress_polyline([start, start_stub, *core[1:-1], end_stub, end])
            blocked, _ = _route_score(candidate, obstacles, occupancy, (start, end))
            if blocked == 0:
                return candidate

    # Emergency perimeter channels.  These are dynamic and can be used by any
    # diagram; they are not tied to a specific node or sample image.
    sx, sy = start_stub
    ex, ey = end_stub
    perimeter = []
    for y in (1.00, 1.08, 1.16, 7.22, 7.14, 7.06):
        perimeter.append(_compress_polyline([start, start_stub, (sx, y), (ex, y), end_stub, end]))
    for x in (0.30, 0.40, 0.50, 13.03, 12.93, 12.83):
        perimeter.append(_compress_polyline([start, start_stub, (x, sy), (x, ey), end_stub, end]))

    perimeter.sort(key=lambda candidate: _route_score(candidate, obstacles, occupancy, (start, end)))
    for candidate in perimeter:
        blocked, _ = _route_score(candidate, obstacles, occupancy, (start, end))
        if blocked == 0:
            return candidate

    # Never drop an edge.  This last fallback is reached only for a geometrically
    # impossible/non-planar case; choose the least-conflicting route while keeping
    # the exact source and destination unchanged.
    if scored:
        return scored[0][1]
    return _compress_polyline([start, start_stub, *_fallback_route(start_stub, end_stub)[1:-1], end_stub, end])

def _segment_conflict_info(a, b, c, d, tolerance: float = 0.012):
    """Return conflict geometry for orthogonal segments.

    In addition to exact overlap/crossings, this also treats *near-parallel*
    segments as a conflict when they run side-by-side inside the same corridor with
    too little spacing. That final separation pass is what keeps visually similar
    routes such as Sump 1 -> Master 1 and Sump 2 -> Master 2 from sitting almost
    on top of each other.
    """
    k1 = _segment_kind(a, b)
    k2 = _segment_kind(c, d)
    if k1 == "d" or k2 == "d":
        return None

    parallel_gap = 0.16

    if k1 == "h" and k2 == "h":
        y_gap = abs(a[1] - c[1])
        if y_gap > parallel_gap:
            return None
        lo = max(min(a[0], b[0]), min(c[0], d[0]))
        hi = min(max(a[0], b[0]), max(c[0], d[0]))
        if hi - lo > tolerance:
            conflict_y = (a[1] + c[1]) / 2.0
            return ("parallel_h", (lo, conflict_y), (hi, conflict_y))
        return None

    if k1 == "v" and k2 == "v":
        x_gap = abs(a[0] - c[0])
        if x_gap > parallel_gap:
            return None
        lo = max(min(a[1], b[1]), min(c[1], d[1]))
        hi = min(max(a[1], b[1]), max(c[1], d[1]))
        if hi - lo > tolerance:
            conflict_x = (a[0] + c[0]) / 2.0
            return ("parallel_v", (conflict_x, lo), (conflict_x, hi))
        return None

    if k1 == "v" and k2 == "h":
        info = _segment_conflict_info(c, d, a, b, tolerance)
        if info is None:
            return None
        if info[0] == "cross":
            return info
        return info

    # first horizontal, second vertical
    x, y = c[0], a[1]
    if (
        min(a[0], b[0]) - tolerance <= x <= max(a[0], b[0]) + tolerance
        and min(c[1], d[1]) - tolerance <= y <= max(c[1], d[1]) + tolerance
    ):
        return ("cross", (x, y), (x, y))
    return None


def _route_nonendpoint_conflicts(points, accepted_segments):
    conflicts = []
    route_ends = (points[0], points[-1])
    for segment_index, (a, b) in enumerate(zip(points, points[1:])):
        for c, d in accepted_segments:
            info = _segment_conflict_info(a, b, c, d)
            if info is None:
                continue
            probe = info[1] if info[0] == "cross" else (
                (info[1][0] + info[2][0]) / 2.0,
                (info[1][1] + info[2][1]) / 2.0,
            )
            if min(
                math.hypot(probe[0] - route_ends[0][0], probe[1] - route_ends[0][1]),
                math.hypot(probe[0] - route_ends[1][0], probe[1] - route_ends[1][1]),
            ) <= 0.13:
                continue
            conflicts.append((segment_index, info))
    return conflicts


def _detour_segment(a, b, info, offset):
    """Replace one straight segment by a small rectangular lane detour."""
    kind, p1, p2 = info
    clearance = 0.055

    if _segment_kind(a, b) == "h":
        direction = 1.0 if b[0] >= a[0] else -1.0
        lo = min(p1[0], p2[0]) - clearance
        hi = max(p1[0], p2[0]) + clearance
        lo = max(min(a[0], b[0]) + 0.025, lo)
        hi = min(max(a[0], b[0]) - 0.025, hi)
        if hi <= lo:
            center = p1[0]
            lo, hi = center - 0.06, center + 0.06
        enter_x, exit_x = (lo, hi) if direction > 0 else (hi, lo)
        lane_y = a[1] + offset
        return [a, (enter_x, a[1]), (enter_x, lane_y), (exit_x, lane_y), (exit_x, a[1]), b]

    direction = 1.0 if b[1] >= a[1] else -1.0
    lo = min(p1[1], p2[1]) - clearance
    hi = max(p1[1], p2[1]) + clearance
    lo = max(min(a[1], b[1]) + 0.025, lo)
    hi = min(max(a[1], b[1]) - 0.025, hi)
    if hi <= lo:
        center = p1[1]
        lo, hi = center - 0.06, center + 0.06
    enter_y, exit_y = (lo, hi) if direction > 0 else (hi, lo)
    lane_x = a[0] + offset
    return [a, (a[0], enter_y), (lane_x, enter_y), (lane_x, exit_y), (a[0], exit_y), b]


def _candidate_route_quality(points, accepted_segments, obstacles):
    conflicts = len(_route_nonendpoint_conflicts(points, accepted_segments))
    hits = 0
    for a, b in zip(points, points[1:]):
        for box in obstacles:
            if _segment_hits_box(a, b, box, clearance=0.04):
                hits += 1
    bends = max(0, len(points) - 2)
    length = sum(_segment_length(a, b) for a, b in zip(points, points[1:]))
    return (hits, conflicts, bends, length)


def _deconflict_routes(routes, diagram: DiagramSpec, boxes):
    """Final deterministic route clean-up with a guaranteed wide-lane fallback."""
    result = list(routes)
    accepted_segments = []

    for edge_index, route in enumerate(result):
        if route is None:
            continue
        points, direction = route
        points = list(points)
        edge = diagram.edges[edge_index]
        obstacles = [box for node_id, box in boxes.items() if node_id not in {edge.source, edge.target}]

        for _ in range(18):
            conflicts = _route_nonendpoint_conflicts(points, accepted_segments)
            if not conflicts:
                break
            segment_index, info = conflicts[0]
            a, b = points[segment_index], points[segment_index + 1]
            best_points = points
            best_quality = _candidate_route_quality(points, accepted_segments, obstacles)

            for magnitude in (0.08, 0.12, 0.17, 0.23, 0.30, 0.39, 0.50):
                for sign in (-1.0, 1.0):
                    replacement = _detour_segment(a, b, info, sign * magnitude)
                    candidate = _compress_polyline(points[:segment_index] + replacement + points[segment_index + 2:])
                    quality = _candidate_route_quality(candidate, accepted_segments, obstacles)
                    if quality < best_quality:
                        best_points, best_quality = candidate, quality
                    if quality[:2] == (0, 0):
                        break
                if best_quality[:2] == (0, 0):
                    break

            if best_points == points:
                break
            points = best_points

        # If local doglegs cannot separate a geometrically trapped route, move the
        # whole internal corridor to a dynamically selected global lane. Endpoints
        # remain unchanged, so topology/source/destination are preserved exactly.
        if _route_nonendpoint_conflicts(points, accepted_segments) and len(points) >= 4:
            start, start_stub = points[0], points[1]
            end_stub, end = points[-2], points[-1]
            global_candidates = []

            # Horizontal corridors, including outer page lanes and the free space
            # between components. Each candidate consists only of straight H/V lines.
            y_values = [1.00 + i * 0.10 for i in range(0, 63)]
            for y in y_values:
                if y > 7.20:
                    break
                candidate = _compress_polyline([
                    start, start_stub,
                    (start_stub[0], y),
                    (end_stub[0], y),
                    end_stub, end,
                ])
                global_candidates.append(candidate)

            x_values = [0.30 + i * 0.12 for i in range(0, 107)]
            for x in x_values:
                if x > 13.03:
                    break
                candidate = _compress_polyline([
                    start, start_stub,
                    (x, start_stub[1]),
                    (x, end_stub[1]),
                    end_stub, end,
                ])
                global_candidates.append(candidate)

            current_quality = _candidate_route_quality(points, accepted_segments, obstacles)
            best_points = points
            best_quality = current_quality
            for candidate in global_candidates:
                quality = _candidate_route_quality(candidate, accepted_segments, obstacles)
                if quality < best_quality:
                    best_points, best_quality = candidate, quality
                if quality[:2] == (0, 0):
                    break
            points = best_points

        result[edge_index] = (points, direction)
        accepted_segments.extend(list(zip(points, points[1:])))

    return result



def _selection_mode_row_groups(diagram: DiagramSpec, boxes):
    """Cluster current component cards into visual rows without moving them."""
    visible = []
    for node in diagram.nodes:
        box = boxes.get(node.id)
        if box is None or node.node_type == "junction":
            continue
        cx, cy = _box_center(box)
        visible.append((node.id, cx, cy, box))

    visible.sort(key=lambda item: (item[2], item[1]))
    rows: list[dict] = []

    # Cards created by Component Selection mode are intentionally layered. A
    # 0.58-inch grouping threshold keeps each existing layer together even after
    # the collision solver makes tiny vertical adjustments.
    threshold = 0.58
    for node_id, cx, cy, box in visible:
        best_row = None
        best_distance = float("inf")
        for row in rows:
            distance = abs(cy - row["center_y"])
            if distance <= threshold and distance < best_distance:
                best_row = row
                best_distance = distance

        if best_row is None:
            rows.append({
                "items": [(node_id, cx, cy, box)],
                "center_y": cy,
                "top": box[1],
                "bottom": box[1] + box[3],
            })
            continue

        best_row["items"].append((node_id, cx, cy, box))
        count = len(best_row["items"])
        best_row["center_y"] = (
            (best_row["center_y"] * (count - 1)) + cy
        ) / count
        best_row["top"] = min(best_row["top"], box[1])
        best_row["bottom"] = max(best_row["bottom"], box[1] + box[3])

    rows.sort(key=lambda row: row["center_y"])
    row_index_by_node: dict[str, int] = {}

    for row_index, row in enumerate(rows):
        row["items"].sort(key=lambda item: item[1])
        row["ids"] = [item[0] for item in row["items"]]
        row["left"] = min(item[3][0] for item in row["items"])
        row["right"] = max(item[3][0] + item[3][2] for item in row["items"])
        for node_id, _cx, _cy, _box in row["items"]:
            row_index_by_node[node_id] = row_index

    return rows, row_index_by_node


def _selection_mode_gap_lanes(
    top: float,
    bottom: float,
    preferred: float | None = None,
    spacing: float = 0.16,
) -> list[float]:
    """Return evenly spaced horizontal lanes inside one clear vertical gap."""
    top = max(CONTENT_TOP + 0.08, float(top))
    bottom = min(CONTENT_BOTTOM - 0.08, float(bottom))
    if bottom <= top:
        return []

    height = bottom - top
    margin = min(0.12, max(0.04, height * 0.16))
    lo = top + margin
    hi = bottom - margin
    if hi < lo:
        lo = hi = (top + bottom) / 2.0

    lanes: list[float] = []
    if abs(hi - lo) < 1e-8:
        lanes = [lo]
    else:
        count = max(1, int((hi - lo) / max(spacing, 0.08)) + 1)
        if count == 1:
            lanes = [(lo + hi) / 2.0]
        else:
            actual = (hi - lo) / (count - 1)
            lanes = [lo + index * actual for index in range(count)]

    if preferred is not None:
        lanes.sort(key=lambda value: abs(value - preferred))

    return lanes


def _selection_mode_candidate_lanes(
    rows,
    source_row: int,
    target_row: int,
) -> list[float]:
    """Return clean lane candidates for one connection.

    Same-row connections prefer the empty space directly below that row (the style
    seen in the clean Master→Display examples), then above it. Cross-row
    connections use every available gap between the source and target rows.
    """
    lanes: list[float] = []

    if source_row == target_row:
        row = rows[source_row]

        # Prefer below for every row except the final row. This produces the same
        # visual language as the user's clean downward Master→Display connections.
        if source_row + 1 < len(rows):
            top = row["bottom"] + 0.12
            bottom = rows[source_row + 1]["top"] - 0.12
            preferred = top + max(0.16, (bottom - top) * 0.30)
            lanes.extend(_selection_mode_gap_lanes(top, bottom, preferred))

        if source_row > 0:
            top = rows[source_row - 1]["bottom"] + 0.12
            bottom = row["top"] - 0.12
            preferred = bottom - max(0.16, (bottom - top) * 0.30)
            lanes.extend(_selection_mode_gap_lanes(top, bottom, preferred))

        # Outer-page lanes are only fallback choices when the normal inter-row
        # corridor cannot provide a conflict-free route.
        lanes.extend([
            max(CONTENT_TOP + 0.12, row["top"] - 0.34),
            min(CONTENT_BOTTOM - 0.12, row["bottom"] + 0.34),
        ])
    else:
        upper = min(source_row, target_row)
        lower = max(source_row, target_row)

        # Use each gap between the involved rows. For adjacent rows this naturally
        # becomes one clean family of parallel horizontal tracks.
        for gap_index in range(upper, lower):
            top = rows[gap_index]["bottom"] + 0.12
            bottom = rows[gap_index + 1]["top"] - 0.12
            preferred = (top + bottom) / 2.0
            lanes.extend(_selection_mode_gap_lanes(top, bottom, preferred))

    # Stable unique order.
    unique: list[float] = []
    seen = set()
    for lane in lanes:
        key = round(lane, 3)
        if key in seen:
            continue
        seen.add(key)
        unique.append(lane)
    return unique


def _selection_mode_segment_conflicts(
    points,
    accepted_segments,
    endpoints,
) -> int:
    """Count non-endpoint overlap/crossing conflicts against accepted routes."""
    count_conflicts = 0
    source_end, target_end = endpoints

    def close_to_own_endpoint(point) -> bool:
        return min(
            math.hypot(point[0] - source_end[0], point[1] - source_end[1]),
            math.hypot(point[0] - target_end[0], point[1] - target_end[1]),
        ) <= 0.11

    for a, b in zip(points, points[1:]):
        if _segment_length(a, b) <= 1e-8:
            continue
        for c, d in accepted_segments:
            conflict, point = _orthogonal_segment_conflict(a, b, c, d, tolerance=0.012)
            if conflict and point is not None and not close_to_own_endpoint(point):
                count_conflicts += 1

    return count_conflicts


def _selection_mode_route_quality(
    points,
    boxes,
    ignore_nodes: set[str],
    accepted_segments,
) -> tuple[int, int, int, int, float]:
    """Score one simple orthogonal route; lower is better.

    In addition to hard intersections, keep visible breathing room between
    horizontal tracks. This prevents the almost-touching parallel lines that can
    still look like an overlap even when their geometry is technically separate.
    """
    component_hits = 0
    for a, b in zip(points, points[1:]):
        for node_id, box in boxes.items():
            if node_id in ignore_nodes:
                continue
            if _segment_hits_box(a, b, box, clearance=0.055):
                component_hits += 1

    conflicts = _selection_mode_segment_conflicts(
        points,
        accepted_segments,
        (points[0], points[-1]),
    )

    spacing_violations = 0
    for a, b in zip(points, points[1:]):
        if _segment_kind(a, b) != "h" or _segment_length(a, b) < 0.28:
            continue
        for c, d in accepted_segments:
            if _segment_kind(c, d) != "h" or _segment_length(c, d) < 0.28:
                continue
            # Reserve a clear lane only when the two horizontal spans actually
            # overlap. Adjacent chain segments such as Master <-> Repeater and
            # Repeater <-> Transmitter may stay on the same y-axis because they
            # meet only at the component between them and do not share a path.
            if abs(a[1] - c[1]) < 0.22:
                overlap_left = max(min(a[0], b[0]), min(c[0], d[0]))
                overlap_right = min(max(a[0], b[0]), max(c[0], d[0]))
                if overlap_right - overlap_left > 0.03:
                    spacing_violations += 1

    bends = max(0, len(points) - 2)
    length = sum(_segment_length(a, b) for a, b in zip(points, points[1:]))
    return component_hits, conflicts, spacing_violations, bends, length


def _selection_mode_endpoint_direction_valid(
    points,
    source_side: str,
    target_side: str,
    tolerance: float = 0.015,
) -> bool:
    """Ensure a route leaves/enters each card through the selected side.

    This prevents a compressed candidate from making a U-turn through the source
    or target card after an outward stub has been added.
    """
    if len(points) < 2:
        return False

    start = points[0]
    next_point = points[1]
    previous_point = points[-2]
    end = points[-1]

    if source_side == "left" and next_point[0] > start[0] + tolerance:
        return False
    if source_side == "right" and next_point[0] < start[0] - tolerance:
        return False
    if source_side == "top" and next_point[1] > start[1] + tolerance:
        return False
    if source_side == "bottom" and next_point[1] < start[1] - tolerance:
        return False

    if target_side == "left" and previous_point[0] > end[0] + tolerance:
        return False
    if target_side == "right" and previous_point[0] < end[0] - tolerance:
        return False
    if target_side == "top" and previous_point[1] > end[1] + tolerance:
        return False
    if target_side == "bottom" and previous_point[1] < end[1] - tolerance:
        return False

    return True


def _selection_mode_port_plan(diagram: DiagramSpec, boxes):
    """Universal side-center port planner for Component Selection mode.

    This planner is intentionally component-independent.  It never checks labels,
    codes, node types, or connection-rule names.  Side choice is based only on:
      * final source/destination geometry,
      * how well a side faces the other component,
      * route obstruction / path length,
      * current side usage at both endpoints, and
      * endpoint degree / nearby connection pressure.

    Every visible endpoint remains exactly at the center of the selected side.
    """
    side_info: list[tuple[str, str] | None] = [None] * len(diagram.edges)
    fractions: dict[tuple[str, int, str, str], float] = {}
    totals: dict[tuple[str, int, str, str], tuple[int, int]] = {}
    all_sides = ("top", "bottom", "left", "right")

    def center_port(box, side: str):
        return _port_point(box, side, 0, 1, 0.5)

    def side_vector(side: str) -> tuple[float, float]:
        return {
            "left": (-1.0, 0.0),
            "right": (1.0, 0.0),
            "top": (0.0, -1.0),
            "bottom": (0.0, 1.0),
        }[side]

    def side_rank_for_vector(box, other_box) -> list[str]:
        sx, sy = _box_center(box)
        tx, ty = _box_center(other_box)
        dx = tx - sx
        dy = ty - sy
        distance = max(1e-9, math.hypot(dx, dy))
        ux, uy = dx / distance, dy / distance

        # Highest dot product points most directly toward the other component.
        return sorted(
            all_sides,
            key=lambda side: (
                -(side_vector(side)[0] * ux + side_vector(side)[1] * uy),
                all_sides.index(side),
            ),
        )

    def pair_score(
        source_id: str,
        target_id: str,
        source_box,
        target_box,
        source_side: str,
        target_side: str,
        preference_penalty: float,
    ):
        start = center_port(source_box, source_side)
        end = center_port(target_box, target_side)
        start_stub = _stub_point(start, source_side, 0.11)
        end_stub = _stub_point(end, target_side, 0.11)
        candidates = _selection_mode_direct_candidates(start, start_stub, end_stub, end)
        obstacles = [
            box for node_id, box in boxes.items()
            if node_id not in {source_id, target_id}
        ]

        best = None
        for candidate in candidates:
            if not _selection_mode_endpoint_direction_valid(candidate, source_side, target_side):
                continue
            hits = 0
            near_components = 0
            for a, b in zip(candidate, candidate[1:]):
                for box in obstacles:
                    if _segment_hits_box(a, b, box, clearance=0.05):
                        hits += 1
                    elif _segment_hits_box(a, b, box, clearance=0.15):
                        near_components += 1
            length = sum(_segment_length(a, b) for a, b in zip(candidate, candidate[1:]))
            bends = max(0, len(candidate) - 2)
            score = (
                hits,
                length + bends * 0.11 + near_components * 0.42 + preference_penalty,
                bends,
            )
            if best is None or score < best:
                best = score
        return best or (999, 999.0, 999)

    degree: dict[str, int] = {node_id: 0 for node_id in boxes}
    for edge in diagram.edges:
        if edge.source in degree:
            degree[edge.source] += 1
        if edge.target in degree:
            degree[edge.target] += 1

    # High-degree endpoints may need more distinct sides to prevent many links
    # from sharing one exit corridor.  The limit grows gradually and never exceeds
    # the four physical card sides.
    max_side_count = {
        # Use as many distinct physical sides as the endpoint degree justifies,
        # up to the four real card sides. This avoids forcing two unrelated links
        # to share one center-port corridor when another clean side is available.
        node_id: min(4, max(1, int(value)))
        for node_id, value in degree.items()
    }

    used_sides: dict[str, set[str]] = {node_id: set() for node_id in boxes}
    side_use_count: dict[tuple[str, str], int] = {}

    def route_order_key(edge_index: int):
        edge = diagram.edges[edge_index]
        source_box = boxes.get(edge.source)
        target_box = boxes.get(edge.target)
        if source_box is None or target_box is None:
            return (999, 999.0, edge_index)
        sx, sy = _box_center(source_box)
        tx, ty = _box_center(target_box)
        pressure = max(degree.get(edge.source, 0), degree.get(edge.target, 0))
        distance = abs(tx - sx) + abs(ty - sy)
        return (-pressure, distance, edge_index)

    for edge_index in sorted(range(len(diagram.edges)), key=route_order_key):
        edge = diagram.edges[edge_index]
        source_box = boxes.get(edge.source)
        target_box = boxes.get(edge.target)
        if source_box is None or target_box is None:
            continue

        source_rank = side_rank_for_vector(source_box, target_box)
        target_rank = side_rank_for_vector(target_box, source_box)
        source_used = used_sides.setdefault(edge.source, set())
        target_used = used_sides.setdefault(edge.target, set())
        source_limit = max_side_count.get(edge.source, 2)
        target_limit = max_side_count.get(edge.target, 2)

        candidates: list[tuple[tuple, str, str]] = []
        for source_side in all_sides:
            if source_side not in source_used and len(source_used) >= source_limit:
                continue
            for target_side in all_sides:
                if target_side not in target_used and len(target_used) >= target_limit:
                    continue

                source_pref = source_rank.index(source_side)
                target_pref = target_rank.index(target_side)
                facing_penalty = source_pref * 0.15 + target_pref * 0.15

                source_count = side_use_count.get((edge.source, source_side), 0)
                target_count = side_use_count.get((edge.target, target_side), 0)

                # Prefer a clean unused facing side when available. As degree grows,
                # reuse becomes increasingly expensive, which naturally distributes
                # bulk fan-in/fan-out across the available card sides.
                reuse_penalty = (source_count ** 1.45 + target_count ** 1.45) * 0.20
                if source_side not in source_used and len(source_used) < source_limit:
                    reuse_penalty -= 0.06
                if target_side not in target_used and len(target_used) < target_limit:
                    reuse_penalty -= 0.06

                base = pair_score(
                    edge.source, edge.target, source_box, target_box,
                    source_side, target_side, facing_penalty,
                )
                score = (
                    base[0],
                    base[1] + reuse_penalty,
                    base[2],
                    source_count + target_count,
                    source_pref + target_pref,
                )
                candidates.append((score, source_side, target_side))

        if candidates:
            candidates.sort(key=lambda item: item[0])
            _score, source_side, target_side = candidates[0]
        else:
            # Defensive fallback: choose the best-facing already-used side.
            source_pool = list(source_used) or source_rank
            target_pool = list(target_used) or target_rank
            source_side = min(source_pool, key=lambda side: source_rank.index(side))
            target_side = min(target_pool, key=lambda side: target_rank.index(side))

        side_info[edge_index] = (source_side, target_side)
        source_used.add(source_side)
        target_used.add(target_side)
        side_use_count[(edge.source, source_side)] = side_use_count.get((edge.source, source_side), 0) + 1
        side_use_count[(edge.target, target_side)] = side_use_count.get((edge.target, target_side), 0) + 1

        source_key = ("source", edge_index, edge.source, source_side)
        target_key = ("target", edge_index, edge.target, target_side)
        fractions[source_key] = 0.5
        fractions[target_key] = 0.5
        totals[source_key] = (0, 1)
        totals[target_key] = (0, 1)

    return side_info, fractions, totals

def _selection_mode_route_for_lane(
    start,
    start_stub,
    end_stub,
    end,
    lane_y: float,
):
    return _compress_polyline([
        start,
        start_stub,
        (start_stub[0], lane_y),
        (end_stub[0], lane_y),
        end_stub,
        end,
    ])


def _selection_mode_choose_separated_lanes(
    candidate_lanes: list[float],
    count: int,
    minimum_gap: float = 0.22,
) -> list[float]:
    """Choose visibly separate lanes for edges sharing one component endpoint.

    The returned lanes are deterministic and never intentionally reuse the same
    horizontal track. This is what prevents two Sump→Display connections from
    visually merging into one line.
    """
    if count <= 0:
        return []

    unique: list[float] = []
    for lane in candidate_lanes:
        if all(abs(lane - existing) >= minimum_gap for existing in unique):
            unique.append(lane)
            if len(unique) >= count:
                return unique

    # If the available corridor is narrow, keep the paths separate with a smaller
    # but still clearly visible gap before considering synthetic lanes.
    for gap in (0.20, 0.18, 0.16):
        unique = []
        for lane in candidate_lanes:
            if all(abs(lane - existing) >= gap for existing in unique):
                unique.append(lane)
                if len(unique) >= count:
                    return unique

    if candidate_lanes:
        center = candidate_lanes[0]
    else:
        center = (CONTENT_TOP + CONTENT_BOTTOM) / 2.0

    # Last-resort synthetic family centred around the preferred corridor.
    gap = 0.18
    start = center - gap * (count - 1) / 2.0
    lanes = [start + index * gap for index in range(count)]
    return [
        min(max(lane, CONTENT_TOP + 0.10), CONTENT_BOTTOM - 0.10)
        for lane in lanes
    ]


def _selection_mode_shared_endpoint_lane_plan(
    diagram: DiagramSpec,
    boxes,
    rows,
    row_index_by_node,
    side_info,
) -> dict[int, float]:
    """Reserve one independent lane for each edge in a shared-endpoint bundle.

    Edges are bundled only when they share the same source/target component side
    *and* use the same visual row corridor. This keeps unrelated connections
    untouched while guaranteeing that repeated connections such as:

        Sump 1 -> Display 1
        Sump 2 -> Display 1

    never collapse onto a single visible path.
    """
    groups: dict[tuple, list[int]] = {}

    for edge_index, edge in enumerate(diagram.edges):
        if edge_index >= len(side_info) or side_info[edge_index] is None:
            continue
        source_side, target_side = side_info[edge_index]
        source_row = row_index_by_node.get(edge.source)
        target_row = row_index_by_node.get(edge.target)
        if source_row is None or target_row is None:
            continue

        corridor = (min(source_row, target_row), max(source_row, target_row))
        groups.setdefault(
            ("source", edge.source, source_side, corridor), []
        ).append(edge_index)
        groups.setdefault(
            ("target", edge.target, target_side, corridor), []
        ).append(edge_index)

    # An edge can belong to both a shared-source and shared-target group. Prefer
    # whichever group contains more edges; ties prefer target bundling because
    # visually merged paths most often occur while converging into one target.
    chosen_group_for_edge: dict[int, tuple] = {}
    for key, indices in groups.items():
        if len(indices) <= 1:
            continue
        for edge_index in indices:
            previous = chosen_group_for_edge.get(edge_index)
            if previous is None:
                chosen_group_for_edge[edge_index] = key
                continue
            previous_count = len(groups[previous])
            current_count = len(indices)
            previous_role = previous[0]
            current_role = key[0]
            if (
                current_count > previous_count
                or (current_count == previous_count and current_role == "target" and previous_role != "target")
            ):
                chosen_group_for_edge[edge_index] = key

    active_groups: dict[tuple, list[int]] = {}
    for edge_index, key in chosen_group_for_edge.items():
        active_groups.setdefault(key, []).append(edge_index)

    reserved_lane_by_edge: dict[int, float] = {}

    for key, indices in active_groups.items():
        role, node_id, side, corridor = key
        if len(indices) <= 1:
            continue

        def other_axis(edge_index: int) -> tuple[float, float, int]:
            edge = diagram.edges[edge_index]
            other_id = edge.target if role == "source" else edge.source
            other_box = boxes.get(other_id)
            if other_box is None:
                return 0.0, 0.0, edge_index
            ox, oy = _box_center(other_box)
            # For top/bottom ports, left-to-right order is the cleanest fan-out.
            # For left/right ports, top-to-bottom order is the cleanest fan-out.
            primary = ox if side in {"top", "bottom"} else oy
            secondary = oy if side in {"top", "bottom"} else ox
            return primary, secondary, edge_index

        ordered = sorted(indices, key=other_axis)
        first_edge = diagram.edges[ordered[0]]
        sr = row_index_by_node.get(first_edge.source, corridor[0])
        tr = row_index_by_node.get(first_edge.target, corridor[1])
        candidates = _selection_mode_candidate_lanes(rows, sr, tr)
        separated = _selection_mode_choose_separated_lanes(candidates, len(ordered))

        # Preserve spatial order: left/top source gets the first reserved lane,
        # right/bottom source gets the next. The lanes themselves remain distinct.
        for position, edge_index in enumerate(ordered):
            reserved_lane_by_edge[edge_index] = separated[position]

    return reserved_lane_by_edge


def _selection_mode_strict_astar(
    start,
    end,
    obstacles,
    occupancy,
    spacing_radius: int = 1,
):
    """Fast orthogonal A* used only to repair a conflicting selection route."""
    min_gx = _snap(0.24)
    max_gx = _snap(SLIDE_WIDTH_IN - 0.24)
    min_gy = _snap(0.96)
    max_gy = _snap(SLIDE_HEIGHT_IN - 0.24)
    start_cell = (_snap(start[0]), _snap(start[1]))
    end_cell = (_snap(end[0]), _snap(end[1]))

    blocked = set()
    for gx in range(min_gx, max_gx + 1):
        for gy in range(min_gy, max_gy + 1):
            cell = (gx, gy)
            if cell in {start_cell, end_cell}:
                continue
            point = (_unsnap(gx), _unsnap(gy))
            if any(_point_inside_expanded_box(point, box, clearance=0.07) for box in obstacles):
                blocked.add(cell)
                continue
            occupied = False
            for dx in range(-spacing_radius, spacing_radius + 1):
                for dy in range(-spacing_radius, spacing_radius + 1):
                    if abs(dx) + abs(dy) > spacing_radius:
                        continue
                    use = occupancy.get((gx + dx, gy + dy), {"h": 0, "v": 0})
                    if use.get("h", 0) > 0 or use.get("v", 0) > 0:
                        occupied = True
                        break
                if occupied:
                    break
            if occupied:
                blocked.add(cell)

    start_state = (start_cell[0], start_cell[1], -1)
    serial = count()
    queue = [(0.0, next(serial), start_state)]
    best = {start_state: 0.0}
    parent = {}
    directions = [(1, 0, 0), (-1, 0, 1), (0, 1, 2), (0, -1, 3)]
    final_state = None

    while queue:
        _estimated, _serial, state = heapq.heappop(queue)
        gx, gy, previous_direction = state
        if (gx, gy) == end_cell:
            final_state = state
            break
        current_cost = best[state]
        for dx, dy, direction_index in directions:
            nx, ny = gx + dx, gy + dy
            if nx < min_gx or nx > max_gx or ny < min_gy or ny > max_gy:
                continue
            if (nx, ny) in blocked and (nx, ny) != end_cell:
                continue
            turn_penalty = 0.0 if previous_direction in {-1, direction_index} else 1.25
            next_cost = current_cost + 1.0 + turn_penalty
            next_state = (nx, ny, direction_index)
            if next_cost >= best.get(next_state, float("inf")):
                continue
            best[next_state] = next_cost
            parent[next_state] = state
            heuristic = abs(end_cell[0] - nx) + abs(end_cell[1] - ny)
            heapq.heappush(queue, (next_cost + heuristic, next(serial), next_state))

    if final_state is None:
        return None

    cells = []
    state = final_state
    while True:
        cells.append((state[0], state[1]))
        if state == start_state:
            break
        state = parent[state]
    cells.reverse()
    return _compress_polyline([(_unsnap(gx), _unsnap(gy)) for gx, gy in cells])


def _selection_mode_route_has_conflict(points, accepted_segments) -> bool:
    for a, b in zip(points, points[1:]):
        for c, d in accepted_segments:
            conflict, _point = _orthogonal_segment_conflict(a, b, c, d, tolerance=0.012)
            if conflict:
                return True
    return False


def _selection_mode_repair_conflicts(diagram: DiagramSpec, boxes, routes):
    """Reroute only edges that still cross/merge after the normal clean router."""
    side_info, fractions, totals = _selection_mode_port_plan(diagram, boxes)
    repaired = [None] * len(routes)
    occupancy = {}
    accepted_segments = []

    for edge_index, route in enumerate(routes):
        if route is None:
            continue
        points, direction = route
        edge = diagram.edges[edge_index]

        if not _selection_mode_route_has_conflict(points, accepted_segments):
            repaired[edge_index] = route
            _mark_route_occupancy(points, occupancy)
            accepted_segments.extend(zip(points, points[1:]))
            continue

        source_box = boxes.get(edge.source)
        target_box = boxes.get(edge.target)
        info = side_info[edge_index] if edge_index < len(side_info) else None
        if source_box is None or target_box is None or info is None:
            repaired[edge_index] = route
            _mark_route_occupancy(points, occupancy)
            accepted_segments.extend(zip(points, points[1:]))
            continue

        source_side, target_side = info
        source_position, source_total = totals[("source", edge_index, edge.source, source_side)]
        target_position, target_total = totals[("target", edge_index, edge.target, target_side)]
        start = _port_point(
            source_box, source_side, source_position, source_total,
            fractions[("source", edge_index, edge.source, source_side)],
        )
        end = _port_point(
            target_box, target_side, target_position, target_total,
            fractions[("target", edge_index, edge.target, target_side)],
        )

        all_obstacles = list(boxes.values())
        best_candidate = None
        best_quality = None

        for source_distance in (0.16, 0.28, 0.40, 0.52, 0.64):
            start_escape = _stub_point(start, source_side, source_distance)
            if _selection_mode_route_has_conflict([start, start_escape], accepted_segments):
                continue
            for target_distance in (0.16, 0.28, 0.40, 0.52, 0.64):
                end_escape = _stub_point(end, target_side, target_distance)
                if _selection_mode_route_has_conflict([end_escape, end], accepted_segments):
                    continue

                core = None
                for radius in (1, 0):
                    core = _selection_mode_strict_astar(
                        start_escape, end_escape, all_obstacles, occupancy, radius
                    )
                    if core is not None:
                        break
                if core is None:
                    continue

                candidate = _compress_polyline([
                    start, start_escape, *core[1:-1], end_escape, end
                ])
                candidate = _orthogonalize_polyline(candidate, all_obstacles, occupancy)
                candidate = _compress_polyline(candidate)
                if not _selection_mode_endpoint_direction_valid(
                    candidate, source_side, target_side
                ):
                    continue
                quality = _selection_mode_route_quality(
                    candidate, boxes, {edge.source, edge.target}, accepted_segments
                )
                if best_quality is None or quality < best_quality:
                    best_quality = quality
                    best_candidate = candidate
                if quality[0] == 0 and quality[1] == 0 and quality[2] == 0:
                    break
            if best_quality is not None and best_quality[0] == 0 and best_quality[1] == 0 and best_quality[2] == 0:
                break

        chosen = best_candidate if best_candidate is not None else points
        repaired_direction = (
            _final_direction(list(reversed(chosen)))
            if edge.direction == "target_to_source"
            else _final_direction(chosen)
        )
        repaired[edge_index] = (chosen, repaired_direction)
        _mark_route_occupancy(chosen, occupancy)
        accepted_segments.extend(zip(chosen, chosen[1:]))

    return repaired


def _selection_mode_direct_candidates(
    start,
    start_stub,
    end_stub,
    end,
) -> list[list[tuple[float, float]]]:
    """Return short orthogonal routes contained near the source/target pair."""
    sx, sy = start_stub
    ex, ey = end_stub
    candidates: list[list[tuple[float, float]]] = []

    # Exact straight line whenever the selected ports line up.
    if abs(sy - ey) <= 0.012 or abs(sx - ex) <= 0.012:
        candidates.append(_compress_polyline([start, start_stub, end_stub, end]))

    # Shortest possible Manhattan alternatives.  The first version continues in
    # the source port's natural outward direction before turning; stable scoring
    # preserves this preference when two alternatives have identical length.
    candidates.extend([
        _compress_polyline([start, start_stub, (ex, sy), end_stub, end]),
        _compress_polyline([start, start_stub, (sx, ey), end_stub, end]),
    ])

    mid_x = (sx + ex) / 2.0
    mid_y = (sy + ey) / 2.0
    candidates.extend([
        _compress_polyline([start, start_stub, (mid_x, sy), (mid_x, ey), end_stub, end]),
        _compress_polyline([start, start_stub, (sx, mid_y), (ex, mid_y), end_stub, end]),
    ])

    # A few small local parallel offsets solve ordinary two-line fan-out cases
    # without ever introducing the large perimeter loops seen in older versions.
    for offset in (0.16, 0.26, 0.38):
        for sign in (-1.0, 1.0):
            lane_y = mid_y + sign * offset
            lane_x = mid_x + sign * offset
            candidates.append(_compress_polyline([
                start, start_stub, (sx, lane_y), (ex, lane_y), end_stub, end
            ]))
            candidates.append(_compress_polyline([
                start, start_stub, (lane_x, sy), (lane_x, ey), end_stub, end
            ]))

    unique = []
    seen = set()
    for candidate in candidates:
        key = tuple((round(x, 3), round(y, 3)) for x, y in candidate)
        if key in seen:
            continue
        seen.add(key)
        unique.append(candidate)
    return unique


def _selection_mode_local_detour_candidates(
    start,
    start_stub,
    end_stub,
    end,
    boxes,
    ignore_nodes: set[str],
    accepted_segments,
) -> list[list[tuple[float, float]]]:
    """Build short obstacle-edge lanes around actual local blockers.

    Instead of falling back to a page-top/page-bottom channel, candidate lanes are
    taken just outside component edges and existing route segments.  Therefore a
    blocked direct connection normally detours only a few tenths of an inch around
    the thing that blocks it.
    """
    sx, sy = start_stub
    ex, ey = end_stub
    lane_ys: list[float] = []
    lane_xs: list[float] = []
    component_clearance = 0.12
    route_clearance = 0.14

    for node_id, box in boxes.items():
        if node_id in ignore_nodes:
            continue
        x, y, w, h = box
        lane_ys.extend([y - component_clearance, y + h + component_clearance])
        lane_xs.extend([x - component_clearance, x + w + component_clearance])

    for a, b in accepted_segments:
        min_x, max_x = sorted((a[0], b[0]))
        min_y, max_y = sorted((a[1], b[1]))
        if _segment_kind(a, b) == "v":
            lane_ys.extend([min_y - route_clearance, max_y + route_clearance])
            lane_xs.extend([a[0] - route_clearance, a[0] + route_clearance])
        elif _segment_kind(a, b) == "h":
            lane_ys.extend([a[1] - route_clearance, a[1] + route_clearance])
            lane_xs.extend([min_x - route_clearance, max_x + route_clearance])

    candidates: list[list[tuple[float, float]]] = []

    # Most useful first: lanes closest to the natural midpoint of this pair.
    lane_ys = sorted(
        {
            round(value, 4): value
            for value in lane_ys
            if CONTENT_TOP + 0.08 <= value <= CONTENT_BOTTOM - 0.08
        }.values(),
        key=lambda value: abs(value - (sy + ey) / 2.0),
    )
    lane_xs = sorted(
        {
            round(value, 4): value
            for value in lane_xs
            if CONTENT_LEFT + 0.08 <= value <= CONTENT_RIGHT - 0.08
        }.values(),
        key=lambda value: abs(value - (sx + ex) / 2.0),
    )

    # Limit the search to the nearest useful local corridors.  This keeps routing
    # deterministic and avoids considering distant page edges unless truly needed.
    for lane_y in lane_ys[:18]:
        candidates.append(_compress_polyline([
            start, start_stub, (sx, lane_y), (ex, lane_y), end_stub, end
        ]))
    for lane_x in lane_xs[:18]:
        candidates.append(_compress_polyline([
            start, start_stub, (lane_x, sy), (lane_x, ey), end_stub, end
        ]))

    unique = []
    seen = set()
    for candidate in candidates:
        key = tuple((round(x, 3), round(y, 3)) for x, y in candidate)
        if key in seen:
            continue
        seen.add(key)
        unique.append(candidate)
    return unique


def _selection_mode_direct_quality(
    points,
    boxes,
    ignore_nodes: set[str],
    accepted_segments,
    source_side: str,
    target_side: str,
):
    """Score a route with hard collision checks and a strong short-path bias."""
    if not _selection_mode_endpoint_direction_valid(
        points, source_side, target_side
    ):
        return 999, 999, 9999.0, 999, 999, 9999.0

    component_hits = 0
    for a, b in zip(points, points[1:]):
        for node_id, box in boxes.items():
            if node_id in ignore_nodes:
                continue
            if _segment_hits_box(a, b, box, clearance=0.05):
                component_hits += 1

    conflicts = _selection_mode_segment_conflicts(
        points,
        accepted_segments,
        (points[0], points[-1]),
    )

    spacing_violations = 0
    for a, b in zip(points, points[1:]):
        kind = _segment_kind(a, b)
        if kind not in {"h", "v"} or _segment_length(a, b) < 0.28:
            continue
        for c, d in accepted_segments:
            if _segment_kind(c, d) != kind or _segment_length(c, d) < 0.28:
                continue
            if kind == "h":
                gap = abs(a[1] - c[1])
                overlap = min(max(a[0], b[0]), max(c[0], d[0])) - max(min(a[0], b[0]), min(c[0], d[0]))
            else:
                gap = abs(a[0] - c[0])
                overlap = min(max(a[1], b[1]), max(c[1], d[1])) - max(min(a[1], b[1]), min(c[1], d[1]))
            if gap < 0.16 and overlap > 0.05:
                spacing_violations += 1

    bends = max(0, len(points) - 2)
    length = sum(_segment_length(a, b) for a, b in zip(points, points[1:]))

    # Component intersections and true line crossings/overlaps are hard failures.
    # After those are avoided, physical distance dominates.  Bend/spacing costs
    # are intentionally small so they cannot justify a multi-inch loop.
    soft_cost = length + bends * 0.10 + spacing_violations * 0.22
    return component_hits, conflicts, soft_cost, bends, spacing_violations, length


def _build_selection_mode_edge_routes(diagram: DiagramSpec, boxes):
    """Direct-first connection router for Component Selection mode only.

    Existing layout, UI, connection intent, component images and allowed rules are
    untouched.  The router uses the current component positions and chooses the
    shortest clean horizontal/vertical Manhattan path from the nearest card sides.
    Local obstacle-edge detours are considered before the existing A* fallback.
    """
    side_info, fractions, totals = _selection_mode_port_plan(diagram, boxes)
    routes = [None] * len(diagram.edges)
    accepted_segments: list[tuple[tuple[float, float], tuple[float, float]]] = []
    occupancy: dict = {}
    topology_managed = any(
        node.node_type == "junction"
        and str(getattr(node, "topology_role", "")).endswith("_junction")
        for node in diagram.nodes
    )

    def geometric_distance(edge_index: int):
        edge = diagram.edges[edge_index]
        source_box = boxes.get(edge.source)
        target_box = boxes.get(edge.target)
        if source_box is None or target_box is None:
            return (float("inf"), edge_index)
        sx, sy = _box_center(source_box)
        tx, ty = _box_center(target_box)
        # Establish short local relationships first; long links then route around
        # them instead of forcing a nearby connection into an outer-page corridor.
        return (abs(tx - sx) + abs(ty - sy), edge_index)

    route_order = sorted(range(len(diagram.edges)), key=geometric_distance)

    for edge_index in route_order:
        edge = diagram.edges[edge_index]
        info = side_info[edge_index] if edge_index < len(side_info) else None
        source_box = boxes.get(edge.source)
        target_box = boxes.get(edge.target)
        if info is None or source_box is None or target_box is None:
            continue

        source_side, target_side = info
        source_position, source_total = totals[("source", edge_index, edge.source, source_side)]
        target_position, target_total = totals[("target", edge_index, edge.target, target_side)]

        start = _port_point(
            source_box,
            source_side,
            source_position,
            source_total,
            fractions[("source", edge_index, edge.source, source_side)],
        )
        end = _port_point(
            target_box,
            target_side,
            target_position,
            target_total,
            fractions[("target", edge_index, edge.target, target_side)],
        )

        # Short outward stubs keep the connection clear of card borders/labels and
        # guarantee approach from the intended nearest side.
        start_stub = _stub_point(start, source_side, 0.10)
        end_stub = _stub_point(end, target_side, 0.10)
        ignore_nodes = {edge.source, edge.target}

        candidates = _selection_mode_direct_candidates(
            start,
            start_stub,
            end_stub,
            end,
        )
        candidates.extend(
            _selection_mode_local_detour_candidates(
                start,
                start_stub,
                end_stub,
                end,
                boxes,
                ignore_nodes,
                accepted_segments,
            )
        )

        scored = [
            (
                _selection_mode_direct_quality(
                    candidate,
                    boxes,
                    ignore_nodes,
                    accepted_segments,
                    source_side,
                    target_side,
                ),
                candidate,
            )
            for candidate in candidates
        ]
        # Stable key-only sort preserves the deliberate candidate ordering when
        # two paths have the same geometry score.
        scored.sort(key=lambda item: item[0])

        chosen = scored[0][1] if scored else _fallback_route(start, end)
        chosen_quality = scored[0][0] if scored else (999, 999, 999.0, 999, 999, 999.0)

        # A* is now a true last resort.  It is invoked only if all local direct and
        # obstacle-edge candidates still hit a component or another connection.
        if (not topology_managed) and (chosen_quality[0] > 0 or chosen_quality[1] > 0):
            all_obstacles = list(boxes.values())
            core = _selection_mode_strict_astar(
                start_stub,
                end_stub,
                all_obstacles,
                occupancy,
                spacing_radius=0,
            )
            if core is not None:
                astar_candidate = _compress_polyline([
                    start,
                    start_stub,
                    *core[1:-1],
                    end_stub,
                    end,
                ])
                astar_candidate = _orthogonalize_polyline(
                    astar_candidate,
                    all_obstacles,
                    occupancy,
                )
                astar_candidate = _compress_polyline(astar_candidate)
                astar_quality = _selection_mode_direct_quality(
                    astar_candidate,
                    boxes,
                    ignore_nodes,
                    accepted_segments,
                    source_side,
                    target_side,
                )
                if astar_quality < chosen_quality:
                    chosen = astar_candidate
                    chosen_quality = astar_quality

        direction = (
            _final_direction(list(reversed(chosen)))
            if edge.direction == "target_to_source"
            else _final_direction(chosen)
        )
        routes[edge_index] = (chosen, direction)
        _mark_route_occupancy(chosen, occupancy)
        accepted_segments.extend(list(zip(chosen, chosen[1:])))

    # Topology-managed graphs already contain explicit distribution/collection
    # spines. Running the expensive global A*/repair pass over every tiny trunk
    # segment can turn a 100-edge diagram into a minutes-long operation and can
    # undo the clean shared-main-line structure. Direct/local routing is therefore
    # authoritative for those graphs. Legacy selection diagrams retain the repair
    # pass unchanged.
    if topology_managed:
        return routes
    return _selection_mode_repair_conflicts(diagram, boxes, routes)


def _visually_separate_shared_connection_segments(routes, lane_gap: float = 0.18):
    """Return visually separated copies of routed polylines.

    The router remains the source of truth.  This post-processing step only
    separates *collinear overlapping* segments that belong to different edges so
    two logical one-to-one connections never collapse into one visible shared
    stroke.  Endpoints are preserved exactly; only internal display geometry is
    offset into deterministic parallel lanes.

    ``routes`` is the existing list of ``(points, direction)`` values.
    """
    if not routes:
        return routes

    eps = 1e-6
    prepared = []
    segment_refs = []

    for edge_index, route in enumerate(routes):
        if route is None:
            prepared.append(None)
            continue
        points, direction = route
        clean = [(float(x), float(y)) for x, y in points]
        prepared.append((clean, direction))
        for seg_index, (a, b) in enumerate(zip(clean, clean[1:])):
            ax, ay = a
            bx, by = b
            if abs(ay - by) <= eps and abs(ax - bx) > eps:
                lo, hi = sorted((ax, bx))
                segment_refs.append({
                    "edge": edge_index, "seg": seg_index, "ori": "H",
                    "coord": (ay + by) / 2.0, "lo": lo, "hi": hi,
                })
            elif abs(ax - bx) <= eps and abs(ay - by) > eps:
                lo, hi = sorted((ay, by))
                segment_refs.append({
                    "edge": edge_index, "seg": seg_index, "ori": "V",
                    "coord": (ax + bx) / 2.0, "lo": lo, "hi": hi,
                })

    # Build conflict components only between segments that are on the same
    # centreline and overlap by a real distance.  Merely touching at a bend does
    # not create a visual lane.
    offsets = {}
    by_line = {}
    coord_precision = 5
    for ref in segment_refs:
        key = (ref["ori"], round(ref["coord"], coord_precision))
        by_line.setdefault(key, []).append(ref)

    for _line_key, refs in by_line.items():
        if len(refs) < 2:
            continue

        n = len(refs)
        adjacency = [set() for _ in range(n)]
        for i in range(n):
            ri = refs[i]
            for j in range(i + 1, n):
                rj = refs[j]
                if ri["edge"] == rj["edge"]:
                    continue
                overlap = min(ri["hi"], rj["hi"]) - max(ri["lo"], rj["lo"])
                if overlap > eps:
                    adjacency[i].add(j)
                    adjacency[j].add(i)

        seen = set()
        for start in range(n):
            if start in seen or not adjacency[start]:
                continue
            stack = [start]
            component = []
            seen.add(start)
            while stack:
                current = stack.pop()
                component.append(current)
                for nxt in adjacency[current]:
                    if nxt not in seen:
                        seen.add(nxt)
                        stack.append(nxt)

            edge_ids = sorted({refs[idx]["edge"] for idx in component})
            if len(edge_ids) < 2:
                continue
            lane_by_edge = {
                edge_id: (pos - (len(edge_ids) - 1) / 2.0) * float(lane_gap)
                for pos, edge_id in enumerate(edge_ids)
            }
            for idx in component:
                ref = refs[idx]
                offsets[(ref["edge"], ref["seg"])] = lane_by_edge[ref["edge"]]

    if not offsets:
        return prepared

    result = []
    for edge_index, route in enumerate(prepared):
        if route is None:
            result.append(None)
            continue
        points, direction = route
        if len(points) < 2:
            result.append((points, direction))
            continue

        seg_lines = []
        for seg_index, (a, b) in enumerate(zip(points, points[1:])):
            ax, ay = a
            bx, by = b
            delta = float(offsets.get((edge_index, seg_index), 0.0))
            if abs(ay - by) <= eps:
                seg_lines.append(("H", ay + delta))
            elif abs(ax - bx) <= eps:
                seg_lines.append(("V", ax + delta))
            else:
                # Existing routes are orthogonal. Keep an unexpected diagonal
                # segment unchanged rather than altering router semantics.
                seg_lines.append(("D", (a, b)))

        new_points = [points[0]]

        # Join the fixed source endpoint to the potentially offset first lane.
        first_ori, first_value = seg_lines[0]
        sx, sy = points[0]
        if first_ori == "H" and abs(sy - first_value) > eps:
            new_points.append((sx, first_value))
        elif first_ori == "V" and abs(sx - first_value) > eps:
            new_points.append((first_value, sy))

        for seg_index in range(len(seg_lines) - 1):
            ori_a, value_a = seg_lines[seg_index]
            ori_b, value_b = seg_lines[seg_index + 1]
            original_vertex = points[seg_index + 1]

            if ori_a == "H" and ori_b == "V":
                vertex = (value_b, value_a)
            elif ori_a == "V" and ori_b == "H":
                vertex = (value_a, value_b)
            else:
                # This should only occur for an unexpected diagonal or two
                # collinear segments. Preserve the original vertex exactly.
                vertex = original_vertex
            if not new_points or (abs(new_points[-1][0] - vertex[0]) > eps or abs(new_points[-1][1] - vertex[1]) > eps):
                new_points.append(vertex)

        # Join the final lane back to the fixed target endpoint.
        ex, ey = points[-1]
        last_ori, last_value = seg_lines[-1]
        if last_ori == "H" and abs(ey - last_value) > eps:
            bridge = (ex, last_value)
            if abs(new_points[-1][0] - bridge[0]) > eps or abs(new_points[-1][1] - bridge[1]) > eps:
                new_points.append(bridge)
        elif last_ori == "V" and abs(ex - last_value) > eps:
            bridge = (last_value, ey)
            if abs(new_points[-1][0] - bridge[0]) > eps or abs(new_points[-1][1] - bridge[1]) > eps:
                new_points.append(bridge)

        if abs(new_points[-1][0] - ex) > eps or abs(new_points[-1][1] - ey) > eps:
            new_points.append((ex, ey))

        result.append((_compress_polyline(new_points), direction))

    return result

def build_edge_routes(diagram: DiagramSpec, boxes):
    """Route every DiagramEdge once with separate lanes and preserved topology.

    Component Selection mode uses the universal reserved-channel router.  The
    sketch/AI workflow keeps its existing waypoint-preserving renderer so exact
    hand-drawn geometry remains backward compatible.
    """
    if "Component Selection mode" in (getattr(diagram, "style_notes", []) or []):
        return plan_connection_routes(
            diagram,
            boxes,
            (CONTENT_LEFT, CONTENT_RIGHT, CONTENT_TOP, CONTENT_BOTTOM),
        )

    node_lookup = {node.id: node for node in diagram.nodes}
    degree = {node.id: 0 for node in diagram.nodes}
    for edge in diagram.edges:
        degree[edge.source] = degree.get(edge.source, 0) + 1
        degree[edge.target] = degree.get(edge.target, 0) + 1

    usage = {}
    side_info = []

    for edge_index, edge in enumerate(diagram.edges):
        source_box = boxes.get(edge.source)
        target_box = boxes.get(edge.target)
        if source_box is None or target_box is None:
            side_info.append(None)
            continue

        fallback_source, fallback_target = _choose_sides(source_box, target_box)
        source_side = edge.source_side or fallback_source
        target_side = edge.target_side or fallback_target
        side_info.append((source_side, target_side))
        usage.setdefault(("source", edge.source, source_side), []).append(edge_index)
        usage.setdefault(("target", edge.target, target_side), []).append(edge_index)

    port_index = {}
    fractions = {}
    for (role, node_id, side), indices in usage.items():
        ordered = sorted(
            indices,
            key=lambda idx: _port_sort_key(diagram.edges[idx], role, side, boxes),
        )
        total = len(ordered)
        node = node_lookup.get(node_id)
        values = []

        for position, edge_index in enumerate(ordered):
            edge = diagram.edges[edge_index]
            anchor = edge.source_anchor if role == "source" else edge.target_anchor
            raw = _anchor_fraction(node, side, anchor) if node is not None else None
            if raw is None:
                raw = PORT_MARGIN + (1.0 - 2.0 * PORT_MARGIN) * ((position + 1) / (total + 1))
            values.append(float(raw))

        if total > 1:
            usable = 1.0 - 2.0 * PORT_MARGIN
            min_gap = min(0.14, usable / max(total + 1, 2) * 0.84)
            for position in range(1, total):
                values[position] = max(values[position], values[position - 1] + min_gap)
            overflow = values[-1] - (1.0 - PORT_MARGIN)
            if overflow > 0:
                values = [value - overflow for value in values]
            for position in range(total - 2, -1, -1):
                values[position] = min(values[position], values[position + 1] - min_gap)
            underflow = PORT_MARGIN - values[0]
            if underflow > 0:
                values = [value + underflow for value in values]

        for position, edge_index in enumerate(ordered):
            port_index[(role, edge_index, node_id, side)] = (position, total)
            fractions[(role, edge_index, node_id, side)] = min(
                max(values[position], PORT_MARGIN),
                1.0 - PORT_MARGIN,
            )

    routes = [None] * len(diagram.edges)
    occupancy = {}
    lane_offsets = _parallel_route_offsets(diagram)

    # Route the most constrained edges first: explicit sketch geometry, high-degree
    # manifold connections, then long connections. This prevents local edges from
    # consuming the only clean corridor needed by a long physical pipe.
    def priority(edge_index: int):
        edge = diagram.edges[edge_index]
        return (
            1 if edge.waypoints else 0,
            len(edge.waypoints),
            degree.get(edge.source, 0) + degree.get(edge.target, 0),
            _edge_route_priority_distance(edge, boxes),
            -edge_index,
        )

    route_order = sorted(range(len(diagram.edges)), key=priority, reverse=True)

    for edge_index in route_order:
        edge = diagram.edges[edge_index]
        info = side_info[edge_index]
        if info is None:
            continue

        source_side, target_side = info
        source_box = boxes[edge.source]
        target_box = boxes[edge.target]
        source_index, source_total = port_index[("source", edge_index, edge.source, source_side)]
        target_index, target_total = port_index[("target", edge_index, edge.target, target_side)]

        start = _port_point(
            source_box,
            source_side,
            source_index,
            source_total,
            fractions[("source", edge_index, edge.source, source_side)],
        )
        end = _port_point(
            target_box,
            target_side,
            target_index,
            target_total,
            fractions[("target", edge_index, edge.target, target_side)],
        )
        start_stub = _stub_point(
            start,
            source_side,
            _stub_distance(source_index, source_total) + 0.12,
        )
        end_stub = _stub_point(
            end,
            target_side,
            _stub_distance(target_index, target_total) + 0.12,
        )

        obstacles = [
            box
            for node_id, box in boxes.items()
            if node_id not in {edge.source, edge.target}
            and (node_lookup.get(node_id) is None or node_lookup[node_id].node_type != "junction")
        ]

        chosen = _route_from_final_layout(
            edge,
            start,
            start_stub,
            end_stub,
            end,
            source_side,
            target_side,
            obstacles,
            occupancy,
            boxes,
            lane_offsets.get(edge_index, 0.0),
        )
        chosen = _orthogonalize_polyline(chosen, obstacles, occupancy)
        chosen = _compress_polyline(chosen)

        _mark_route_occupancy(chosen, occupancy)
        direction = (
            _final_direction(list(reversed(chosen)))
            if edge.direction == "target_to_source"
            else _final_direction(chosen)
        )
        routes[edge_index] = (chosen, direction)

    # Exact geometric deconfliction catches the remaining non-grid intersections
    # and moves only the conflicting local segment into a nearby independent lane.
    routes = _deconflict_routes(routes, diagram, boxes)
    return routes

def _segment_kind(a, b):
    if abs(a[1] - b[1]) < 1e-8 and abs(a[0] - b[0]) >= 1e-8:
        return "h"
    if abs(a[0] - b[0]) < 1e-8 and abs(a[1] - b[1]) >= 1e-8:
        return "v"
    return "d"



def _separate_exact_parallel_routes(routes, diagram):
    """Separate only genuinely coincident internal pipe corridors.

    Routes that cross are not moved. Ports are never moved. A small lane offset is
    applied only when two or more edges have the same long internal segment.
    """
    result = list(routes)
    groups = {}

    for i, route in enumerate(result):
        if route is None:
            continue

        points, direction = route

        for si, (a, b) in enumerate(zip(points, points[1:])):
            kind = _segment_kind(a, b)
            if kind == "d" or _segment_length(a, b) < 0.30:
                continue

            if si == 0 or si == len(points) - 2:
                continue

            if kind == "h":
                key = (
                    kind,
                    round((a[1] + b[1]) / 2, 3),
                    round(min(a[0], b[0]), 3),
                    round(max(a[0], b[0]), 3),
                )
            else:
                key = (
                    kind,
                    round((a[0] + b[0]) / 2, 3),
                    round(min(a[1], b[1]), 3),
                    round(max(a[1], b[1]), 3),
                )

            groups.setdefault(key, []).append((i, si))

    for members in groups.values():
        if len(members) < 2:
            continue

        center = (len(members) - 1) / 2.0

        for lane, (route_index, segment_index) in enumerate(members):
            offset = (lane - center) * 0.040
            if abs(offset) < 1e-9:
                continue

            points, direction = result[route_index]
            points = list(points)
            a = points[segment_index]
            b = points[segment_index + 1]

            if _segment_kind(a, b) == "h":
                points[segment_index] = (
                    a[0],
                    a[1] + offset,
                )
                points[segment_index + 1] = (
                    b[0],
                    b[1] + offset,
                )
            else:
                points[segment_index] = (
                    a[0] + offset,
                    a[1],
                )
                points[segment_index + 1] = (
                    b[0] + offset,
                    b[1],
                )

            result[route_index] = (
                _compress_polyline(points),
                direction,
            )

    return result


# =============================================================================
# LABEL PLACEMENT
# =============================================================================

def _rect_overlap(a, b, padding=0.0) -> bool:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return not (
        ax + aw + padding <= bx
        or bx + bw + padding <= ax
        or ay + ah + padding <= by
        or by + bh + padding <= ay
    )


def _segment_length(a, b) -> float:
    return abs(b[0] - a[0]) + abs(b[1] - a[1])


def _label_size(text: str) -> tuple[float, float]:
    clean = (text or "").strip()
    if not clean:
        return 0.0, 0.0
    if len(clean) <= 10:
        return 1.08, 0.34
    if len(clean) <= 18:
        return 1.48, 0.38
    if len(clean) <= 28:
        return 2.02, 0.42
    if len(clean) <= 42:
        return 2.48, 0.48
    return 2.90, 0.56


def _segment_intersects_rect(a, b, rect) -> bool:
    rx, ry, rw, rh = rect
    min_x = min(a[0], b[0])
    max_x = max(a[0], b[0])
    min_y = min(a[1], b[1])
    max_y = max(a[1], b[1])
    return not (
        max_x < rx
        or min_x > rx + rw
        or max_y < ry
        or min_y > ry + rh
    )

def _detail_tag_rects(
    diagram: DiagramSpec,
    boxes,
) -> list[tuple[float, float, float, float]]:
    rects = []

    for node in diagram.nodes:
        box = boxes.get(node.id)
        if box is None or node.node_type == "junction":
            continue

        x, y, w, _h = box
        for index, detail in enumerate(_node_details(node)):
            tag_w = min(1.30, max(0.55, 0.050 * len(detail) + 0.24))
            tag_h = 0.19
            tag_x = x + w - tag_w
            tag_y = max(0.90, y - 0.20 - index * 0.20)
            rects.append((tag_x, tag_y, tag_w, tag_h))

    return rects


def _point_to_segment_distance(point, a, b):
    px, py = point
    ax, ay = a
    bx, by = b

    dx = bx - ax
    dy = by - ay
    length_sq = dx * dx + dy * dy

    if length_sq <= 1e-12:
        return math.hypot(px - ax, py - ay), a

    t = ((px - ax) * dx + (py - ay) * dy) / length_sq
    t = min(1.0, max(0.0, t))
    closest = (ax + t * dx, ay + t * dy)
    return math.hypot(px - closest[0], py - closest[1]), closest


def _closest_route_point(point, points):
    best_distance = float("inf")
    best_point = points[0]

    for a, b in zip(points, points[1:]):
        distance, candidate = _point_to_segment_distance(point, a, b)
        if distance < best_distance:
            best_distance = distance
            best_point = candidate

    return best_distance, best_point


def _label_rect_is_free(
    rect,
    component_rects,
    tag_rects,
    used_labels,
    all_route_segments,
    own_edge_index,
) -> bool:
    if (
        rect[0] < CONTENT_LEFT
        or rect[1] < CONTENT_TOP
        or rect[0] + rect[2] > CONTENT_RIGHT
        or rect[1] + rect[3] > CONTENT_BOTTOM
    ):
        return False

    if any(_rect_overlap(rect, component, 0.055) for component in component_rects):
        return False

    if any(_rect_overlap(rect, tag, 0.035) for tag in tag_rects):
        return False

    if any(_rect_overlap(rect, other, 0.075) for other in used_labels):
        return False

    for other_index, segments in enumerate(all_route_segments):
        if other_index == own_edge_index:
            continue
        for a, b in segments:
            if _segment_intersects_rect(a, b, rect):
                return False

    return True


def place_edge_labels(
    diagram: DiagramSpec,
    routes,
    boxes,
):
    """
    Always place labels in non-overlapping free space.

    Return:
        None
        or (text, rect, leader_anchor)

    A leader_anchor is used when the nearest collision-free position is not
    directly beside the owning pipe, so the reader can still identify the pipe.
    """

    component_rects = [
        box
        for node_id, box in boxes.items()
        if next(
            (node.node_type for node in diagram.nodes if node.id == node_id),
            "other",
        ) != "junction"
    ]
    tag_rects = _detail_tag_rects(diagram, boxes)

    all_route_segments = []
    for route in routes:
        if route is None:
            all_route_segments.append([])
        else:
            points, _direction = route
            all_route_segments.append(list(zip(points, points[1:])))

    used_labels = []
    placements = []

    for edge_index, (edge, route) in enumerate(zip(diagram.edges, routes)):
        text = format_edge_label(edge)
        if not text or route is None:
            placements.append(None)
            continue

        points, _direction = route
        label_w, label_h = _label_size(text)
        candidates = []

        # Candidate positions around long segments first.
        segment_records = sorted(
            [
                (_segment_length(a, b), a, b)
                for a, b in zip(points, points[1:])
                if _segment_length(a, b) >= 0.25
            ],
            key=lambda item: item[0],
            reverse=True,
        )

        for length, start, end in segment_records:
            horizontal = abs(end[0] - start[0]) >= abs(end[1] - start[1])

            for t in (0.50, 0.30, 0.70, 0.18, 0.82):
                route_x = start[0] + (end[0] - start[0]) * t
                route_y = start[1] + (end[1] - start[1]) * t

                offsets = (
                    (0.20, -0.20, 0.31, -0.31, 0.43, -0.43, 0.56, -0.56)
                    if horizontal
                    else (0.22, -0.22, 0.34, -0.34, 0.47, -0.47, 0.60, -0.60)
                )

                for offset in offsets:
                    if horizontal:
                        center_x = route_x
                        center_y = route_y + offset
                    else:
                        center_x = route_x + offset
                        center_y = route_y

                    rect = (
                        center_x - label_w / 2,
                        center_y - label_h / 2,
                        label_w,
                        label_h,
                    )

                    if _label_rect_is_free(
                        rect,
                        component_rects,
                        tag_rects,
                        used_labels,
                        all_route_segments,
                        edge_index,
                    ):
                        distance, anchor = _closest_route_point((center_x, center_y), points)
                        score = distance + abs(offset) * 0.20 - min(length, 2.0) * 0.10
                        candidates.append((score, rect, anchor))

        # If pipe-adjacent candidates are full, search a wider free-space grid.
        if not candidates:
            route_center_x = sum(point[0] for point in points) / len(points)
            route_center_y = sum(point[1] for point in points) / len(points)

            for radius in (0.55, 0.80, 1.05, 1.35, 1.70, 2.10, 2.55, 3.00):
                for angle_index in range(16):
                    angle = (math.pi * 2 * angle_index) / 16
                    center_x = route_center_x + math.cos(angle) * radius
                    center_y = route_center_y + math.sin(angle) * radius * 0.65

                    rect = (
                        center_x - label_w / 2,
                        center_y - label_h / 2,
                        label_w,
                        label_h,
                    )

                    if not _label_rect_is_free(
                        rect,
                        component_rects,
                        tag_rects,
                        used_labels,
                        all_route_segments,
                        edge_index,
                    ):
                        continue

                    distance, anchor = _closest_route_point((center_x, center_y), points)
                    candidates.append((distance + radius * 0.05, rect, anchor))

                if candidates:
                    break

        # Absolute fallback: scan the drawable area. This prevents dropped labels.
        if not candidates:
            y = CONTENT_TOP + 0.20
            while y < CONTENT_BOTTOM - 0.20 and not candidates:
                x = CONTENT_LEFT + 0.20
                while x < CONTENT_RIGHT - 0.20:
                    rect = (x - label_w / 2, y - label_h / 2, label_w, label_h)
                    if _label_rect_is_free(
                        rect,
                        component_rects,
                        tag_rects,
                        used_labels,
                        all_route_segments,
                        edge_index,
                    ):
                        distance, anchor = _closest_route_point((x, y), points)
                        candidates.append((distance + 2.0, rect, anchor))
                        break
                    x += 0.18
                y += 0.16

        if not candidates:
            placements.append(None)
            continue

        _score, rect, anchor = min(candidates, key=lambda item: item[0])
        used_labels.append(rect)
        placements.append((text, rect, anchor))

    return placements


# =============================================================================
# PPT HEADER
# =============================================================================

def add_reference_header_ppt(
    slide,
    title: str,
) -> None:
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(SLIDE_WIDTH_IN), Inches(0.78),
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(7, 60, 120)
    header.line.fill.background()

    brand = slide.shapes.add_textbox(Inches(0.45), Inches(0.16), Inches(3.6), Inches(0.36))
    p = brand.text_frame.paragraphs[0]
    p.text = "Agentic Sketch-to-PPT"
    p.font.name = "Aptos Display"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*WHITE)

    title_box = slide.shapes.add_textbox(Inches(4.05), Inches(0.14), Inches(7.1), Inches(0.42))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.size = Pt(14.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*WHITE)
    p.alignment = PP_ALIGN.CENTER

    ready = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.45), Inches(0.18), Inches(1.45), Inches(0.34))
    ready.fill.solid()
    ready.fill.fore_color.rgb = RGBColor(229, 249, 238)
    ready.line.color.rgb = RGBColor(176, 232, 198)
    p = ready.text_frame.paragraphs[0]
    p.text = "✓ Diagram Ready"
    p.font.name = "Aptos"
    p.font.size = Pt(7.8)
    p.font.bold = True
    p.font.color.rgb = RGBColor(19, 138, 75)
    p.alignment = PP_ALIGN.CENTER


# =============================================================================
# PPT PIPE / LABEL
# =============================================================================

def add_ppt_segment(
    slide,
    start,
    end,
    *,
    color: tuple[int, int, int] = WATER_FLOW_COLOR,
    dotted: bool = False,
) -> None:
    """Draw one routed PPT segment without changing its geometry.

    Water-flow edges retain the existing blue solid engineering-pipe treatment.
    Control/signal edges use a black dotted stroke over the same white clearance
    underlay so they remain readable on top of the slide background.
    """
    white_line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(start[0]),
        Inches(start[1]),
        Inches(end[0]),
        Inches(end[1]),
    )
    white_line.line.color.rgb = RGBColor(*WHITE)
    white_line.line.width = Pt(1.8)

    if not dotted:
        halo = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(start[0]),
            Inches(start[1]),
            Inches(end[0]),
            Inches(end[1]),
        )
        halo.line.color.rgb = RGBColor(*WATER_FLOW_DARK)
        halo.line.width = Pt(1.15)

    pipe = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(start[0]),
        Inches(start[1]),
        Inches(end[0]),
        Inches(end[1]),
    )
    pipe.line.color.rgb = RGBColor(*color)
    pipe.line.width = Pt(0.80 if not dotted else 0.85)
    pipe.line.dash_style = (
        MSO_LINE_DASH_STYLE.ROUND_DOT if dotted else MSO_LINE_DASH_STYLE.SOLID
    )


def add_ppt_arrow(
    slide,
    point,
    direction: str,
) -> None:
    """Draw a larger, high-contrast flow-direction marker."""
    size = 0.28

    arrow = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(point[0] - size / 2),
        Inches(point[1] - size / 2),
        Inches(size),
        Inches(size),
    )

    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(*ARROW_COLOR)
    arrow.line.color.rgb = RGBColor(*ARROW_OUTLINE)
    arrow.line.width = Pt(1.25)

    # Keep the arrowhead in its native upright orientation.
    # The route geometry/direction is still calculated normally, but the
    # visual marker must never rotate with the route.
    arrow.rotation = 0


def add_ppt_edge_label(
    slide,
    text: str,
    rect,
    leader_anchor=None,
) -> None:
    x, y, w, h = rect
    center = (x + w / 2, y + h / 2)

    if leader_anchor is not None:
        distance = math.hypot(center[0] - leader_anchor[0], center[1] - leader_anchor[1])
        if distance > 0.34:
            leader = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(leader_anchor[0]),
                Inches(leader_anchor[1]),
                Inches(center[0]),
                Inches(center[1]),
            )
            leader.line.color.rgb = RGBColor(120, 120, 120)
            leader.line.width = Pt(0.65)

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )

    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(243, 248, 255)
    box.line.color.rgb = RGBColor(174, 205, 242)
    box.line.width = Pt(0.65)

    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.025)
    frame.margin_right = Inches(0.025)
    frame.margin_top = Inches(0.012)
    frame.margin_bottom = Inches(0.012)

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = "Aptos"
    paragraph.font.bold = True
    paragraph.font.size = Pt(13.0 if len(text) <= 24 else 11.2)
    paragraph.font.color.rgb = RGBColor(*TITLE_BLUE)


# =============================================================================
# PPT COMPONENT
# =============================================================================

def _node_details(
    node: DiagramNode,
) -> list[str]:
    return [
        str(detail).strip()
        for detail in node.details[:3]
        if str(detail).strip()
    ]


def add_ppt_detail_tags(
    slide,
    node: DiagramNode,
    x: float,
    y: float,
    width: float,
) -> None:
    for index, detail in enumerate(
        _node_details(node)
    ):
        tag_w = min(
            1.30,
            max(
                0.55,
                0.050 * len(detail)
                + 0.24,
            ),
        )

        tag_h = 0.19

        tag_x = (
            x + width - tag_w
        )

        tag_y = max(
            0.90,
            y - 0.20
            - index * 0.20,
        )

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(tag_x),
            Inches(tag_y),
            Inches(tag_w),
            Inches(tag_h),
        )

        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(
            *TAG_BG
        )

        shape.line.color.rgb = RGBColor(
            *TAG_BORDER
        )

        shape.line.width = Pt(0.5)

        paragraph = (
            shape.text_frame.paragraphs[0]
        )

        paragraph.text = detail
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(7.2)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(
            *BLACK
        )


def add_ppt_component(
    slide,
    node: DiagramNode,
    box,
    assets_dir: Path,
) -> None:
    x, y, w, h = box

    if node.node_type == "junction":
        size = 0.085
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + (w - size) / 2),
            Inches(y + (h - size) / 2),
            Inches(size),
            Inches(size),
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor(*PIPE_DARK)
        dot.line.fill.background()
        return

    caption_h = min(0.40, max(0.24, h * 0.32))
    image_h = max(0.18, h - caption_h)

    # Subtle shadow card.
    shadow = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x + 0.04), Inches(y + 0.04), Inches(w), Inches(h),
    )
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = RGBColor(226, 235, 246)
    shadow.line.fill.background()

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(*WHITE)
    card.line.color.rgb = RGBColor(*CARD_BORDER)
    card.line.width = Pt(0.9)

    image_path = node_component_asset_path(assets_dir, node)
    if image_path:
        add_picture_contained_ppt(
            slide, image_path,
            x + 0.05, y + 0.04,
            max(0.05, w - 0.10), max(0.05, image_h - 0.07),
        )
    else:
        # Vector fallback for selectable components without dedicated artwork.
        # This keeps every catalog item visually complete while allowing optional
        # real PNG/JPG assets to override the fallback automatically.
        badge_w = min(max(0.42, w * 0.52), 0.72)
        badge_h = min(max(0.30, image_h * 0.55), 0.50)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x + (w - badge_w) / 2),
            Inches(y + max(0.04, (image_h - badge_h) / 2)),
            Inches(badge_w),
            Inches(badge_h),
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(233, 244, 255)
        badge.line.color.rgb = RGBColor(*LABEL_BLUE)
        badge.line.width = Pt(1.2)
        tf = badge.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = _component_code(node)
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Aptos Display"
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(*TITLE_BLUE)

    caption_y = y + h - caption_h
    caption = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x + 0.04), Inches(caption_y - 0.02), Inches(max(0.05, w - 0.08)), Inches(max(0.05, caption_h - 0.04)),
    )
    caption.fill.solid()
    caption.fill.fore_color.rgb = RGBColor(*LABEL_BLUE)
    caption.line.fill.background()

    frame = caption.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.025)
    frame.margin_right = Inches(0.025)
    frame.margin_top = Inches(0.012)
    frame.margin_bottom = Inches(0.012)

    lines = wrap_component_label(component_display_label(node), 20)[:2]
    details = _node_details(node)[:1]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.name = "Aptos"
        paragraph.font.bold = True
        paragraph.font.size = Pt(12.4 if len(lines) == 1 else 10.8)
        paragraph.font.color.rgb = RGBColor(*WHITE)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)

    if details:
        paragraph = frame.add_paragraph()
        paragraph.text = details[0]
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(7.7)
        paragraph.font.color.rgb = RGBColor(226, 239, 255)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)


def add_ppt_wifi_badge(slide, box) -> None:
    """Overlay one clear Wi-Fi badge at a component's top-right corner."""
    x, y, w, h = box
    size = min(0.30, max(0.16, min(w, h) * 0.30))
    badge_x = x + w - size - 0.025
    badge_y = y + 0.025
    stream = BytesIO(_wifi_badge_png_bytes())
    slide.shapes.add_picture(
        stream,
        Inches(badge_x),
        Inches(badge_y),
        Inches(size),
        Inches(size),
    )


# =============================================================================
# POWERPOINT TEMPLATE SUPPORT
# =============================================================================

EMU_PER_INCH = 914400.0
PPT_TEMPLATE_MARKERS = (
    "{{DIAGRAM}}",
    "[[DIAGRAM]]",
    "DIAGRAM_PLACEHOLDER",
    "DIAGRAM_AREA",
)


@dataclass(frozen=True)
class _PptContentTransform:
    """Uniformly map the existing engineering drawing area into a template area."""

    origin_x: float
    origin_y: float
    scale: float

    def point(self, point):
        return (
            self.origin_x + (point[0] - CONTENT_LEFT) * self.scale,
            self.origin_y + (point[1] - CONTENT_TOP) * self.scale,
        )

    def box(self, box):
        x, y, w, h = box
        px, py = self.point((x, y))
        return (
            px,
            py,
            w * self.scale,
            h * self.scale,
        )


def _pptx_source_to_presentation(template_source):
    """Open a PowerPoint template from bytes, BytesIO, pathlib.Path, or string path."""
    if template_source is None:
        return None

    if isinstance(template_source, (bytes, bytearray)):
        return Presentation(BytesIO(bytes(template_source)))

    if isinstance(template_source, BytesIO):
        template_source.seek(0)
        return Presentation(template_source)

    if hasattr(template_source, "read") and not isinstance(template_source, (str, Path)):
        data = template_source.read()
        try:
            template_source.seek(0)
        except Exception:
            pass
        return Presentation(BytesIO(data))

    path = Path(template_source)
    if not path.exists():
        raise FileNotFoundError(f"PowerPoint template not found: {path}")
    return Presentation(str(path))


def validate_powerpoint_template(template_source) -> tuple[bool, str]:
    """Validate that a supplied file is a readable PPTX template."""
    try:
        presentation = _pptx_source_to_presentation(template_source)
        if presentation is None:
            return False, "No PowerPoint template was supplied."
        slide_count = len(presentation.slides)
        layout_count = len(presentation.slide_layouts)
        return (
            True,
            f"Valid PPTX template ({slide_count} slide{'s' if slide_count != 1 else ''}, "
            f"{layout_count} layout{'s' if layout_count != 1 else ''}).",
        )
    except Exception as exc:
        return False, f"Invalid or unreadable PPTX template: {exc}"


def _shape_text_value(shape) -> str:
    try:
        if getattr(shape, "has_text_frame", False):
            return str(shape.text or "").strip()
    except Exception:
        return ""
    return ""


def _shape_name_value(shape) -> str:
    try:
        return str(shape.name or "").strip()
    except Exception:
        return ""


def _contains_template_marker(shape) -> bool:
    text_value = _shape_text_value(shape).upper()
    name_value = _shape_name_value(shape).upper()
    return any(marker.upper() in text_value or marker.upper() in name_value for marker in PPT_TEMPLATE_MARKERS)


def _shape_bounds_inches(shape) -> tuple[float, float, float, float]:
    return (
        float(shape.left) / EMU_PER_INCH,
        float(shape.top) / EMU_PER_INCH,
        float(shape.width) / EMU_PER_INCH,
        float(shape.height) / EMU_PER_INCH,
    )


def _remove_shape(shape) -> None:
    try:
        element = shape._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
    except Exception:
        pass


def _remove_slide_at(presentation, index: int) -> None:
    """Remove one slide while preserving the presentation masters/layouts/themes."""
    slide_ids = presentation.slides._sldIdLst
    slide_id = slide_ids[index]
    try:
        presentation.part.drop_rel(slide_id.rId)
    except Exception:
        pass
    slide_ids.remove(slide_id)


def _find_marker_slide(presentation):
    for slide_index, slide in enumerate(presentation.slides):
        for shape in slide.shapes:
            if _contains_template_marker(shape):
                return slide_index, shape
    return None, None


def _layout_score(layout) -> tuple[int, int]:
    name = str(getattr(layout, "name", "") or "").lower()
    score = 0
    if "diagram" in name:
        score += 100
    if "content" in name:
        score += 70
    if "object" in name:
        score += 55
    if "blank" in name:
        score += 35
    if "title" in name and "content" not in name:
        score -= 15

    placeholder_area = 0
    try:
        for shape in layout.placeholders:
            placeholder_area += int(shape.width) * int(shape.height)
    except Exception:
        pass
    return score, placeholder_area


def _choose_template_layout(presentation):
    if len(presentation.slide_layouts) == 0:
        raise ValueError("The PowerPoint template does not contain any slide layouts.")

    # First prefer an explicitly useful content/diagram layout from the template.
    layouts = list(presentation.slide_layouts)
    best = max(layouts, key=_layout_score)

    # If the template already has a slide with a meaningful non-title layout, use
    # that layout as a stronger hint because it is the layout the template author
    # actually used.
    if len(presentation.slides) > 0:
        try:
            first_layout = presentation.slides[0].slide_layout
            first_name = str(getattr(first_layout, "name", "") or "").lower()
            if "content" in first_name or "diagram" in first_name or "blank" in first_name:
                best = first_layout
        except Exception:
            pass
    return best


def _is_title_like_placeholder(shape) -> bool:
    name = _shape_name_value(shape).lower()
    if "title" in name:
        return True
    text = _shape_text_value(shape).lower()
    return text in {"click to add title", "click to add subtitle"}


def _find_content_placeholder(slide, slide_width_in: float, slide_height_in: float):
    candidates = []
    for shape in slide.shapes:
        if not getattr(shape, "is_placeholder", False):
            continue
        if _is_title_like_placeholder(shape):
            continue

        x, y, w, h = _shape_bounds_inches(shape)
        area = w * h
        # Ignore footer/date/page-number sized placeholders.
        if w < slide_width_in * 0.28 or h < slide_height_in * 0.20:
            continue
        candidates.append((area, shape, (x, y, w, h)))

    if not candidates:
        return None, None

    _area, shape, bounds = max(candidates, key=lambda item: item[0])
    return shape, bounds


def _set_template_title_if_available(slide, title: str) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "is_placeholder", False):
            continue
        if not _is_title_like_placeholder(shape):
            continue
        try:
            frame = shape.text_frame
            frame.clear()
            frame.paragraphs[0].text = title
            return
        except Exception:
            continue


def _default_template_content_rect(slide_width_in: float, slide_height_in: float):
    # Preserve likely template header/footer regions when no explicit content
    # placeholder exists.
    left = slide_width_in * 0.055
    top = slide_height_in * 0.16
    width = slide_width_in * 0.89
    height = slide_height_in * 0.75
    return left, top, width, height


def _make_content_transform(target_rect) -> _PptContentTransform:
    target_x, target_y, target_w, target_h = target_rect
    base_w = CONTENT_RIGHT - CONTENT_LEFT
    base_h = CONTENT_BOTTOM - CONTENT_TOP

    scale = min(target_w / base_w, target_h / base_h)
    scale = max(scale, 0.05)

    mapped_w = base_w * scale
    mapped_h = base_h * scale
    origin_x = target_x + (target_w - mapped_w) / 2.0
    origin_y = target_y + (target_h - mapped_h) / 2.0
    return _PptContentTransform(origin_x, origin_y, scale)


def _prepare_template_presentation(template_source, title: str):
    """Return (presentation, slide, transform) for a user-provided PPTX template.

    Two template styles are supported:
    1. Exact slide-level template: place a text box/shape named DIAGRAM_AREA or
       containing {{DIAGRAM}} where the diagram should be inserted. That exact
       slide is retained and only the marker shape is removed.
    2. Normal PowerPoint theme/layout template: if no marker exists, a fresh slide
       is created from the template's best content layout, preserving the template
       master, theme, fonts, background and inherited formatting.
    """
    presentation = _pptx_source_to_presentation(template_source)
    if presentation is None:
        raise ValueError("No PowerPoint template was supplied.")

    slide_width_in = float(presentation.slide_width) / EMU_PER_INCH
    slide_height_in = float(presentation.slide_height) / EMU_PER_INCH

    marker_slide_index, marker_shape = _find_marker_slide(presentation)

    if marker_slide_index is not None and marker_shape is not None:
        # Keep the exact authored slide so slide-level artwork is preserved.
        for index in range(len(presentation.slides) - 1, -1, -1):
            if index != marker_slide_index:
                _remove_slide_at(presentation, index)

        slide = presentation.slides[0]
        # Re-find the marker because slide indices/elements may have shifted.
        actual_marker = None
        for shape in slide.shapes:
            if _contains_template_marker(shape):
                actual_marker = shape
                break

        if actual_marker is not None:
            target_rect = _shape_bounds_inches(actual_marker)
            _remove_shape(actual_marker)
        else:
            target_rect = _default_template_content_rect(slide_width_in, slide_height_in)

        transform = _make_content_transform(target_rect)
        return presentation, slide, transform

    # No explicit marker: preserve theme/master/layout, but avoid retaining sample
    # slide content by creating a new clean slide from the template layout.
    layout = _choose_template_layout(presentation)
    for index in range(len(presentation.slides) - 1, -1, -1):
        _remove_slide_at(presentation, index)

    slide = presentation.slides.add_slide(layout)
    _set_template_title_if_available(slide, title)

    content_placeholder, target_rect = _find_content_placeholder(
        slide,
        slide_width_in,
        slide_height_in,
    )
    if target_rect is None:
        target_rect = _default_template_content_rect(slide_width_in, slide_height_in)
    elif content_placeholder is not None:
        try:
            # Keep placeholder geometry and styling inherited from the template but
            # remove prompt/sample text so the generated diagram occupies the area.
            content_placeholder.text_frame.clear()
        except Exception:
            pass

    transform = _make_content_transform(target_rect)
    return presentation, slide, transform


# =============================================================================
# FINAL AUTOMATIC GEOMETRY CACHE
# =============================================================================


def _diagram_geometry_signature(diagram: DiagramSpec) -> str:
    """Stable, exact signature for geometry-affecting diagram state."""
    try:
        return diagram.model_dump_json()
    except AttributeError:
        return diagram.json()


@lru_cache(maxsize=48)
def _cached_final_auto_geometry(
    diagram_json: str,
    canvas_key: tuple[float, float, float, float, float, float],
):
    """Compute layout + routes + labels once for an unchanged finalized diagram.

    This does not alter any routing rule.  It only reuses the exact deterministic
    result when PNG/PPT/editor rendering asks for the same geometry again.
    """
    del canvas_key  # included only to invalidate when the dynamic canvas changes
    try:
        cached_diagram = DiagramSpec.model_validate_json(diagram_json)
    except AttributeError:
        cached_diagram = DiagramSpec.parse_raw(diagram_json)

    boxes = compute_layout_boxes(cached_diagram)
    routes = build_edge_routes(cached_diagram, boxes)
    routes = _visually_separate_shared_connection_segments(routes)

    if "Component Selection mode" in (getattr(cached_diagram, "style_notes", []) or []):
        labels = [None] * len(cached_diagram.edges)
    else:
        labels = place_edge_labels(cached_diagram, routes, boxes)

    return _center_final_layout_horizontally(boxes, routes, labels)


def _final_auto_geometry(diagram: DiagramSpec):
    canvas_key = (
        float(SLIDE_WIDTH_IN),
        float(SLIDE_HEIGHT_IN),
        float(CONTENT_LEFT),
        float(CONTENT_RIGHT),
        float(CONTENT_TOP),
        float(CONTENT_BOTTOM),
    )
    boxes, routes, labels = _cached_final_auto_geometry(
        _diagram_geometry_signature(diagram),
        canvas_key,
    )
    # Callers receive lightweight containers of the cached immutable tuples; the
    # cached routing result itself is never recalculated or modified.
    return dict(boxes), list(routes), list(labels)


# =============================================================================
# POWERPOINT
# =============================================================================

def _create_powerpoint_impl(
    diagram: DiagramSpec,
    assets_dir: Path,
    template_source=None,
) -> bytes:
    """Create the downloadable PowerPoint.

    When ``template_source`` is supplied, the template is used as the PowerPoint
    base. Its masters, layouts, theme, background and slide-level artwork are
    preserved wherever applicable. The engineering diagram is mapped into the
    template's diagram/content area while keeping the existing editable PowerPoint
    shapes used by this project.

    When no template is supplied, the original Agentic Sketch-to-PPT PowerPoint
    output is produced exactly as before.
    """
    ensure_component_assets(assets_dir)

    template_active = template_source is not None
    transform = None

    if template_active:
        presentation, slide, transform = _prepare_template_presentation(
            template_source,
            get_display_title(diagram),
        )
    else:
        presentation = Presentation()
        presentation.slide_width = Inches(SLIDE_WIDTH_IN)
        presentation.slide_height = Inches(SLIDE_HEIGHT_IN)

        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*WHITE)

        add_reference_header_ppt(slide, get_display_title(diagram))

    boxes, routes, labels = _final_auto_geometry(diagram)

    # Apply one uniform coordinate transform only for template-based PPT output.
    # PNG/PDF rendering and every diagram-layout algorithm remain untouched.
    if transform is not None:
        ppt_boxes = {
            node_id: transform.box(box)
            for node_id, box in boxes.items()
        }
        ppt_routes = []
        for route in routes:
            if route is None:
                ppt_routes.append(None)
                continue
            points, direction = route
            ppt_routes.append(
                ([transform.point(point) for point in points], direction)
            )

        ppt_labels = []
        for placement in labels:
            if placement is None:
                ppt_labels.append(None)
                continue
            text_value, rect, leader_anchor = placement
            ppt_labels.append(
                (
                    text_value,
                    transform.box(rect),
                    transform.point(leader_anchor) if leader_anchor is not None else None,
                )
            )
    else:
        ppt_boxes = boxes
        ppt_routes = routes
        ppt_labels = labels

    # Pipes first. Route geometry is unchanged; only visual stroke style comes
    # from the edge topology metadata.
    for edge, route in zip(diagram.edges, ppt_routes):
        if route is None:
            continue

        line_color, dotted = _edge_line_style(edge)
        points, _direction = route
        for start_point, end_point in zip(points, points[1:]):
            add_ppt_segment(
                slide,
                start_point,
                end_point,
                color=line_color,
                dotted=dotted,
            )

    # Components.
    for node in diagram.nodes:
        box = ppt_boxes.get(node.id)
        if box is not None:
            add_ppt_component(slide, node, box, assets_dir)

    # Arrowheads.
    for edge, route in zip(diagram.edges, ppt_routes):
        if route is None:
            continue
        if edge.direction == "unknown":
            continue

        points, direction = route
        if edge.direction == "target_to_source":
            reverse_direction = _final_direction(list(reversed(points)))
            # Keep the arrowhead on the final visible segment, just before the
            # endpoint/component boundary.  This matches the PNG renderer and
            # prevents the triangle from being clipped or appearing rotated
            # around the component connection point.
            if len(points) >= 2:
                add_ppt_arrow(slide, points[1], reverse_direction)
        else:
            # Use the last segment's inner point instead of the exact endpoint.
            # The endpoint remains the connection anchor; the arrow is a visual
            # marker placed on the segment, matching the reference image.
            if len(points) >= 2:
                add_ppt_arrow(slide, points[-2], direction)

    # Wireless destination markers. These are visual overlays only and do not
    # affect routing, spacing, topology, arrows, line colors or component layout.
    for node_id in sorted(_wireless_destination_node_ids(diagram)):
        box = ppt_boxes.get(node_id)
        if box is not None:
            add_ppt_wifi_badge(slide, box)

    # Labels last.
    for placement in ppt_labels:
        if placement is None:
            continue
        text_value, rect, leader_anchor = placement
        add_ppt_edge_label(slide, text_value, rect, leader_anchor)

    # Preserve the original Agentic Sketch-to-PPT footer only for the original
    # non-template output. A user template owns its own footer/background design.
    if not template_active:
        footer = slide.shapes.add_textbox(
            Inches(11.25),
            Inches(6.91),
            Inches(1.35),
            Inches(0.20),
        )

        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = "iTank"
        paragraph.alignment = PP_ALIGN.RIGHT
        paragraph.font.name = "Aptos Display"
        paragraph.font.size = Pt(11)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(38, 122, 196)

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


# =============================================================================
# PNG HELPERS
# =============================================================================

def ppt_box_to_pixels(
    box,
    width: int,
    height: int,
):
    x, y, w, h = box

    return (
        int(
            x / SLIDE_WIDTH_IN
            * width
        ),
        int(
            y / SLIDE_HEIGHT_IN
            * height
        ),
        int(
            (x + w)
            / SLIDE_WIDTH_IN
            * width
        ),
        int(
            (y + h)
            / SLIDE_HEIGHT_IN
            * height
        ),
    )


def ppt_point_to_pixels(
    point,
    width: int,
    height: int,
):
    return (
        int(
            point[0]
            / SLIDE_WIDTH_IN
            * width
        ),
        int(
            point[1]
            / SLIDE_HEIGHT_IN
            * height
        ),
    )


def draw_png_header(
    draw,
    width: int,
    height: int,
    title: str,
) -> None:
    # Premium reference-style header and white content panel.
    draw.rectangle((0, 0, width, 170), fill=(7, 60, 120))

    draw.text(
        (86, 56),
        "Agentic Sketch-to-PPT",
        anchor="lm",
        fill=WHITE,
        font=load_font(42, True),
    )
    draw.text(
        (86, 112),
        "Professional Water Automation Diagram",
        anchor="lm",
        fill=(220, 236, 255),
        font=load_font(24, False),
    )

    # Success pill.
    pill_w = 330
    _safe_rounded_rectangle(
        draw,
        (width - pill_w - 70, 48, width - 70, 116),
        radius=30,
        fill=(229, 249, 238),
        outline=(176, 232, 198),
        width=1,
    )
    draw.text(
        (width - pill_w // 2 - 70, 82),
        "✓  Diagram Ready",
        anchor="mm",
        fill=(19, 138, 75),
        font=load_font(24, True),
    )

    # Main content card.
    _safe_rounded_rectangle(
        draw,
        (46, 195, width - 46, height - 44),
        radius=30,
        fill=WHITE,
        outline=BORDER_GRAY,
        width=3,
    )

    title_lines = split_title_lines(title, 70)
    title_font = load_font(36 if len(title_lines) == 1 else 31, True)
    y = 242
    for idx, line in enumerate(title_lines):
        draw.text(
            (width // 2, y + idx * 42),
            line,
            anchor="mm",
            fill=TITLE_BLUE,
            font=title_font,
        )

    rule_y = 294 if len(title_lines) == 1 else 332
    draw.line((120, rule_y, width - 120, rule_y), fill=(215, 226, 240), width=3)


def _draw_png_dotted_segment(
    draw,
    start,
    end,
    *,
    color=NON_WATER_COLOR,
    dot_radius: int = 6,
    gap: int = 22,
) -> None:
    """Draw a true dotted segment with round dots and stable endpoint coverage."""
    x1, y1 = float(start[0]), float(start[1])
    x2, y2 = float(end[0]), float(end[1])
    dx = x2 - x1
    dy = y2 - y1
    length = math.hypot(dx, dy)
    if length <= 1e-6:
        return

    step = max(float(gap), float(dot_radius * 2 + 4))
    count = max(1, int(length // step))
    for index in range(count + 1):
        distance = min(length, index * step)
        ratio = distance / length
        cx = x1 + dx * ratio
        cy = y1 + dy * ratio
        draw.ellipse(
            (
                int(round(cx - dot_radius)),
                int(round(cy - dot_radius)),
                int(round(cx + dot_radius)),
                int(round(cy + dot_radius)),
            ),
            fill=color,
        )

    # Ensure the segment endpoint is visibly represented even when it falls
    # between the normal dot spacing interval.
    draw.ellipse(
        (
            int(round(x2 - dot_radius)),
            int(round(y2 - dot_radius)),
            int(round(x2 + dot_radius)),
            int(round(y2 + dot_radius)),
        ),
        fill=color,
    )


def draw_png_arrow(
    draw,
    start,
    end,
) -> None:
    """Draw a larger orange direction marker with a white contrast outline."""
    _x1, _y1 = start
    x2, y2 = end

    # IMPORTANT: keep the arrowhead visually upright.
    # Do not derive its rotation from the route angle.  The routing path
    # remains unchanged; only the arrowhead orientation is fixed.
    size = 52
    spread = 0.60

    # Pillow coordinates: Y increases downward, so these points create an
    # upright triangle whose tip is at (x2, y2).
    point_a = (
        x2 - size * math.sin(spread),
        y2 + size * math.cos(spread),
    )
    point_b = (
        x2 + size * math.sin(spread),
        y2 + size * math.cos(spread),
    )

    polygon = [(x2, y2), point_a, point_b]
    draw.polygon(
        polygon,
        fill=ARROW_COLOR,
        outline=ARROW_OUTLINE,
        width=1,
    )


def draw_png_wifi_badge(draw, box) -> None:
    """Draw one Wi-Fi badge at the exact top-right corner of a component card."""
    left, top, right, bottom = _normalize_rect(box)
    width = max(1, right - left)
    height = max(1, bottom - top)
    size = max(44, min(104, int(min(width, height) * 0.30)))
    pad = max(6, size // 12)

    badge_left = right - size - pad
    badge_top = top + pad
    badge_right = badge_left + size
    badge_bottom = badge_top + size

    border_width = max(2, size // 28)
    draw.ellipse(
        (badge_left, badge_top, badge_right, badge_bottom),
        fill=WIFI_SYMBOL_BG,
        outline=WIFI_SYMBOL_BORDER,
        width=border_width,
    )

    cx = (badge_left + badge_right) // 2
    base_y = badge_top + int(size * 0.76)
    stroke = max(4, size // 14)
    for radius in (int(size * 0.36), int(size * 0.26), int(size * 0.17)):
        draw.arc(
            (cx - radius, base_y - radius, cx + radius, base_y + radius),
            start=205,
            end=335,
            fill=WIFI_SYMBOL_COLOR,
            width=stroke,
        )

    dot_radius = max(4, size // 20)
    draw.ellipse(
        (
            cx - dot_radius,
            base_y - dot_radius,
            cx + dot_radius,
            base_y + dot_radius,
        ),
        fill=WIFI_SYMBOL_COLOR,
    )


def draw_png_component(
    canvas: Image.Image,
    draw,
    node: DiagramNode,
    box,
    assets_dir: Path,
) -> None:
    left, top, right, bottom = _normalize_rect(box)
    width = max(1, right - left)
    height = max(1, bottom - top)

    if node.node_type == "junction":
        radius = max(6, min(width, height) // 3)
        cx = (left + right) // 2
        cy = (top + bottom) // 2
        draw.ellipse((cx - radius, cy - radius, cx + radius, cy + radius), fill=PIPE_DARK)
        return

    # Premium card shadow.
    shadow = 12
    _safe_rounded_rectangle(
        draw,
        (left + shadow, top + shadow, right + shadow, bottom + shadow),
        radius=20,
        fill=(226, 235, 246),
    )
    _safe_rounded_rectangle(
        draw,
        (left, top, right, bottom),
        radius=20,
        fill=WHITE,
        outline=CARD_BORDER,
        width=3,
    )

    caption_h = max(92, min(142, int(height * 0.31)))
    image_bottom = max(top + 20, bottom - caption_h)
    image_height = max(1, image_bottom - top)

    path = node_component_asset_path(assets_dir, node)
    if path is not None and width > 16 and image_height > 16:
        target_w = max(1, width - 32)
        target_h = max(1, image_height - 30)
        fitted = _contained_asset_for_path(path, target_w, target_h)
        paste_x = left + (width - fitted.width) // 2
        paste_y = top + 12 + max(0, (image_height - fitted.height - 12) // 2)
        canvas.paste(fitted, (paste_x, paste_y))
    elif width > 16 and image_height > 16:
        badge_w = max(92, min(int(width * 0.55), 230))
        badge_h = max(70, min(int(image_height * 0.55), 150))
        badge_left = left + (width - badge_w) // 2
        badge_top = top + max(12, (image_height - badge_h) // 2)
        _safe_rounded_rectangle(
            draw,
            (badge_left, badge_top, badge_left + badge_w, badge_top + badge_h),
            radius=26,
            fill=(233, 244, 255),
            outline=LABEL_BLUE,
            width=1,
        )
        code = _component_code(node)
        code_size = max(28, min(64, int(badge_h * 0.38)))
        code_font = load_font(code_size, True)
        while code_size > 24:
            bbox = draw.textbbox((0, 0), code, font=code_font)
            if bbox[2] - bbox[0] <= badge_w - 28:
                break
            code_size -= 2
            code_font = load_font(code_size, True)
        draw.text(
            (badge_left + badge_w // 2, badge_top + badge_h // 2),
            code,
            anchor="mm",
            fill=TITLE_BLUE,
            font=code_font,
        )

    caption_top = image_bottom
    _safe_rounded_rectangle(
        draw,
        (left + 8, caption_top - 6, right - 8, bottom - 8),
        radius=18,
        fill=LABEL_BLUE,
    )

    label = component_display_label(node)
    lines = wrap_component_label(label, 20)[:2]
    detail_lines = _node_details(node)[:2]

    # Large white title and compact metadata, like the supplied reference UI.
    title_size = 54 if len(lines) == 1 else 46
    title_font = load_font(title_size, True)
    while title_size > 28:
        boxes = [draw.textbbox((0, 0), line, font=title_font) for line in lines]
        if max((b[2] - b[0] for b in boxes), default=0) <= width - 42:
            break
        title_size -= 2
        title_font = load_font(title_size, True)

    cx = (left + right) // 2
    if len(lines) == 1:
        title_y = caption_top + 36
        draw.text((cx, title_y), lines[0], anchor="mm", fill=WHITE, font=title_font)
        meta_y = title_y + 44
    else:
        draw.text((cx, caption_top + 27), lines[0], anchor="mm", fill=WHITE, font=title_font)
        draw.text((cx, caption_top + 68), lines[1], anchor="mm", fill=WHITE, font=title_font)
        meta_y = caption_top + 105

    if detail_lines:
        meta_font = load_font(28, False)
        meta_text = "  •  ".join(detail_lines)
        if len(meta_text) > 54:
            meta_text = meta_text[:51] + "..."
        draw.text((cx, min(bottom - 28, meta_y)), meta_text, anchor="mm", fill=(228, 239, 255), font=meta_font)


def draw_png_edge_label(
    draw,
    text: str,
    rect,
    canvas_width: int,
    canvas_height: int,
    leader_anchor=None,
) -> None:
    x, y, w, h = rect
    left = int(x / SLIDE_WIDTH_IN * canvas_width)
    top = int(y / SLIDE_HEIGHT_IN * canvas_height)
    right = int((x + w) / SLIDE_WIDTH_IN * canvas_width)
    bottom = int((y + h) / SLIDE_HEIGHT_IN * canvas_height)

    if leader_anchor is not None:
        anchor_px = ppt_point_to_pixels(leader_anchor, canvas_width, canvas_height)
        center_px = ((left + right) // 2, (top + bottom) // 2)
        if math.hypot(center_px[0] - anchor_px[0], center_px[1] - anchor_px[1]) > 55:
            draw.line((anchor_px[0], anchor_px[1], center_px[0], center_px[1]), fill=(122, 158, 202), width=3)

    _safe_rounded_rectangle(
        draw,
        (left, top, right, bottom),
        radius=14,
        fill=(243, 248, 255),
        outline=(174, 205, 242),
        width=3,
    )

    lines = textwrap.wrap(text, width=26, break_long_words=False, break_on_hyphens=False) or [text]
    lines = lines[:2]
    font = load_font(50 if len(lines) == 1 else 42, True)
    cx = (left + right) // 2
    cy = (top + bottom) // 2
    if len(lines) == 1:
        draw.text((cx, cy), lines[0], anchor="mm", fill=TITLE_BLUE, font=font)
    else:
        draw.text((cx, cy - 20), lines[0], anchor="mm", fill=TITLE_BLUE, font=font)
        draw.text((cx, cy + 20), lines[1], anchor="mm", fill=TITLE_BLUE, font=font)


def create_powerpoint(
    diagram: DiagramSpec,
    assets_dir: Path,
    template_source=None,
) -> bytes:
    """Canvas-aware wrapper around the existing PowerPoint renderer."""
    with _canvas_scope(diagram):
        return _create_powerpoint_impl(
            diagram=diagram,
            assets_dir=assets_dir,
            template_source=template_source,
        )


# =============================================================================
# PNG PREVIEW
# =============================================================================

def _create_preview_png_impl(
    diagram: DiagramSpec,
    assets_dir: Path,
    manual_route_overrides: dict | None = None,
    manual_component_overrides: dict | None = None,
    suppress_connection_lines: bool = False,
    suppress_components: bool = False,
    suppress_edge_labels: bool = False,
) -> bytes:
    ensure_component_assets(
        assets_dir
    )

    # Pixel dimensions follow the logical canvas.  Large diagrams receive more
    # actual pixels, while a max dimension keeps memory use bounded.
    pixels_per_inch = min(
        420.0,
        9600.0 / max(SLIDE_WIDTH_IN, 1.0),
        7600.0 / max(SLIDE_HEIGHT_IN, 1.0),
    )
    pixels_per_inch = max(180.0, pixels_per_inch)
    width = max(1600, int(round(SLIDE_WIDTH_IN * pixels_per_inch)))
    height = max(900, int(round(SLIDE_HEIGHT_IN * pixels_per_inch)))

    canvas = Image.new(
        "RGB",
        (
            width,
            height,
        ),
        (247, 251, 255),
    )

    draw = ImageDraw.Draw(
        canvas
    )

    draw_png_header(
        draw,
        width,
        height,
        get_display_title(
            diagram
        ),
    )

    boxes, routes, labels = _final_auto_geometry(diagram)

    boxes, hidden_component_ids = _apply_manual_component_overrides(
        boxes, manual_component_overrides
    )

    # Manual connection editing is a display-only override applied after the
    # normal router has finished.  With no override, output is byte-for-byte
    # governed by the existing routing path exactly as before.
    routes = _apply_manual_route_overrides(
        diagram,
        routes,
        manual_route_overrides,
    )

    pixel_boxes = {
        node_id: ppt_box_to_pixels(
            box,
            width,
            height,
        )
        for node_id, box
        in boxes.items()
    }

    pixel_routes = []
    pixel_route_styles = []

    # -------------------------------------------------------------------------
    # 1. PIPES
    # -------------------------------------------------------------------------

    # Convert every routed edge first.  All white clearance underlays are drawn
    # in one batch before any coloured stroke.  This is important for one-to-one
    # connections: an underlay from a later edge must never erase a previously
    # drawn parallel connection and make two logical edges look like one line.
    for edge, route in zip(diagram.edges, routes):
        if route is None:
            pixel_routes.append(None)
            pixel_route_styles.append(None)
            continue

        points, direction = route
        line_color, dotted = _edge_line_style(edge)
        pixel_points = [
            ppt_point_to_pixels(point, width, height)
            for point in points
        ]
        pixel_routes.append((pixel_points, direction))
        pixel_route_styles.append((line_color, dotted))

    if not suppress_connection_lines:
        # Pass 1: clear background beneath every connection.
        for route in pixel_routes:
            if route is None:
                continue
            pixel_points, _direction = route
            for start_point, end_point in zip(pixel_points, pixel_points[1:]):
                draw.line(
                    (
                        start_point[0],
                        start_point[1],
                        end_point[0],
                        end_point[1],
                    ),
                    fill=WHITE,
                    width=2,
                )

        # Pass 2: draw every actual connection.  Because all underlays are
        # already complete, adjacent/parallel one-to-one lines remain visible.
        for route, style in zip(pixel_routes, pixel_route_styles):
            if route is None or style is None:
                continue
            pixel_points, _direction = route
            line_color, dotted = style

            if dotted:
                for start_point, end_point in zip(pixel_points, pixel_points[1:]):
                    _draw_png_dotted_segment(
                        draw,
                        start_point,
                        end_point,
                        color=line_color,
                        dot_radius=2,
                        gap=24,
                    )
            else:
                for start_point, end_point in zip(pixel_points, pixel_points[1:]):
                    draw.line(
                        (
                            start_point[0],
                            start_point[1],
                            end_point[0],
                            end_point[1],
                        ),
                        fill=WATER_FLOW_DARK,
                        width=2,
                    )

                for start_point, end_point in zip(pixel_points, pixel_points[1:]):
                    draw.line(
                        (
                            start_point[0],
                            start_point[1],
                            end_point[0],
                            end_point[1],
                        ),
                        fill=line_color,
                        width=2,
                    )

    # Manually created visual connections are drawn with the exact same line
    # widths and arrow treatment as normal solid/dotted connections.
    custom_pixel_routes = []
    if not suppress_connection_lines:
        for custom_id, raw in dict(manual_route_overrides or {}).items():
            if not isinstance(raw, dict) or not bool(raw.get("custom", False)) or bool(raw.get("hidden", False)):
                continue
            clean = []
            for value in list(raw.get("points", []) or []):
                try:
                    x, y = value
                    clean.append((float(x), float(y)))
                except Exception:
                    continue
            if len(clean) < 2:
                continue
            pixel_points = [ppt_point_to_pixels(point, width, height) for point in clean]
            color_hex = str(raw.get("color", "#1473E6") or "#1473E6").lstrip("#")
            try:
                if len(color_hex) == 6:
                    line_color = tuple(int(color_hex[i:i+2], 16) for i in (0, 2, 4))
                else:
                    line_color = WATER_FLOW_DARK
            except Exception:
                line_color = WATER_FLOW_DARK
            dotted = bool(raw.get("dotted", False))
            for a, b in zip(pixel_points, pixel_points[1:]):
                draw.line((a[0], a[1], b[0], b[1]), fill=WHITE, width=5)
            if dotted:
                for a, b in zip(pixel_points, pixel_points[1:]):
                    _draw_png_dotted_segment(draw, a, b, color=line_color, dot_radius=2, gap=20)
            else:
                for a, b in zip(pixel_points, pixel_points[1:]):
                    draw.line((a[0], a[1], b[0], b[1]), fill=WATER_FLOW_DARK, width=3)
                for a, b in zip(pixel_points, pixel_points[1:]):
                    draw.line((a[0], a[1], b[0], b[1]), fill=line_color, width=2)
            custom_pixel_routes.append((pixel_points, str(raw.get("direction", "source_to_target") or "source_to_target")))

    # -------------------------------------------------------------------------
    # 2. COMPONENTS
    # -------------------------------------------------------------------------

    if not suppress_components:
        for node in diagram.nodes:
            if str(node.id) in hidden_component_ids:
                continue
            box = pixel_boxes.get(
                node.id
            )

            if box is None:
                continue

            draw_png_component(
                canvas,
                draw,
                node,
                box,
                assets_dir,
            )

    # -------------------------------------------------------------------------
    # 3. ARROWS
    # -------------------------------------------------------------------------

    if not suppress_connection_lines:
        for edge, route in zip(
            diagram.edges,
            pixel_routes,
        ):
            effective_direction = _effective_edge_direction(edge, manual_route_overrides)
            if route is None or effective_direction == "unknown":
                continue

            points, _direction = route

            if len(points) >= 2:
                if effective_direction == "target_to_source":
                    draw_png_arrow(draw, points[1], points[0])
                else:
                    draw_png_arrow(draw, points[-2], points[-1])

        for points, effective_direction in custom_pixel_routes:
            if effective_direction == "unknown" or len(points) < 2:
                continue
            if effective_direction == "target_to_source":
                draw_png_arrow(draw, points[1], points[0])
            else:
                draw_png_arrow(draw, points[-2], points[-1])

    # -------------------------------------------------------------------------
    # 4. WIRELESS DESTINATION BADGES
    # -------------------------------------------------------------------------

    if not suppress_components:
        for node_id in sorted(_wireless_destination_node_ids(diagram)):
            if str(node_id) in hidden_component_ids:
                continue
            box = pixel_boxes.get(node_id)
            if box is not None:
                draw_png_wifi_badge(draw, box)

    # -------------------------------------------------------------------------
    # 5. LABELS
    # -------------------------------------------------------------------------

    if not suppress_edge_labels:
        for placement in labels:
            if placement is None:
                continue

            text, rect, leader_anchor = placement

            draw_png_edge_label(
                draw,
                text,
                rect,
                width,
                height,
                leader_anchor,
            )

    draw.text(
        (
            width - 70,
            height - 37,
        ),
        "iTank",
        anchor="rm",
        fill=(
            36,
            122,
            196,
        ),
        font=load_font(
            26,
            True,
        ),
    )

    output = BytesIO()

    # PNG optimization performs an expensive second compression search.  It does
    # not change a single rendered pixel, so disable it for much faster generation.
    canvas.save(
        output,
        format="PNG",
        optimize=False,
        compress_level=3,
    )

    return output.getvalue()



def _manual_component_config(manual_component_overrides, node_id: str) -> dict:
    if not manual_component_overrides:
        return {}
    value = dict(manual_component_overrides or {}).get(str(node_id), {})
    return dict(value) if isinstance(value, dict) else {}


def _apply_manual_component_overrides(boxes, manual_component_overrides=None):
    """Apply visual-only component box overrides after automatic layout."""
    result = dict(boxes or {})
    hidden = set()
    if not manual_component_overrides:
        return result, hidden

    for node_id, raw in dict(manual_component_overrides or {}).items():
        if str(node_id) not in result or not isinstance(raw, dict):
            continue
        if bool(raw.get("hidden", False)):
            hidden.add(str(node_id))
            continue
        raw_box = raw.get("box")
        if not isinstance(raw_box, (list, tuple)) or len(raw_box) != 4:
            continue
        try:
            x, y, w, h = [float(v) for v in raw_box]
        except Exception:
            continue
        if w <= 0 or h <= 0:
            continue
        result[str(node_id)] = (x, y, w, h)
    return result, hidden



def _manual_route_config(manual_route_overrides, edge_id: str) -> dict:
    if not manual_route_overrides:
        return {}
    value = dict(manual_route_overrides or {}).get(str(edge_id), {})
    if isinstance(value, dict):
        return dict(value)
    if isinstance(value, (list, tuple)):
        return {"points": value}
    return {}


def _apply_manual_route_overrides(diagram: DiagramSpec, routes, manual_route_overrides=None):
    """Apply visual-only user route overrides after the normal router finishes.

    The logical edge and component positions remain untouched. Manual editing may
    move the visual endpoint connection points along the existing component boundary,
    while the underlying source/target relationship is preserved. A manual delete is
    represented by ``hidden=True`` and only suppresses drawing of that edge.
    """
    if not manual_route_overrides:
        return routes

    result = list(routes)
    edge_index_by_id = {
        str(getattr(edge, "id", "")): index
        for index, edge in enumerate(getattr(diagram, "edges", []) or [])
    }

    for edge_id, raw_config in dict(manual_route_overrides or {}).items():
        index = edge_index_by_id.get(str(edge_id))
        if index is None or index >= len(result):
            continue
        original = result[index]
        if original is None:
            continue

        config = raw_config if isinstance(raw_config, dict) else {"points": raw_config}
        if bool(config.get("hidden", False)):
            result[index] = None
            continue

        original_points, direction = original
        if len(original_points) < 2:
            continue

        raw_points = config.get("points", []) or []
        clean = []
        for value in raw_points:
            try:
                x, y = value
                clean.append((float(x), float(y)))
            except Exception:
                continue

        if len(clean) < 2:
            continue

        # Preserve the logical edge while honoring the user's visual endpoint and
        # waypoint edits. Component positions are never changed here.
        result[index] = (_compress_polyline(clean), direction)

    return result


def _effective_edge_direction(edge, manual_route_overrides=None) -> str:
    config = _manual_route_config(manual_route_overrides, str(getattr(edge, "id", "")))
    direction = str(config.get("direction", "") or "").strip()
    if direction in {"source_to_target", "target_to_source", "unknown"}:
        return direction
    return str(getattr(edge, "direction", "source_to_target") or "source_to_target")


def get_preview_route_geometry(
    diagram: DiagramSpec,
    manual_route_overrides: dict | None = None,
    manual_component_overrides: dict | None = None,
) -> dict:
    """Return exact final route geometry plus immutable router originals for editing."""
    with _canvas_scope(diagram):
        boxes, original_routes, labels = _final_auto_geometry(diagram)
        boxes, hidden_component_ids = _apply_manual_component_overrides(
            boxes, manual_component_overrides
        )
        edited_routes = _apply_manual_route_overrides(
            diagram, original_routes, manual_route_overrides
        )

        route_map = {}
        for index, edge in enumerate(diagram.edges):
            original = original_routes[index] if index < len(original_routes) else None
            if original is None:
                continue
            original_points, original_router_direction = original
            edited = edited_routes[index] if index < len(edited_routes) else None
            config = _manual_route_config(manual_route_overrides, str(edge.id))
            hidden = bool(config.get("hidden", False))

            if edited is None:
                points = original_points
                router_direction = original_router_direction
            else:
                points, router_direction = edited

            line_color, dotted = _edge_line_style(edge)
            source_id = str(config.get("source", edge.source) or edge.source)
            target_id = str(config.get("target", edge.target) or edge.target)
            route_map[str(edge.id)] = {
                "points": [[float(x), float(y)] for x, y in points],
                "original_points": [[float(x), float(y)] for x, y in original_points],
                "direction": router_direction,
                "source": source_id,
                "target": target_id,
                "color": "#%02X%02X%02X" % tuple(int(v) for v in line_color),
                "dotted": bool(dotted),
                "hidden": hidden,
                "custom": False,
                "edge_direction": _effective_edge_direction(edge, manual_route_overrides),
                "original_edge_direction": str(
                    getattr(edge, "direction", "source_to_target") or "source_to_target"
                ),
            }

        # Visual-only manually created connections. They never alter DiagramSpec
        # or the automatic router; they are persisted only in the editor override map.
        for custom_id, raw in dict(manual_route_overrides or {}).items():
            if not isinstance(raw, dict) or not bool(raw.get("custom", False)):
                continue
            pts = []
            for value in list(raw.get("points", []) or []):
                try:
                    x, y = value
                    pts.append([float(x), float(y)])
                except Exception:
                    continue
            if len(pts) < 2:
                continue
            route_map[str(custom_id)] = {
                "points": pts,
                "original_points": pts,
                "direction": "manual",
                "source": str(raw.get("source", "") or ""),
                "target": str(raw.get("target", "") or ""),
                "color": str(raw.get("color", "#1473E6") or "#1473E6"),
                "dotted": bool(raw.get("dotted", False)),
                "hidden": bool(raw.get("hidden", False)),
                "custom": True,
                "edge_direction": str(raw.get("direction", "source_to_target") or "source_to_target"),
                "original_edge_direction": str(raw.get("direction", "source_to_target") or "source_to_target"),
            }

        box_map = {
            str(node_id): [float(box[0]), float(box[1]), float(box[2]), float(box[3])]
            for node_id, box in boxes.items()
        }

        return {
            "width": float(SLIDE_WIDTH_IN),
            "height": float(SLIDE_HEIGHT_IN),
            "routes": route_map,
            "boxes": box_map,
            "hidden_components": sorted(hidden_component_ids),
        }


@lru_cache(maxsize=64)
def _cached_editor_base_png(
    slide_width_in: float,
    slide_height_in: float,
    display_title: str,
) -> bytes:
    """Render the component-free Worksheet background once per canvas/title.

    This is pixel-equivalent to ``_create_preview_png_impl`` when connections,
    components and edge labels are all suppressed. The old path still computed
    full layout/routing and converted every route to pixels even though none of
    those values were drawn. Keeping this specialized background cache separate
    leaves every routing, arrow, component and full-preview code path untouched.
    """
    pixels_per_inch = min(
        420.0,
        9600.0 / max(float(slide_width_in), 1.0),
        7600.0 / max(float(slide_height_in), 1.0),
    )
    pixels_per_inch = max(180.0, pixels_per_inch)
    width = max(1600, int(round(float(slide_width_in) * pixels_per_inch)))
    height = max(900, int(round(float(slide_height_in) * pixels_per_inch)))

    canvas = Image.new(
        "RGB",
        (width, height),
        (247, 251, 255),
    )
    draw = ImageDraw.Draw(canvas)
    draw_png_header(draw, width, height, display_title)
    draw.text(
        (width - 70, height - 37),
        "iTank",
        anchor="rm",
        fill=(36, 122, 196),
        font=load_font(26, True),
    )

    output = BytesIO()
    canvas.save(
        output,
        format="PNG",
        optimize=False,
        compress_level=3,
    )
    return output.getvalue()


def create_editor_base_png(
    diagram: DiagramSpec,
    assets_dir: Path,
    manual_route_overrides: dict | None = None,
    manual_component_overrides: dict | None = None,
) -> bytes:
    """Render the exact Worksheet background without duplicate geometry work."""
    # Keep the public signature and caller behavior unchanged. These parameters
    # never affect the component/connection-free background pixels.
    del assets_dir, manual_route_overrides, manual_component_overrides
    with _canvas_scope(diagram):
        return _cached_editor_base_png(
            float(SLIDE_WIDTH_IN),
            float(SLIDE_HEIGHT_IN),
            str(get_display_title(diagram)),
        )


def create_editor_component_sprites(
    diagram: DiagramSpec,
    assets_dir: Path,
    manual_route_overrides: dict | None = None,
    manual_component_overrides: dict | None = None,
) -> list[dict]:
    """Return raster component cards and their logical boxes for the direct editor."""
    with _canvas_scope(diagram):
        geometry = get_preview_route_geometry(
            diagram,
            manual_route_overrides=manual_route_overrides,
            manual_component_overrides=manual_component_overrides,
        )
        source_bytes = _create_preview_png_impl(
            diagram=diagram,
            assets_dir=assets_dir,
            manual_route_overrides=manual_route_overrides,
            manual_component_overrides=manual_component_overrides,
            suppress_connection_lines=True,
            suppress_edge_labels=True,
        )
        image = Image.open(BytesIO(source_bytes)).convert("RGBA")
        iw, ih = image.size
        cw = max(float(geometry.get("width", 1.0)), 1e-9)
        ch = max(float(geometry.get("height", 1.0)), 1e-9)
        hidden = set(geometry.get("hidden_components", []) or [])
        sprites = []
        node_by_id = {str(n.id): n for n in (getattr(diagram, "nodes", []) or [])}
        for node_id, raw_box in dict(geometry.get("boxes", {}) or {}).items():
            if node_id in hidden or node_id not in node_by_id:
                continue
            node = node_by_id[node_id]
            if getattr(node, "node_type", "") == "junction":
                continue
            x, y, w, h = [float(v) for v in raw_box]
            left = max(0, min(iw, int(round(x / cw * iw))))
            top = max(0, min(ih, int(round(y / ch * ih))))
            right = max(left + 1, min(iw, int(round((x + w) / cw * iw))))
            bottom = max(top + 1, min(ih, int(round((y + h) / ch * ih))))
            crop = image.crop((left, top, right, bottom))
            out = BytesIO()
            crop.save(out, format="PNG")
            import base64 as _base64
            sprites.append({
                "instance_id": node_id,
                "title": str(getattr(node, "label", node_id) or node_id),
                "box": [x, y, w, h],
                "image_b64": _base64.b64encode(out.getvalue()).decode("ascii"),
            })
        return sprites


def create_preview_png(
    diagram: DiagramSpec,
    assets_dir: Path,
    manual_route_overrides: dict | None = None,
    manual_component_overrides: dict | None = None,
) -> bytes:
    """Canvas-aware wrapper around the existing PNG renderer."""
    with _canvas_scope(diagram):
        return _create_preview_png_impl(
            diagram=diagram,
            assets_dir=assets_dir,
            manual_route_overrides=manual_route_overrides,
            manual_component_overrides=manual_component_overrides,
        )


# =============================================================================
# PDF
# =============================================================================

def png_to_pdf(
    png_bytes: bytes,
) -> bytes:
    image = Image.open(
        BytesIO(
            png_bytes
        )
    ).convert(
        "RGB"
    )

    output = BytesIO()

    image.save(
        output,
        format="PDF",
        resolution=150.0,
    )

    return output.getvalue()
